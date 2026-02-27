from __future__ import annotations

import csv
import json
import re
import shutil
import subprocess
import time
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

ARCHIVE_PAGE_URL = "https://kaspi.kz/mc/#/orders-new?status=ARCHIVED"
ASTANA_TZ = ZoneInfo("Asia/Almaty")
DEFAULT_STORE_LABELS = ("Universal", "Acmewear", "store-d", "Store-C", "STORE-B")

_MANIFEST_FIELDS = [
    "run_id",
    "store_label",
    "window_index",
    "block_index",
    "start_date",
    "end_date",
    "status",
    "attempt",
    "source_path",
    "dest_path",
    "copied_path",
    "exported_at",
    "error",
]


@dataclass(frozen=True)
class DateBlock:
    index: int
    start: date
    end: date


def _parse_date(value: Any) -> date:
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    if not text:
        raise ValueError("date value is required")
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
        return date.fromisoformat(text)
    if re.fullmatch(r"\d{2}\.\d{2}\.\d{4}", text):
        dd, mm, yyyy = text.split(".")
        return date(int(yyyy), int(mm), int(dd))
    raise ValueError(f"unsupported date format: {text}")


def _to_kz_date_text(value: date) -> str:
    return value.strftime("%d.%m.%Y")


def build_date_blocks(start_date: date | str, end_date: date | str, *, block_days: int = 90) -> list[DateBlock]:
    if block_days <= 0:
        raise ValueError("block_days must be > 0")

    start = _parse_date(start_date)
    end = _parse_date(end_date)
    if end < start:
        raise ValueError("end_date must be >= start_date")

    blocks: list[DateBlock] = []
    cursor = start
    idx = 1
    span = timedelta(days=block_days - 1)
    while cursor <= end:
        block_end = min(cursor + span, end)
        blocks.append(DateBlock(index=idx, start=cursor, end=block_end))
        cursor = block_end + timedelta(days=1)
        idx += 1
    return blocks


def normalize_store_label(value: Any) -> str:
    text = str(value or "").strip()
    norm = re.sub(r"[^a-z0-9]+", "", text.lower())
    if not norm:
        return text

    aliases = {
        "universal": "Universal",
        "acmewear": "Acmewear",
        "store-d": "store-d",
        "store-c": "Store-C",
        "storeb": "STORE-B",
    }
    return aliases.get(norm, text)


def parse_store_labels(value: str | list[str] | None, *, window_count: int) -> list[str]:
    if isinstance(value, list):
        raw = [str(v).strip() for v in value if str(v).strip()]
    else:
        text = str(value or "").strip()
        raw = [chunk.strip() for chunk in text.split(",") if chunk.strip()] if text else list(DEFAULT_STORE_LABELS)
    normalized = [normalize_store_label(v) for v in raw]
    if len(normalized) < window_count:
        for idx in range(len(normalized) + 1, window_count + 1):
            normalized.append(f"Window{idx}")
    return normalized[:window_count]


def _slug_store_label(label: str) -> str:
    raw = str(label or "").strip().lower()
    slug = re.sub(r"[^a-z0-9]+", "-", raw).strip("-")
    return slug or "store"


def build_export_filename(store_label: str, start_date: date, end_date: date) -> str:
    return f"{_slug_store_label(store_label)}_archive_{start_date.isoformat()}_{end_date.isoformat()}.xlsx"


def _is_temporary_download(path: Path) -> bool:
    name = path.name.lower()
    return any(
        name.endswith(ext)
        for ext in (".crdownload", ".download", ".part", ".tmp")
    )


def detect_new_completed_xlsx(downloads_dir: Path, known_paths: set[Path]) -> Path | None:
    if not downloads_dir.exists():
        return None
    known = {p.resolve() for p in known_paths}

    candidates: list[Path] = []
    for path in downloads_dir.iterdir():
        if path.is_dir():
            continue
        if _is_temporary_download(path):
            continue
        if path.suffix.lower() != ".xlsx":
            continue
        if path.name.startswith("~$"):
            continue
        resolved = path.resolve()
        if resolved in known:
            continue
        try:
            size = resolved.stat().st_size
        except OSError:
            continue
        if size <= 0:
            continue
        candidates.append(resolved)

    if not candidates:
        return None
    candidates.sort(key=lambda p: p.stat().st_mtime, reverse=True)
    return candidates[0]


def wait_for_new_download(
    downloads_dir: Path,
    known_paths: set[Path],
    *,
    timeout_seconds: int,
    poll_seconds: float = 1.0,
) -> Path:
    deadline = time.time() + max(timeout_seconds, 1)
    while time.time() <= deadline:
        candidate = detect_new_completed_xlsx(downloads_dir, known_paths)
        if candidate is not None:
            size1 = candidate.stat().st_size
            time.sleep(0.8)
            if candidate.exists():
                size2 = candidate.stat().st_size
                if size1 > 0 and size1 == size2:
                    return candidate
        time.sleep(max(poll_seconds, 0.1))
    raise TimeoutError(f"Timed out waiting for new xlsx in {downloads_dir}")


def validate_export_workbook(path: Path) -> dict[str, Any]:
    out: dict[str, Any] = {
        "ok": False,
        "sheet_names": [],
        "rows_detected": 0,
        "error": "",
    }
    if not path.exists():
        out["error"] = "file_missing"
        return out

    wb = None
    try:
        wb = load_workbook(filename=str(path), read_only=True, data_only=True)
        out["sheet_names"] = list(wb.sheetnames)
        if not wb.sheetnames:
            out["error"] = "no_sheets"
            return out

        ws = wb[wb.sheetnames[0]]
        header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), tuple())
        header_has_values = any(str(v).strip() for v in header_row if v is not None)
        if not header_has_values:
            out["error"] = "empty_header"
            return out

        rows_detected = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            if any(str(v).strip() for v in row if v is not None):
                rows_detected += 1
        out["rows_detected"] = rows_detected
        out["ok"] = rows_detected > 0
        if not out["ok"]:
            out["error"] = "no_data_rows"
        return out
    except Exception as exc:
        out["error"] = f"openpyxl_error:{exc}"
        return out
    finally:
        if wb is not None:
            wb.close()


def append_manifest_row(manifest_path: Path, row: dict[str, Any]) -> None:
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    exists = manifest_path.exists()
    with manifest_path.open("a", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=_MANIFEST_FIELDS)
        if not exists:
            writer.writeheader()
        payload = {key: str(row.get(key, "")) for key in _MANIFEST_FIELDS}
        writer.writerow(payload)


def load_manifest_success_keys(manifest_path: Path) -> set[tuple[str, str, str]]:
    if not manifest_path.exists():
        return set()
    out: set[tuple[str, str, str]] = set()
    with manifest_path.open("r", newline="", encoding="utf-8") as fh:
        for row in csv.DictReader(fh):
            if str(row.get("status", "")).strip().lower() != "success":
                continue
            store = str(row.get("store_label", "")).strip()
            start = str(row.get("start_date", "")).strip()
            end = str(row.get("end_date", "")).strip()
            if store and start and end:
                out.add((store, start, end))
    return out


def append_journal_event(journal_path: Path, event: dict[str, Any]) -> None:
    journal_path.parent.mkdir(parents=True, exist_ok=True)
    with journal_path.open("a", encoding="utf-8") as fh:
        fh.write(json.dumps(event, ensure_ascii=False) + "\n")


def _run_osascript_applescript(lines: list[str]) -> str:
    cmd = ["osascript"]
    for line in lines:
        cmd.extend(["-e", line])
    proc = subprocess.run(cmd, capture_output=True, text=True)
    if proc.returncode != 0:
        raise RuntimeError((proc.stderr or proc.stdout).strip() or "osascript failed")
    return proc.stdout.strip()


def _run_osascript_jxa(script: str) -> str:
    proc = subprocess.run(["osascript", "-l", "JavaScript"], input=script, capture_output=True, text=True)
    if proc.returncode != 0:
        raise RuntimeError((proc.stderr or proc.stdout).strip() or "osascript jxa failed")
    return proc.stdout.strip()


def list_chrome_archive_windows() -> list[dict[str, Any]]:
    script = r'''
const Chrome = Application('Google Chrome');
const windows = Chrome.windows();
const out = [];
for (let i = 0; i < windows.length; i++) {
  const w = windows[i];
  const t = w.activeTab();
  out.push({
    index: i + 1,
    title: String(t.title() || ''),
    url: String(t.url() || ''),
  });
}
JSON.stringify(out);
'''
    raw = _run_osascript_jxa(script)
    data = json.loads(raw or "[]")
    if not isinstance(data, list):
        return []
    return data


def check_appleevent_javascript_enabled(window_index: int = 1) -> tuple[bool, str]:
    try:
        result = execute_tab_javascript(window_index, "document.title")
        return True, result
    except Exception as exc:
        return False, str(exc)


def execute_tab_javascript(window_index: int, js_code: str) -> str:
    script = f"""
const Chrome = Application('Google Chrome');
const jsCode = {json.dumps(js_code, ensure_ascii=False)};
const out = Chrome.windows[{int(window_index) - 1}].activeTab().execute({{javascript: jsCode}});
if (typeof out === 'undefined' || out === null) {{
  '';
}} else if (typeof out === 'string') {{
  out;
}} else {{
  JSON.stringify(out);
}}
""".strip()
    return _run_osascript_jxa(script)


def _build_export_js(start_date: date, end_date: date) -> str:
    start_text = _to_kz_date_text(start_date)
    end_text = _to_kz_date_text(end_date)
    return f"""
(() => {{
  const result = {{
    ok: false,
    dateInputs: 0,
    applyClicked: false,
    exportClicked: false,
    warnings: [],
    start: {json.dumps(start_text)},
    end: {json.dumps(end_text)}
  }};

  const visible = (el) => {{
    if (!el) return false;
    const st = window.getComputedStyle(el);
    const r = el.getBoundingClientRect();
    return st && st.display !== 'none' && st.visibility !== 'hidden' && r.width > 0 && r.height > 0;
  }};

  const normalizedText = (el) => String(el?.innerText || el?.textContent || el?.value || '').trim().toLowerCase();
  const dateValueRe = /^\\d{{2}}\\.\\d{{2}}\\.\\d{{4}}$/;

  const inputs = Array.from(document.querySelectorAll('input')).filter(visible);
  const candidates = inputs.filter((el) => {{
    const meta = `${{el.type || ''}} ${{el.placeholder || ''}} ${{el.name || ''}} ${{el.id || ''}} ${{el.className || ''}}`.toLowerCase();
    return el.type === 'date'
      || /дд\\.мм\\.гггг|dd\\.mm|yyyy|date|период|period/.test(meta)
      || dateValueRe.test(String(el.value || '').trim());
  }});

  candidates.sort((a, b) => {{
    const ar = a.getBoundingClientRect();
    const br = b.getBoundingClientRect();
    if (Math.abs(ar.top - br.top) > 2) return ar.top - br.top;
    return ar.left - br.left;
  }});

  const selected = candidates.slice(0, 2);
  result.dateInputs = selected.length;

  const setValue = (el, value) => {{
    el.focus();
    el.value = value;
    ['input', 'change', 'keyup', 'blur'].forEach((name) => {{
      el.dispatchEvent(new Event(name, {{ bubbles: true }}));
    }});
  }};

  if (selected[0]) setValue(selected[0], result.start);
  if (selected[1]) setValue(selected[1], result.end);

  const clickable = Array.from(document.querySelectorAll('button,a,span,div')).filter(visible);

  const clickByPattern = (pattern) => {{
    for (const el of clickable) {{
      const txt = normalizedText(el);
      if (!txt) continue;
      if (pattern.test(txt)) {{
        try {{
          el.click();
          return true;
        }} catch (_err) {{}}
      }}
    }}
    return false;
  }};

  result.applyClicked = clickByPattern(/применить|показать|найти|обновить/);
  result.exportClicked = clickByPattern(/экспорт|скачать|выгруз|xlsx|xls/);

  result.ok = result.dateInputs >= 2 && result.exportClicked;
  if (!result.ok) result.warnings.push('date or export controls were not detected reliably');
  return JSON.stringify(result);
}})();
""".strip()


def _collect_identity(window_index: int) -> dict[str, Any]:
    js = """
(() => {
  const out = {
    title: document.title || '',
    url: location.href || '',
    text_hint: '',
    storage_hints: {}
  };
  const pick = (sel) => {
    const el = document.querySelector(sel);
    if (!el) return '';
    return String(el.textContent || el.innerText || '').trim();
  };
  out.text_hint = pick('header') || pick('.header') || pick('.merchant') || '';

  const tryRead = (store, key) => {
    try { return store.getItem(key) || ''; } catch (_e) { return ''; }
  };

  const keys = [];
  try { keys.push(...Object.keys(localStorage)); } catch (_e) {}
  try { keys.push(...Object.keys(sessionStorage)); } catch (_e) {}

  for (const key of keys) {
    if (!/merchant|store|shop|seller|account/i.test(key)) continue;
    const v = tryRead(localStorage, key) || tryRead(sessionStorage, key);
    if (!v) continue;
    if (String(v).length > 240) continue;
    out.storage_hints[key] = String(v);
  }

  return JSON.stringify(out);
})();
""".strip()
    try:
        raw = execute_tab_javascript(window_index, js)
        parsed = json.loads(raw or "{}")
        if isinstance(parsed, dict):
            return parsed
    except Exception:
        pass
    return {}


def _write_report(
    report_path: Path,
    *,
    run_id: str,
    start_date: date,
    end_date: date,
    blocks: list[DateBlock],
    windows: list[dict[str, Any]],
    assignments: list[dict[str, Any]],
    summary: dict[str, Any],
) -> None:
    report_path.parent.mkdir(parents=True, exist_ok=True)
    lines = [
        f"# Kaspi Archive Export Report — {run_id}",
        "",
        "## Scope",
        f"- Period: {start_date.isoformat()} .. {end_date.isoformat()}",
        f"- Block size: 90 days (inclusive windows)",
        f"- Blocks: {len(blocks)}",
        f"- Expected exports: {len(blocks) * len(assignments)}",
        "",
        "## Windows",
    ]
    for win in windows:
        lines.append(f"- window {win.get('index')}: {win.get('url')}")
    lines.append("")
    lines.append("## Store Assignment")
    for item in assignments:
        lines.append(
            f"- window {item.get('window_index')} -> {item.get('store_label')}"
            f" (identity_hint={item.get('identity_hint', '')})"
        )
    lines.append("")
    lines.append("## Result")
    lines.append(f"- status: {summary.get('status')}")
    lines.append(f"- exports_expected: {summary.get('exports_expected')}")
    lines.append(f"- exports_succeeded: {summary.get('exports_succeeded')}")
    lines.append(f"- exports_failed: {summary.get('exports_failed')}")
    if summary.get("errors"):
        lines.append("- errors:")
        for err in summary["errors"]:
            lines.append(f"  - {err}")

    report_path.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")


def _copy_run_folder(src: Path, copy_dst_root: Path) -> Path:
    copy_dst_root.mkdir(parents=True, exist_ok=True)
    target = copy_dst_root / src.name
    if target.exists():
        suffix = datetime.now(ASTANA_TZ).strftime("%Y%m%d_%H%M%S")
        target = copy_dst_root / f"{src.name}_{suffix}"
    shutil.copytree(src, target)
    return target


def run_kaspi_archive_sales_export(
    *,
    start_date: str,
    end_date: str,
    block_days: int,
    window_count: int,
    store_labels: str,
    downloads_dir: str,
    output_root: str,
    copy_dst: str,
    timeout_seconds: int,
    retries: int,
    dry_run: bool,
    confirm: bool,
) -> dict[str, Any]:
    start = _parse_date(start_date)
    end = _parse_date(end_date)
    blocks = build_date_blocks(start, end, block_days=block_days)

    windows = list_chrome_archive_windows()
    run_id = datetime.now(ASTANA_TZ).strftime("%Y%m%d_%H%M%S")
    summary: dict[str, Any] = {
        "run_id": run_id,
        "status": "dry_run" if dry_run or not confirm else "success",
        "stores_total": window_count,
        "blocks_total": len(blocks),
        "exports_expected": 0,
        "exports_succeeded": 0,
        "exports_failed": 0,
        "run_dir": "",
        "copy_dir": "",
        "manifest_path": "",
        "journal_path": "",
        "report_path": "",
        "errors": [],
    }

    if len(windows) < window_count:
        summary["status"] = "preflight_failed"
        summary["errors"].append(f"expected {window_count} Chrome windows, got {len(windows)}")
        return summary

    target_windows = windows[:window_count]
    for item in target_windows:
        url = str(item.get("url") or "")
        if ARCHIVE_PAGE_URL not in url:
            summary["status"] = "preflight_failed"
            summary["errors"].append(
                f"window {item.get('index')} is not on archive page: {url}"
            )

    js_ok, js_msg = check_appleevent_javascript_enabled(int(target_windows[0].get("index", 1)))
    if not js_ok:
        summary["status"] = "preflight_failed"
        summary["errors"].append(
            "Chrome AppleEvent JavaScript is disabled. Enable View > Developer > Allow JavaScript from Apple Events."
        )
        summary["errors"].append(js_msg)

    labels = parse_store_labels(store_labels, window_count=window_count)
    assignments: list[dict[str, Any]] = []
    for idx, win in enumerate(target_windows):
        label = labels[idx]
        identity = _collect_identity(int(win.get("index", idx + 1))) if js_ok else {}
        assignments.append(
            {
                "window_index": int(win.get("index", idx + 1)),
                "store_label": label,
                "identity": identity,
                "identity_hint": identity.get("text_hint", "") if isinstance(identity, dict) else "",
            }
        )

    summary["exports_expected"] = len(assignments) * len(blocks)

    if summary["status"] == "preflight_failed":
        return summary

    if dry_run or not confirm:
        return summary

    output_root_path = Path(output_root)
    run_dir = output_root_path / run_id
    manifest_path = run_dir / "manifest.csv"
    journal_path = run_dir / "journal.jsonl"
    report_path = run_dir / "run_report.md"
    downloads_path = Path(downloads_dir).expanduser()

    run_dir.mkdir(parents=True, exist_ok=True)
    summary["run_dir"] = str(run_dir)
    summary["manifest_path"] = str(manifest_path)
    summary["journal_path"] = str(journal_path)
    summary["report_path"] = str(report_path)

    success_keys = load_manifest_success_keys(manifest_path)

    for assignment in assignments:
        store_label = assignment["store_label"]
        window_index = assignment["window_index"]
        store_dir = run_dir / store_label / "raw"
        store_dir.mkdir(parents=True, exist_ok=True)
        meta_path = run_dir / store_label / "meta" / "store_identity.json"
        meta_path.parent.mkdir(parents=True, exist_ok=True)
        meta_path.write_text(json.dumps(assignment.get("identity", {}), ensure_ascii=False, indent=2), encoding="utf-8")

        for block in blocks:
            key = (store_label, block.start.isoformat(), block.end.isoformat())
            if key in success_keys:
                continue

            dest_name = build_export_filename(store_label, block.start, block.end)
            dest_path = store_dir / dest_name

            attempts_total = max(int(retries), 0) + 1
            last_error = ""
            for attempt in range(1, attempts_total + 1):
                exported_at = datetime.now(ASTANA_TZ).isoformat()
                source_path = ""
                copied_path = ""
                try:
                    known = {p.resolve() for p in downloads_path.glob("*.xlsx")}
                    js_result_raw = execute_tab_javascript(window_index, _build_export_js(block.start, block.end))
                    js_result = json.loads(js_result_raw or "{}")
                    if not bool(js_result.get("exportClicked")):
                        raise RuntimeError(f"export button not clicked: {js_result}")

                    downloaded = wait_for_new_download(downloads_path, known, timeout_seconds=timeout_seconds)
                    source_path = str(downloaded)

                    if dest_path.exists():
                        stamp = datetime.now(ASTANA_TZ).strftime("%H%M%S")
                        dest_path = dest_path.with_name(dest_path.stem + f"_{stamp}" + dest_path.suffix)

                    shutil.move(str(downloaded), str(dest_path))
                    valid = validate_export_workbook(dest_path)
                    if not valid.get("ok"):
                        raise RuntimeError(f"downloaded workbook validation failed: {valid}")

                    append_manifest_row(
                        manifest_path,
                        {
                            "run_id": run_id,
                            "store_label": store_label,
                            "window_index": window_index,
                            "block_index": block.index,
                            "start_date": block.start.isoformat(),
                            "end_date": block.end.isoformat(),
                            "status": "success",
                            "attempt": attempt,
                            "source_path": source_path,
                            "dest_path": str(dest_path),
                            "copied_path": copied_path,
                            "exported_at": exported_at,
                            "error": "",
                        },
                    )
                    append_journal_event(
                        journal_path,
                        {
                            "ts": exported_at,
                            "event": "export_success",
                            "run_id": run_id,
                            "store_label": store_label,
                            "window_index": window_index,
                            "block_index": block.index,
                            "start_date": block.start.isoformat(),
                            "end_date": block.end.isoformat(),
                            "source_path": source_path,
                            "dest_path": str(dest_path),
                            "attempt": attempt,
                        },
                    )
                    summary["exports_succeeded"] += 1
                    success_keys.add(key)
                    last_error = ""
                    break
                except Exception as exc:
                    last_error = str(exc)
                    if attempt >= attempts_total:
                        append_manifest_row(
                            manifest_path,
                            {
                                "run_id": run_id,
                                "store_label": store_label,
                                "window_index": window_index,
                                "block_index": block.index,
                                "start_date": block.start.isoformat(),
                                "end_date": block.end.isoformat(),
                                "status": "failed",
                                "attempt": attempt,
                                "source_path": source_path,
                                "dest_path": str(dest_path),
                                "copied_path": copied_path,
                                "exported_at": exported_at,
                                "error": last_error,
                            },
                        )
                        append_journal_event(
                            journal_path,
                            {
                                "ts": exported_at,
                                "event": "export_failed",
                                "run_id": run_id,
                                "store_label": store_label,
                                "window_index": window_index,
                                "block_index": block.index,
                                "start_date": block.start.isoformat(),
                                "end_date": block.end.isoformat(),
                                "error": last_error,
                                "attempt": attempt,
                            },
                        )
                        summary["exports_failed"] += 1
                        summary["errors"].append(
                            f"{store_label} {block.start.isoformat()}..{block.end.isoformat()}: {last_error}"
                        )

    if summary["exports_failed"] > 0:
        summary["status"] = "partial"

    try:
        copied_dir = _copy_run_folder(run_dir, Path(copy_dst).expanduser())
        summary["copy_dir"] = str(copied_dir)
    except Exception as exc:
        summary["status"] = "partial"
        summary["errors"].append(f"copy_failed: {exc}")

    _write_report(
        report_path,
        run_id=run_id,
        start_date=start,
        end_date=end,
        blocks=blocks,
        windows=target_windows,
        assignments=assignments,
        summary=summary,
    )

    return summary
