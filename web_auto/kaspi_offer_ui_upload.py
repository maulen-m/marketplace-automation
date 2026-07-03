from __future__ import annotations

import argparse
import csv
import json
import os
import re
import subprocess
import tempfile
import time
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from .kaspi_forbidden_cards import (
    FORBIDDEN_KASPI_OFFER_PRODUCT_CODES,
    FORBIDDEN_KASPI_OFFER_URL_FRAGMENTS,
    forbidden_card_reason_for_row,
)

ASTANA_TZ = ZoneInfo("Asia/Almaty")
DEFAULT_ENTRY_URL = "https://kaspi.kz/mc/#/add-product/v2"

REQUIRED_COLUMNS = [
    "upload_enabled",
    "store_code",
    "merchant_id",
    "variant_url",
    "expected_kaspi_heading",
    "size_rus",
    "merchant_sku_article",
    "merchant_offer_name",
    "price_kzt",
    "stock_pp1",
    "stock_pp2",
]

OPTIONAL_COLUMNS = [
    "barcode",
    "product_code",
    "resolved_sku_key",
    "resolved_sku_id",
    "final_attached_size",
    "ui_entry_url",
    "ingest_method",
    "notes",
]

_STORE_ALIASES = {
    "universal": "UNIVERSAL",
    "storeb": "STOREB",
    "store-b": "STOREB",
    "acmewear": "ACMEWEAR",
    "only-fit": "ACMEWEAR",
    "onlyfeed": "ACMEWEAR",
    "only-feed": "ACMEWEAR",
    "store-c": "MELVIS",
    "store-d": "11KZ",
}

ALLOWED_UPLOAD_STORE_CODES = {"UNIVERSAL", "STOREB"}
ALLOWED_UPLOAD_MERCHANT_IDS_BY_STORE = {
    "UNIVERSAL": "30000001",
    "STOREB": "30000002",
}
FORBIDDEN_UPLOAD_STORE_CODES = {"ACMEWEAR"}
FORBIDDEN_UPLOAD_MERCHANT_IDS = {"30137883"}
FORBIDDEN_UPLOAD_PRODUCT_CODES = FORBIDDEN_KASPI_OFFER_PRODUCT_CODES
FORBIDDEN_UPLOAD_URL_FRAGMENTS = FORBIDDEN_KASPI_OFFER_URL_FRAGMENTS

_COLOR_ALIASES = {
    "black": {"black", "chernyi", "черный", "чёрный"},
    "white": {"white", "belyi", "белый"},
    "grey": {"grey", "gray", "seryi", "серый"},
}


def _norm_text(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().split())


def _to_int(value: Any) -> int:
    if value is None:
        raise ValueError("empty numeric value")
    text = str(value).strip()
    if not text:
        raise ValueError("empty numeric value")
    text = text.replace(",", ".")
    # Handle Excel style float-ish values.
    return int(float(text))


def _merchant_id_text(value: Any) -> str:
    try:
        return str(_to_int(value))
    except Exception:
        return re.sub(r"\D", "", _norm_text(value))


def _normalize_barcode(value: Any) -> str:
    text = _norm_text(value)
    if not text:
        return ""
    digits = re.sub(r"\D", "", text)
    if 8 <= len(digits) <= 14:
        return digits
    return ""


def _normalize_size_rus(value: Any) -> str:
    text = _norm_text(value)
    if not text:
        return ""
    num_text = text.replace(",", ".")
    if re.fullmatch(r"\d+(\.\d+)?", num_text):
        return str(int(float(num_text)))
    return text.upper()


def _extract_offer_code_from_variant_url(variant_url: Any) -> str:
    text = _norm_text(variant_url)
    if not text:
        return ""
    m = re.search(r"-(\d{6,})(?:[/?#]|$)", text)
    return m.group(1) if m else ""


def _product_code_for_search(row: dict[str, Any]) -> str:
    product_code = re.sub(r"\D", "", _norm_text(row.get("product_code")))
    if product_code:
        return product_code
    return _extract_offer_code_from_variant_url(row.get("variant_url"))


def _is_unsafe_link_catalog_entry_url(value: Any) -> bool:
    text = _norm_text(value).lower()
    return "link-catalog" in text and "code=" not in text


def _forbidden_upload_product_reason(row: dict[str, Any]) -> str:
    return forbidden_card_reason_for_row(row)


def _infer_color_token(*values: Any) -> str:
    haystack = " ".join(_norm_text(v).lower() for v in values if _norm_text(v))
    if not haystack:
        return ""
    for canonical, aliases in _COLOR_ALIASES.items():
        if any(alias in haystack for alias in aliases):
            return canonical
    return ""


def _normalize_color_token(value: Any) -> str:
    text = _norm_text(value).lower()
    if not text:
        return ""
    for canonical, aliases in _COLOR_ALIASES.items():
        if text == canonical or text in aliases:
            return canonical
    return ""


def _has_semantic_partner_article(article: Any, resolved_sku_key: Any, resolved_sku_id: Any) -> bool:
    article_text = _norm_text(article)
    if not article_text:
        return False
    sku_id = _norm_text(resolved_sku_id)
    if sku_id and article_text.startswith(sku_id):
        return True
    sku_key = _norm_text(resolved_sku_key)
    if sku_key and article_text.startswith(sku_key):
        return True
    return False


def normalize_store_code(value: Any) -> str:
    raw = str(value or "").strip()
    if not raw:
        return ""
    key = raw.lower().replace("_", "-")
    key = re.sub(r"\s+", "", key)
    if key in _STORE_ALIASES:
        return _STORE_ALIASES[key]
    # Preserve canonical uppercase for unknown but explicit codes.
    return raw.upper()


def parse_store_codes(value: str | list[str] | None) -> list[str]:
    if value is None:
        return []
    if isinstance(value, list):
        parts = [str(v).strip() for v in value if str(v).strip()]
    else:
        parts = [chunk.strip() for chunk in str(value).split(",") if chunk.strip()]
    out: list[str] = []
    for part in parts:
        code = normalize_store_code(part)
        if code and code not in out:
            out.append(code)
    return out


def _truthy_yes(value: Any) -> bool:
    text = str(value or "").strip().lower()
    return text in {"1", "yes", "y", "true", "да"}


def load_offer_upload_rows(
    workbook_path: str | Path,
    *,
    sheet_name: str = "offer upload",
    store_codes: list[str] | None = None,
) -> list[dict[str, Any]]:
    path = Path(workbook_path)
    wb = load_workbook(path, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"missing sheet: {sheet_name}")
    ws = wb[sheet_name]
    headers = [_norm_text(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)]
    idx = {h: i + 1 for i, h in enumerate(headers) if h}
    missing = [name for name in REQUIRED_COLUMNS if name not in idx]
    if missing:
        raise ValueError(f"missing required columns: {', '.join(missing)}")

    filter_codes = set(store_codes or [])
    rows: list[dict[str, Any]] = []
    for r in range(2, ws.max_row + 1):
        upload_enabled = ws.cell(r, idx["upload_enabled"]).value
        if not _truthy_yes(upload_enabled):
            continue

        store_code = normalize_store_code(ws.cell(r, idx["store_code"]).value)
        if filter_codes and store_code not in filter_codes:
            continue

        row: dict[str, Any] = {"row_number": r, "upload_enabled": "yes", "store_code": store_code}
        for name in REQUIRED_COLUMNS:
            if name in {"upload_enabled", "store_code"}:
                continue
            row[name] = ws.cell(r, idx[name]).value

        # Optional columns used by runtime and validations.
        raw_color = ws.cell(r, idx.get("color", 0)).value if idx.get("color") else ""
        row["color"] = _normalize_color_token(raw_color) or _infer_color_token(
            row.get("variant_url"),
            row.get("merchant_sku_article"),
            row.get("merchant_offer_name"),
            row.get("expected_kaspi_heading"),
        )
        for name in OPTIONAL_COLUMNS:
            if name in {"ui_entry_url", "ingest_method"}:
                continue
            row[name] = ws.cell(r, idx.get(name, 0)).value if idx.get(name) else ""
        row["ui_entry_url"] = ws.cell(r, idx.get("ui_entry_url", 0)).value if idx.get("ui_entry_url") else DEFAULT_ENTRY_URL
        row["ingest_method"] = ws.cell(r, idx.get("ingest_method", 0)).value if idx.get("ingest_method") else ""

        rows.append(row)
    return rows


def validate_upload_rows(
    rows: list[dict[str, Any]],
    *,
    required_store_codes: list[str],
    require_black_coverage: bool = True,
) -> dict[str, Any]:
    errors: list[str] = []
    warnings: list[str] = []
    required = [normalize_store_code(s) for s in required_store_codes if normalize_store_code(s)]

    for store in required:
        if store not in ALLOWED_UPLOAD_STORE_CODES:
            errors.append(
                f"store {store} is not allowed for offer upload; allowed stores: "
                f"{','.join(sorted(ALLOWED_UPLOAD_STORE_CODES))}"
            )

    if not rows:
        errors.append("no active rows to upload")
        return {"ok": False, "errors": errors, "warnings": warnings, "rows_total": 0}

    by_store = Counter(normalize_store_code(r.get("store_code")) for r in rows)
    for store in required:
        if by_store.get(store, 0) == 0:
            errors.append(f"store {store} has no active rows")

    # Ensure black color coverage in required stores.
    if require_black_coverage:
        for store in required:
            has_black = any(
                normalize_store_code(r.get("store_code")) == store
                and str(r.get("color") or "").strip().lower() == "black"
                for r in rows
            )
            if not has_black:
                errors.append(f"black color row missing for store {store}")

    required_fields = [
        "merchant_id",
        "variant_url",
        "expected_kaspi_heading",
        "size_rus",
        "merchant_sku_article",
        "merchant_offer_name",
        "price_kzt",
        "stock_pp1",
        "stock_pp2",
    ]
    for row in rows:
        row_no = row.get("row_number", "?")
        store_code = normalize_store_code(row.get("store_code"))
        merchant_id = _merchant_id_text(row.get("merchant_id"))
        if store_code not in ALLOWED_UPLOAD_STORE_CODES:
            errors.append(
                f"row {row_no}: store {store_code or '<empty>'} is not allowed for offer upload; "
                f"allowed stores: {','.join(sorted(ALLOWED_UPLOAD_STORE_CODES))}"
            )
        if store_code in FORBIDDEN_UPLOAD_STORE_CODES or merchant_id in FORBIDDEN_UPLOAD_MERCHANT_IDS:
            errors.append(
                f"row {row_no}: forbidden upload surface store={store_code or '<empty>'} merchant_id={merchant_id or '<empty>'}"
            )
        expected_mid_for_store = ALLOWED_UPLOAD_MERCHANT_IDS_BY_STORE.get(store_code)
        if expected_mid_for_store and merchant_id and merchant_id != expected_mid_for_store:
            errors.append(
                f"row {row_no}: merchant_id {merchant_id} does not match store {store_code} expected {expected_mid_for_store}"
            )
        for field in required_fields:
            if _norm_text(row.get(field)) == "":
                errors.append(f"row {row_no}: missing {field}")
        for field in ("price_kzt", "stock_pp1", "stock_pp2", "merchant_id"):
            try:
                _to_int(row.get(field))
            except Exception:
                errors.append(f"row {row_no}: invalid numeric {field}={row.get(field)!r}")
        product_code = re.sub(r"\D", "", _norm_text(row.get("product_code")))
        url_offer_code = _extract_offer_code_from_variant_url(row.get("variant_url"))
        if product_code and url_offer_code and product_code != url_offer_code:
            errors.append(
                f"row {row_no}: product_code {product_code} does not match variant_url code {url_offer_code}"
            )
        forbidden_reason = _forbidden_upload_product_reason(row)
        if forbidden_reason:
            errors.append(f"row {row_no}: {forbidden_reason}")
        if _is_unsafe_link_catalog_entry_url(row.get("ui_entry_url")):
            errors.append(
                f"row {row_no}: unsafe ui_entry_url link-catalog without code=; use {DEFAULT_ENTRY_URL} for URL search"
            )
        resolved_sku_key = row.get("resolved_sku_key")
        resolved_sku_id = row.get("resolved_sku_id")
        if (_norm_text(resolved_sku_key) or _norm_text(resolved_sku_id)) and not _has_semantic_partner_article(
            row.get("merchant_sku_article"),
            resolved_sku_key,
            resolved_sku_id,
        ):
            errors.append(
                f"row {row_no}: merchant_sku_article must start with resolved_sku_key/resolved_sku_id semantics"
            )
        if _norm_text(row.get("barcode")):
            warnings.append(f"row {row_no}: barcode ignored; Kaspi barcode field is intentionally left empty")

    seen: set[tuple[str, str, str]] = set()
    for row in rows:
        key = (
            normalize_store_code(row.get("store_code")),
            _norm_text(row.get("merchant_sku_article")),
            _norm_text(row.get("variant_url")),
        )
        if key in seen:
            errors.append(f"duplicate upload key: {key[0]}::{key[1]}::{key[2]}")
        else:
            seen.add(key)

    return {
        "ok": len(errors) == 0,
        "errors": errors,
        "warnings": warnings,
        "rows_total": len(rows),
        "rows_by_store": dict(by_store),
    }


def _run_osascript_jxa(script: str) -> str:
    proc = subprocess.run(["osascript", "-l", "JavaScript"], input=script, capture_output=True, text=True)
    if proc.returncode != 0:
        raise RuntimeError((proc.stderr or proc.stdout).strip() or "osascript jxa failed")
    return proc.stdout.strip()


def _execute_tab_js(window_index: int, js_code: str) -> str:
    script = f"""
const Chrome = Application('Google Chrome');
const jsCode = {json.dumps(js_code, ensure_ascii=False)};
const out = Chrome.windows[{int(window_index) - 1}].activeTab().execute({{javascript: jsCode}});
if (typeof out === 'undefined' || out === null) {{
  '';
}} else if (typeof out === 'string') {{
  out;
}} else {{
  JSON.stringify(out);
}}
""".strip()
    return _run_osascript_jxa(script)


def _set_tab_url(window_index: int, url: str) -> None:
    script = f"""
const Chrome = Application('Google Chrome');
Chrome.windows[{int(window_index) - 1}].activeTab().url = {json.dumps(url, ensure_ascii=False)};
'ok';
""".strip()
    _run_osascript_jxa(script)


def _list_chrome_windows() -> list[dict[str, Any]]:
    script = r'''
const Chrome = Application('Google Chrome');
const out = [];
const wins = Chrome.windows();
for (let i = 0; i < wins.length; i++) {
  const t = wins[i].activeTab();
  out.push({
    index: i + 1,
    title: String(t.title() || ''),
    url: String(t.url() || ''),
  });
}
JSON.stringify(out);
'''
    raw = _run_osascript_jxa(script)
    data = json.loads(raw or "[]")
    return data if isinstance(data, list) else []


def _window_bounds(window_index: int) -> tuple[int, int, int, int]:
    script = f'''
tell application "Google Chrome"
  set b to bounds of window {int(window_index)}
  return (item 1 of b as text) & "," & (item 2 of b as text) & "," & (item 3 of b as text) & "," & (item 4 of b as text)
end tell
'''.strip()
    proc = subprocess.run(["osascript"], input=script, capture_output=True, text=True)
    if proc.returncode != 0:
        raise RuntimeError((proc.stderr or proc.stdout).strip() or "failed to read Chrome window bounds")
    raw = (proc.stdout or "").strip()
    parts = [int(p.strip()) for p in raw.split(",")]
    if len(parts) != 4:
        raise RuntimeError(f"invalid bounds from osascript: {raw}")
    return parts[0], parts[1], parts[2], parts[3]


def _focus_chrome_window(window_index: int) -> int:
    script = f'''
tell application "Google Chrome"
  activate
  set index of window {int(window_index)} to 1
end tell
return "1"
'''.strip()
    proc = subprocess.run(["osascript"], input=script, capture_output=True, text=True)
    if proc.returncode == 0:
        return 1
    return int(window_index)


def _jxa_focus_and_normalize_window(window_index: int) -> dict[str, Any]:
    script = f"""
const Chrome = Application('Google Chrome');
Chrome.activate();
const wins = Chrome.windows();
const w = wins[{int(window_index) - 1}];
if (!w) throw new Error('chrome_window_not_found');
w.index = 1;
w.bounds = {{x: 20, y: 25, width: 1280, height: 1100}};
const b = w.bounds();
JSON.stringify({{
  title: String(w.activeTab().title() || ''),
  url: String(w.activeTab().url() || ''),
  bounds: {{x: Number(b.x), y: Number(b.y), width: Number(b.width), height: Number(b.height)}}
}});
""".strip()
    raw = _run_osascript_jxa(script)
    data = json.loads(raw or "{}")
    return data if isinstance(data, dict) else {}


def _coregraphics_chrome_window_id(title: str, bounds: dict[str, Any]) -> str:
    swift = r'''
import Foundation
import CoreGraphics

let env = ProcessInfo.processInfo.environment
let targetTitle = env["KASPI_CAPTURE_TITLE"] ?? ""
let targetX = Double(env["KASPI_CAPTURE_X"] ?? "") ?? 0
let targetY = Double(env["KASPI_CAPTURE_Y"] ?? "") ?? 0
let targetW = Double(env["KASPI_CAPTURE_W"] ?? "") ?? 0
let targetH = Double(env["KASPI_CAPTURE_H"] ?? "") ?? 0

guard let list = CGWindowListCopyWindowInfo(CGWindowListOption(arrayLiteral: .optionAll), kCGNullWindowID) as? [[String: Any]] else {
    exit(1)
}

var bestID = ""
var bestScore = -Double.greatestFiniteMagnitude

for w in list {
    let owner = w[kCGWindowOwnerName as String] as? String ?? ""
    if owner != "Google Chrome" { continue }
    let layer = w[kCGWindowLayer as String] as? Int ?? 0
    if layer != 0 { continue }
    guard let windowNumber = w[kCGWindowNumber as String] else { continue }
    guard let b = w[kCGWindowBounds as String] as? [String: Any] else { continue }
    let name = w[kCGWindowName as String] as? String ?? ""
    let x = Double("\(b["X"] ?? "0")") ?? 0
    let y = Double("\(b["Y"] ?? "0")") ?? 0
    let width = Double("\(b["Width"] ?? "0")") ?? 0
    let height = Double("\(b["Height"] ?? "0")") ?? 0
    if width < 500 || height < 500 { continue }

    var score = 0.0
    if name == targetTitle { score += 10000 }
    if targetTitle != "" && name.contains(targetTitle) { score += 2500 }
    if w[kCGWindowIsOnscreen as String] != nil { score += 1000 }
    score -= abs(x - targetX)
    score -= abs(y - targetY)
    score -= abs(width - targetW) / 2
    score -= abs(height - targetH) / 2

    if score > bestScore {
        bestScore = score
        bestID = "\(windowNumber)"
    }
}

if bestID != "" {
    print(bestID)
}
'''.strip()
    env = os.environ.copy()
    env["KASPI_CAPTURE_TITLE"] = str(title or "")
    env["KASPI_CAPTURE_X"] = str(bounds.get("x", ""))
    env["KASPI_CAPTURE_Y"] = str(bounds.get("y", ""))
    env["KASPI_CAPTURE_W"] = str(bounds.get("width", ""))
    env["KASPI_CAPTURE_H"] = str(bounds.get("height", ""))
    with tempfile.NamedTemporaryFile("w", suffix=".swift", delete=True) as fh:
        fh.write(swift)
        fh.flush()
        proc = subprocess.run(["swift", fh.name], capture_output=True, text=True, env=env)
    if proc.returncode != 0:
        return ""
    return (proc.stdout or "").strip().splitlines()[-1].strip() if (proc.stdout or "").strip() else ""


def _capture_window_screenshot(window_index: int, out_path: Path) -> None:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        meta = _jxa_focus_and_normalize_window(window_index)
        bounds = meta.get("bounds") if isinstance(meta.get("bounds"), dict) else {}
        time.sleep(0.2)
        cg_window_id = _coregraphics_chrome_window_id(str(meta.get("title") or ""), bounds)
        if cg_window_id:
            proc = subprocess.run(
                ["screencapture", "-x", "-l", cg_window_id, str(out_path)],
                capture_output=True,
                text=True,
            )
            if proc.returncode == 0:
                return
        if bounds:
            left = int(float(bounds.get("x", 0)))
            top = int(float(bounds.get("y", 0)))
            width = max(1, int(float(bounds.get("width", 1))))
            height = max(1, int(float(bounds.get("height", 1))))
        else:
            screenshot_window_index = _focus_chrome_window(window_index)
            left, top, right, bottom = _window_bounds(screenshot_window_index)
            width = max(1, right - left)
            height = max(1, bottom - top)
        rect = f"{left},{top},{width},{height}"
        proc = subprocess.run(
            ["screencapture", "-x", "-R", rect, str(out_path)],
            capture_output=True,
            text=True,
        )
        if proc.returncode != 0:
            raise RuntimeError((proc.stderr or proc.stdout).strip() or "screencapture failed")
    except Exception:
        # Screenshot failures should not block upload run.
        return


def _step_page_state_js() -> str:
    return r"""
(() => {
  const visible = (el) => {
    if (!el) return false;
    const st = window.getComputedStyle(el);
    const r = el.getBoundingClientRect();
    return st.display !== 'none' && st.visibility !== 'hidden' && r.width > 0 && r.height > 0;
  };
  const norm = (v) => String(v || '').trim().toLowerCase();
  const contextText = (el) => {
    const parts = [el.placeholder, el.name, el.id, el.getAttribute && el.getAttribute('aria-label')];
    let p = el.parentElement;
    for (let i = 0; i < 3 && p; i++) {
      parts.push(p.innerText || '');
      p = p.parentElement;
    }
    return String(parts.join(' ') || '').trim();
  };
  const inputs = Array.from(document.querySelectorAll('input')).filter(visible)
    .filter(el => el.type !== 'hidden' && !el.disabled && !el.readOnly)
    .map((el, i) => ({
      index: i,
      type: String(el.type || ''),
      value: String(el.value || ''),
      placeholder: String(el.placeholder || ''),
      label_context: contextText(el).slice(0, 300),
    }));
  const buttons = Array.from(document.querySelectorAll('button')).filter(visible)
    .map((el, i) => ({ index: i, text: String(el.innerText || '').trim().slice(0, 120) }));
	  const bodyText = String(document.body?.innerText || '');
	  const bodyNorm = norm(bodyText);
	  const href = String(location.href || '');
	  const isAddProductPage = href.toLowerCase().includes('#/add-product/v2');
	  const isOrdersPage = /#\/orders-new/i.test(href);
	  const successModalPresent = /ваш товар успешно добавлен/i.test(bodyText);
	  const articleInput = inputs.find(inp => /(^|\s|[\n\r])артикул($|\s|[\n\r])/i.test(inp.label_context));
	  const nameInput = inputs.find(inp => /наименование товара|название товара/i.test(inp.label_context));
	  const marketplaceUrlInArticle = !!(articleInput && /https?:\/\/kaspi\.kz\/shop\/p\//i.test(articleInput.value));
	  const sellerIdentityForm = bodyNorm.includes('информация о товаре') && !!articleInput && !!nameInput;
  const searchContextPresent = /присоединиться|существующей карточк|ссылк|url|поиск/i.test(bodyText);
  const sizeOptions = Array.from(document.querySelectorAll('.matrix__values, .matrix__values *'))
    .filter(el => visible(el) && String(el.innerText || '').trim())
    .map(el => String(el.innerText || '').trim())
    .slice(0, 40);
  return JSON.stringify({
	    ok: true,
	    url: location.href,
	    title: document.title,
	    page_kind: isOrdersPage ? 'orders'
	      : sellerIdentityForm && !searchContextPresent ? 'seller_identity_form'
	      : isAddProductPage ? 'add_product_or_search'
	      : 'unknown_or_search',
	    inputs,
	    buttons,
	    size_options: sizeOptions,
	    hazards: {
	      marketplace_url_in_article_field: marketplaceUrlInArticle,
	      seller_identity_form_without_search_context: sellerIdentityForm && !searchContextPresent,
	      success_modal_present: successModalPresent,
	      not_add_product_page: !isAddProductPage,
	      orders_page: isOrdersPage,
	    },
	    body_excerpt: bodyText.slice(0, 1200),
	  });
	})();
	""".strip()


def _step_assert_not_unsafe_identity_form_js() -> str:
    return r"""
(() => {
  const raw = (() => {
    const visible = (el) => {
      if (!el) return false;
      const st = window.getComputedStyle(el);
      const r = el.getBoundingClientRect();
      return st.display !== 'none' && st.visibility !== 'hidden' && r.width > 0 && r.height > 0;
    };
    const norm = (v) => String(v || '').trim().toLowerCase();
    const contextText = (el) => {
      const parts = [el.placeholder, el.name, el.id, el.getAttribute && el.getAttribute('aria-label')];
      let p = el.parentElement;
      for (let i = 0; i < 3 && p; i++) {
        parts.push(p.innerText || '');
        p = p.parentElement;
      }
      return String(parts.join(' ') || '').trim();
    };
    const inputs = Array.from(document.querySelectorAll('input')).filter(visible)
      .filter(el => el.type !== 'hidden' && !el.disabled && !el.readOnly);
    const bodyText = String(document.body?.innerText || '');
    const articleInput = inputs.find(el => /(^|\s|[\n\r])артикул($|\s|[\n\r])/i.test(contextText(el)));
    const nameInput = inputs.find(el => /наименование товара|название товара/i.test(contextText(el)));
    const searchContextPresent = /присоединиться|существующей карточк|ссылк|url|поиск/i.test(bodyText);
    return {
      articleInput,
      nameInput,
      bodyText,
      bodyNorm: norm(bodyText),
      searchContextPresent,
      articleValue: articleInput ? String(articleInput.value || '') : '',
    };
  })();
  const sellerIdentityForm = raw.bodyNorm.includes('информация о товаре') && !!raw.articleInput && !!raw.nameInput;
  if (sellerIdentityForm && !raw.searchContextPresent) {
    return JSON.stringify({
      ok: false,
      reason: 'unsafe_seller_identity_form_not_search',
      article_value: raw.articleValue,
      marketplace_url_in_article_field: /https?:\/\/kaspi\.kz\/shop\/p\//i.test(raw.articleValue),
    });
  }
  if (/https?:\/\/kaspi\.kz\/shop\/p\//i.test(raw.articleValue)) {
    return JSON.stringify({
      ok: false,
      reason: 'marketplace_url_in_article_field',
      article_value: raw.articleValue,
    });
  }
  return JSON.stringify({ok:true});
})();
""".strip()


def _step_close_stale_success_modal_js() -> str:
    return r"""
(() => {
  const visible = (el) => {
    if (!el) return false;
    const st = window.getComputedStyle(el);
    const r = el.getBoundingClientRect();
    return st.display !== 'none' && st.visibility !== 'hidden' && r.width > 0 && r.height > 0;
  };
  const norm = (v) => String(v || '').trim().toLowerCase();
  const all = Array.from(document.querySelectorAll('[role="dialog"], .modal, .modal-content, .v-dialog, .dialog, div'))
    .filter(visible);
  const dialog = all
    .filter(el => /ваш товар успешно добавлен/i.test(String(el.innerText || '')))
    .map(el => ({el, r: el.getBoundingClientRect()}))
    .filter(x => x.r.width >= 250 && x.r.width <= 900 && x.r.height >= 100 && x.r.height <= 500)
    .sort((a, b) => {
      return (a.r.width * a.r.height) - (b.r.width * b.r.height);
    })[0] || null;
  if (!dialog) return JSON.stringify({ok:true, closed:false, reason:'success_modal_not_present'});

  const clickable = Array.from(dialog.el.querySelectorAll('button, [role="button"], a, svg, path, span, div, i'))
    .filter(visible)
    .filter(el => !/перейти|управление товарами/i.test(norm(el.innerText)));
  const closeControl = clickable.find(el => /закрыть|close|×|x/i.test(norm(el.innerText) || norm(el.getAttribute && el.getAttribute('aria-label'))))
    || clickable.find(el => /close|modal-header__close-button/i.test(String(el.className || '')))
    || clickable
      .map(el => ({el, r: el.getBoundingClientRect()}))
      .filter(x => x.r.width <= 80 && x.r.height <= 80)
      .sort((a, b) => (b.r.right - b.r.top) - (a.r.right - a.r.top))[0]?.el
    || null;
  if (closeControl) {
    closeControl.click();
    return JSON.stringify({ok:true, closed:true, method:'close_control'});
  }

  // Some Kaspi dialogs render the X as a textless icon with no button role.
  // Click a safe point in the dialog's top-right close icon area; this is far
  // away from the "Перейти в управление товарами" action button.
  const x = dialog.r.right - 32;
  const y = dialog.r.top + 58;
  const target = document.elementFromPoint(x, y);
  if (target && !/перейти|управление товарами/i.test(norm(target.innerText))) {
    const ev = {bubbles:true, cancelable:true, view:window, clientX:x, clientY:y};
    for (const type of ['pointerdown', 'mousedown', 'pointerup', 'mouseup', 'click']) {
      try { target.dispatchEvent(new MouseEvent(type, ev)); } catch (e) {}
    }
    return JSON.stringify({
      ok:true,
      closed:true,
      method:'top_right_coordinate',
      target_tag:String(target.tagName || ''),
      target_class:String(target.className || ''),
    });
  }

  document.dispatchEvent(new KeyboardEvent('keydown', {key:'Escape', code:'Escape', bubbles:true}));
  window.dispatchEvent(new KeyboardEvent('keydown', {key:'Escape', code:'Escape', bubbles:true}));
  return JSON.stringify({ok:true, closed:true, method:'escape'});
})();
""".strip()


def _merchant_id_js() -> str:
    return r"""
(() => {
  const txt = document.body ? (document.body.innerText || '') : '';
  const m = txt.match(/ID\s*[-–:]?\s*(\d{6,})/i);
  if (m) return m[1];
  try {
    const lsMid = (localStorage.getItem('merchantUid') || '').trim();
    if (/^\d{6,}$/.test(lsMid)) return lsMid;
  } catch (e) {}
  try {
    const lastMid = (localStorage.getItem('lastSuccessUid') || '').trim();
    if (/^\d{6,}$/.test(lastMid)) return lastMid;
  } catch (e) {}
  const cands = Array.from(document.querySelectorAll('*')).map(el => String(el.textContent || '').trim()).filter(Boolean);
  for (const t of cands) {
    const mm = t.match(/ID\s*[-–:]?\s*(\d{6,})/i);
    if (mm) return mm[1];
  }
  return '';
})();
""".strip()


def parse_window_map(value: str) -> dict[str, int]:
    out: dict[str, int] = {}
    for chunk in str(value or "").split(","):
        chunk = chunk.strip()
        if not chunk:
            continue
        if ":" not in chunk:
            raise ValueError(f"invalid window map item: {chunk}")
        store, win = chunk.split(":", 1)
        code = normalize_store_code(store)
        out[code] = int(win.strip())
    return out


def list_store_window_candidates() -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for win in _list_chrome_windows():
        idx = int(win.get("index", 0))
        merchant_id = ""
        if idx > 0:
            try:
                merchant_id = _execute_tab_js(idx, _merchant_id_js()).strip()
            except Exception:
                merchant_id = ""
        out.append(
            {
                "window_index": idx,
                "merchant_id": merchant_id,
                "title": str(win.get("title") or ""),
                "url": str(win.get("url") or ""),
            }
        )
    return out


def _auto_map_windows(rows: list[dict[str, Any]]) -> dict[str, int]:
    target_mid_by_store: dict[str, str] = {}
    for row in rows:
        code = normalize_store_code(row.get("store_code"))
        if code and code not in target_mid_by_store:
            target_mid_by_store[code] = str(_to_int(row.get("merchant_id")))

    windows = _list_chrome_windows()
    resolved_by_mid: dict[str, int] = {}
    for win in windows:
        widx = int(win.get("index", 0))
        if widx <= 0:
            continue
        try:
            merchant_id = _execute_tab_js(widx, _merchant_id_js()).strip()
        except Exception:
            continue
        if merchant_id:
            resolved_by_mid[merchant_id] = widx

    out: dict[str, int] = {}
    for store, mid in target_mid_by_store.items():
        if mid in resolved_by_mid:
            out[store] = resolved_by_mid[mid]
    return out


def _choose_window_index(mapped_index: int, expected_merchant_id: str, candidates: list[dict[str, Any]]) -> int:
    expected = str(expected_merchant_id or "").strip()
    if not expected:
        return int(mapped_index)
    by_index: dict[int, str] = {}
    for c in candidates:
        try:
            idx = int(c.get("window_index", 0))
        except Exception:
            continue
        if idx <= 0:
            continue
        by_index[idx] = str(c.get("merchant_id") or "").strip()

    mapped_mid = by_index.get(int(mapped_index), "")
    if mapped_mid == expected:
        return int(mapped_index)

    for idx, mid in sorted(by_index.items()):
        if mid == expected:
            return idx
    return int(mapped_index)


def _step_search_js(search_query: str) -> str:
    return f"""
(() => {{
  const needle = {json.dumps(str(search_query), ensure_ascii=False)};
  const visible = (el) => {{
    if (!el) return false;
    const st = window.getComputedStyle(el);
    const r = el.getBoundingClientRect();
    return st.display !== 'none' && st.visibility !== 'hidden' && r.width > 0 && r.height > 0;
  }};
  const norm = (v) => String(v || '').trim().toLowerCase();
  const ctx = (el) => {{
    const parts = [el.placeholder, el.name, el.id, el.getAttribute && el.getAttribute('aria-label')];
    let p = el.parentElement;
    for (let i = 0; i < 3 && p; i++) {{
      parts.push(p.innerText || '');
      p = p.parentElement;
    }}
    return norm(parts.join(' '));
  }};
	  const inputs = Array.from(document.querySelectorAll('input')).filter(visible)
	    .filter(el => el.type !== 'hidden' && !el.disabled && !el.readOnly);
	  const bodyText = String(document.body?.innerText || '');
	  const href = String(location.href || '');
	  if (!href.toLowerCase().includes('#/add-product/v2')) {{
	    return JSON.stringify({{ok:false, reason:'not_add_product_page', url: href}});
	  }}
	  const articleInput = inputs.find(el => /(^|\\s|[\\n\\r])артикул($|\\s|[\\n\\r])/i.test(ctx(el)));
	  const nameInput = inputs.find(el => /наименование товара|название товара/i.test(ctx(el)));
	  const searchContextPresent = /присоединиться|существующей карточк|ссылк|url|поиск/i.test(bodyText);
  if (bodyText.toLowerCase().includes('информация о товаре') && articleInput && nameInput && !searchContextPresent) {{
    return JSON.stringify({{ok:false, reason:'unsafe_seller_identity_form_not_search'}});
  }}
  if (articleInput && /https?:\\/\\/kaspi\\.kz\\/shop\\/p\\//i.test(String(articleInput.value || ''))) {{
    return JSON.stringify({{ok:false, reason:'marketplace_url_in_article_field'}});
  }}
  const target = inputs.find(el => {{
    const text = ctx(el);
    if (/(^|\\s|[\\n\\r])артикул($|\\s|[\\n\\r])/i.test(text) && !/ссылк|url|link|поиск/i.test(text)) return false;
    return /ссылк|url|link|вариант|названию|поиск/i.test(text);
  }});
  if (!target) return JSON.stringify({{ok:false, reason:'search_input_not_found'}});
  if (target === articleInput) {{
    return JSON.stringify({{ok:false, reason:'refusing_article_field_as_search'}});
  }}
  target.focus();
  target.select && target.select();
  target.value = needle;
  ['input','change','keyup'].forEach(ev => target.dispatchEvent(new Event(ev, {{ bubbles:true }})));
  target.dispatchEvent(new KeyboardEvent('keydown', {{ key:'Enter', bubbles:true }}));
  const buttons = Array.from(document.querySelectorAll('button')).filter(visible);
  const searchBtn = buttons.find(btn => /поиск/i.test(norm(btn.innerText)));
  if (searchBtn) searchBtn.click();
  return JSON.stringify({{ok:true}});
}})();
""".strip()


def _step_probe_search_input_js() -> str:
    return r"""
(() => {
  const visible = (el) => {
    if (!el) return false;
    const st = window.getComputedStyle(el);
    const r = el.getBoundingClientRect();
    return st.display !== 'none' && st.visibility !== 'hidden' && r.width > 0 && r.height > 0;
  };
  const norm = (v) => String(v || '').trim().toLowerCase();
  const ctx = (el) => {
    const parts = [el.placeholder, el.name, el.id, el.getAttribute && el.getAttribute('aria-label')];
    let p = el.parentElement;
    for (let i = 0; i < 3 && p; i++) {
      parts.push(p.innerText || '');
      p = p.parentElement;
    }
    return norm(parts.join(' '));
  };
	  const inputs = Array.from(document.querySelectorAll('input')).filter(visible)
	    .filter(el => el.type !== 'hidden' && !el.disabled && !el.readOnly);
	  const bodyText = String(document.body?.innerText || '');
	  const href = String(location.href || '');
	  if (!href.toLowerCase().includes('#/add-product/v2')) {
	    return JSON.stringify({ok:false, ready:false, reason:'not_add_product_page', url: href});
	  }
	  const articleInput = inputs.find(el => /(^|\s|[\n\r])артикул($|\s|[\n\r])/i.test(ctx(el)));
	  const nameInput = inputs.find(el => /наименование товара|название товара/i.test(ctx(el)));
	  const searchContextPresent = /присоединиться|существующей карточк|ссылк|url|поиск/i.test(bodyText);
  if (bodyText.toLowerCase().includes('информация о товаре') && articleInput && nameInput && !searchContextPresent) {
    return JSON.stringify({ok:false, ready:false, reason:'unsafe_seller_identity_form_not_search'});
  }
  if (articleInput && /https?:\/\/kaspi\.kz\/shop\/p\//i.test(String(articleInput.value || ''))) {
    return JSON.stringify({ok:false, ready:false, reason:'marketplace_url_in_article_field'});
  }
  const target = inputs.find(el => {
    const text = ctx(el);
    if (/(^|\s|[\n\r])артикул($|\s|[\n\r])/i.test(text) && !/ссылк|url|link|поиск/i.test(text)) return false;
    return /ссылк|url|link|вариант|названию|поиск/i.test(text);
  });
  return JSON.stringify({ok:true, ready: !!target, reason: target ? '' : 'search_input_not_found'});
})();
""".strip()


def _step_verify_offer_identity_js(expected_heading: str, offer_code: str) -> str:
    return f"""
(() => {{
	  const heading = String({json.dumps(str(expected_heading), ensure_ascii=False)} || '').trim().toLowerCase();
	  const offerCode = String({json.dumps(str(offer_code), ensure_ascii=False)} || '').trim();
	  const href = String(location.href || '');
	  if (!href.toLowerCase().includes('#/add-product/v2')) {{
	    return JSON.stringify({{ok:false, reason:'not_add_product_page', url: href}});
	  }}
	  const bodyTextRaw = String(document.body?.innerText || '');
	  const txt = bodyTextRaw.toLowerCase();
	  if (!/информация о товаре|добавление товара|присоединиться|существующей карточк/i.test(bodyTextRaw)) {{
	    return JSON.stringify({{ok:false, reason:'add_product_context_missing'}});
	  }}
	  if (heading && !txt.includes(heading)) {{
	    return JSON.stringify({{ok:false, reason:'heading_mismatch'}});
	  }}
  if (offerCode) {{
    const rx = new RegExp('\\\\b' + offerCode + '\\\\b');
    if (!rx.test(String(document.body?.innerText || ''))) {{
      return JSON.stringify({{ok:false, reason:'offer_code_mismatch'}});
    }}
  }}
  return JSON.stringify({{ok:true}});
}})();
""".strip()


def _step_choose_card_js(expected_heading: str, offer_code: str = "") -> str:
    return f"""
(() => {{
	  const expected = String({json.dumps(str(expected_heading), ensure_ascii=False)} || '').trim().toLowerCase();
	  const offerCode = String({json.dumps(str(offer_code), ensure_ascii=False)} || '').trim();
	  const visible = (el) => {{
	    if (!el) return false;
    const st = window.getComputedStyle(el);
    const r = el.getBoundingClientRect();
	    return st.display !== 'none' && st.visibility !== 'hidden' && r.width > 0 && r.height > 0;
	  }};
	  const norm = (v) => String(v || '').trim().toLowerCase();
	  const href = String(location.href || '');
	  const bodyTextRaw = String(document.body?.innerText || '');
	  if (!href.toLowerCase().includes('#/add-product/v2')) {{
	    return JSON.stringify({{ok:false, reason:'not_add_product_page', url: href}});
	  }}
	  if (!/добавление товара|информация о товаре|присоединиться|существующей карточк|выбрать/i.test(bodyTextRaw)) {{
	    return JSON.stringify({{ok:false, reason:'add_product_context_missing'}});
	  }}
	  const collectContext = (el) => {{
    const roots = [
      '.product-list__item',
      '.product',
      '[class*="product-list__item"]',
      '[class*="product__"]',
      '[class*="product"]',
      'article',
      'li',
      'tr',
      '[class*="card"]',
      '[class*="item"]',
      '[class*="row"]',
    ];
    for (const selector of roots) {{
      const root = el.closest && el.closest(selector);
      if (root) {{
        const txt = norm(root.innerText);
        if (txt && txt.length <= 300) return txt;
      }}
    }}
    let node = el;
    for (let i = 0; i < 6 && node; i++) {{
      const txt = norm(node.innerText);
      if (txt && txt.length <= 300) return txt;
      node = node.parentElement;
    }}
    return norm(el.innerText);
  }};
  const headingScore = (txt) => {{
    if (!expected) return 0;
    if (txt.includes(expected)) return 100;
    return expected.split(/\\s+/).filter(token => token.length >= 3 && txt.includes(token)).length;
  }};
  const offerCodeScore = (txt) => {{
    if (!offerCode) return 0;
    const rx = new RegExp('\\\\b' + offerCode + '\\\\b');
    return rx.test(txt) ? 1000 : 0;
  }};
  const buttons = Array.from(document.querySelectorAll('button')).filter(visible);
  const chooseButtons = buttons.filter(btn => /выбрать/i.test(norm(btn.innerText)));
  let best = null;
  for (const btn of chooseButtons) {{
    const txt = collectContext(btn);
    const score = headingScore(txt) + offerCodeScore(txt);
    if (!best || score > best.score) {{
      best = {{ btn, score, txt }};
    }}
  }}
  if (best && best.score > 0 && (!offerCode || offerCodeScore(best.txt) > 0)) {{
    best.btn.click();
    return JSON.stringify({{
      ok: true,
      mode: 'scored_match',
      heading_score: headingScore(best.txt),
      offer_code_score: offerCodeScore(best.txt),
    }});
  }}
  if (best && offerCode && chooseButtons.length === 1 && headingScore(best.txt) > 0) {{
    best.btn.click();
    return JSON.stringify({{
      ok: true,
      mode: 'single_heading_match_after_exact_code_search',
      heading_score: headingScore(best.txt),
      offer_code_score: offerCodeScore(best.txt),
    }});
  }}
  if (chooseButtons.length > 0 && offerCode) {{
    return JSON.stringify({{ok:false, reason:'exact_offer_code_card_not_found'}});
  }}
  if (chooseButtons.length > 0) return JSON.stringify({{ok:false, reason:'ambiguous_choose_without_offer_code'}});
	  const bodyText = norm(bodyTextRaw);
  const infoHeadingPresent = bodyText.includes('информация о товаре');
  const continueButtonPresent = buttons.some(btn => /продолжить/i.test(norm(btn.innerText)));
  const sizeOptionsPresent = Array.from(document.querySelectorAll('.matrix__values, .matrix__values *')).some(
    el => visible(el) && norm(el.innerText)
  );
  const offerCodeMatch = offerCode ? new RegExp('\\\\b' + offerCode + '\\\\b').test(bodyTextRaw) : false;
  if (
    infoHeadingPresent
    && continueButtonPresent
    && sizeOptionsPresent
    && (offerCode ? offerCodeMatch : (headingScore(bodyText) > 0 || !expected))
  ) {{
    return JSON.stringify({{
      ok: true,
      mode: 'already_on_product_info',
      heading_score: headingScore(bodyText),
      offer_code_score: offerCodeMatch ? 1000 : 0,
    }});
  }}
  return JSON.stringify({{ok:false, reason:'choose_button_not_found'}});
}})();
""".strip()


def _step_select_size_js(size_rus: Any) -> str:
    return f"""
(() => {{
  const sizeNeedle = String({json.dumps(str(size_rus), ensure_ascii=False)});
  const visible = (el) => {{
    if (!el) return false;
    const st = window.getComputedStyle(el);
    const r = el.getBoundingClientRect();
    return st.display !== 'none' && st.visibility !== 'hidden' && r.width > 0 && r.height > 0;
	  }};
	  const norm = (v) => String(v || '').trim().toLowerCase();
	  const href = String(location.href || '');
	  if (!href.toLowerCase().includes('#/add-product/v2')) {{
	    return JSON.stringify({{ok:false, reason:'not_add_product_page', url: href}});
	  }}
	  const containsSizeToken = (text) => {{
    const hay = norm(text);
    const needle = norm(sizeNeedle);
    if (!hay || !needle) return false;
    if (hay === needle) return true;
    return hay.includes(needle);
  }};
  const querySizeOptions = () => Array.from(document.querySelectorAll('.matrix__values.icon-box, .matrix__values')).filter(visible);
  const sizeOptions = querySizeOptions();
  const availableSizes = sizeOptions.map(el => String(el.innerText || '').trim()).filter(Boolean);
  const getSelected = () => querySizeOptions().find(
    el => /active/i.test(String(el.className || '')) && containsSizeToken(el.innerText)
  );
  const activateSize = (el) => {{
    if (!el) return;
    const targets = [el, el.firstElementChild].filter(Boolean);
    el.scrollIntoView && el.scrollIntoView({{ block: 'center', inline: 'center' }});
    targets.forEach(target => {{
      const rect = target.getBoundingClientRect ? target.getBoundingClientRect() : null;
      const init = {{
        bubbles: true,
        cancelable: true,
        view: window,
        clientX: rect ? rect.left + rect.width / 2 : 0,
        clientY: rect ? rect.top + rect.height / 2 : 0,
      }};
      ['pointerdown', 'mousedown', 'mouseup', 'click'].forEach(ev => {{
        try {{
          target.dispatchEvent(new MouseEvent(ev, init));
        }} catch (e) {{}}
      }});
      try {{
        target.click && target.click();
      }} catch (e) {{}}
    }});
  }};
  const matching = sizeOptions.filter(el => containsSizeToken(el.innerText));
  const selectedBefore = getSelected();
  if (selectedBefore) {{
    return JSON.stringify({{
      ok: true,
      already_selected: true,
      selected_before: String(selectedBefore.innerText || '').trim(),
      available_sizes: availableSizes,
    }});
  }}
  if (matching.length === 0) {{
  const deepNode = Array.from(document.querySelectorAll('.matrix__values, .matrix__values *')).find(
      el => containsSizeToken(el.innerText)
    );
    const fallback = deepNode && (deepNode.closest('.matrix__values') || deepNode);
    if (!fallback) {{
      return JSON.stringify({{
        ok: false,
        reason: 'size_option_not_found',
        available_sizes: availableSizes,
      }});
    }}
    activateSize(fallback);
    return JSON.stringify({{
      ok: true,
      clicked_target: String(fallback.innerText || '').trim(),
      selected_before: '',
      available_sizes: availableSizes,
    }});
  }}
  activateSize(matching[0]);
  return JSON.stringify({{
    ok: true,
    clicked_target: String(matching[0].innerText || '').trim(),
    selected_before: '',
    available_sizes: availableSizes,
  }});
}})();
""".strip()


def _step_verify_selected_size_js(size_rus: Any) -> str:
    return f"""
(() => {{
  const sizeNeedle = String({json.dumps(str(size_rus), ensure_ascii=False)});
  const visible = (el) => {{
    if (!el) return false;
    const st = window.getComputedStyle(el);
    const r = el.getBoundingClientRect();
    return st.display !== 'none' && st.visibility !== 'hidden' && r.width > 0 && r.height > 0;
	  }};
	  const norm = (v) => String(v || '').trim().toLowerCase();
	  const href = String(location.href || '');
	  if (!href.toLowerCase().includes('#/add-product/v2')) {{
	    return JSON.stringify({{ok:false, reason:'not_add_product_page', url: href}});
	  }}
	  const containsSizeToken = (text) => {{
    const hay = norm(text);
    const needle = norm(sizeNeedle);
    if (!hay || !needle) return false;
    if (hay === needle) return true;
    return hay.includes(needle);
  }};
  const sizeOptions = Array.from(document.querySelectorAll('.matrix__values.icon-box, .matrix__values')).filter(visible);
  const availableSizes = sizeOptions.map(el => String(el.innerText || '').trim()).filter(Boolean);
  const selected = sizeOptions.find(el => /active/i.test(String(el.className || '')));
  const selectedSize = selected ? String(selected.innerText || '').trim() : '';
  if (selected && containsSizeToken(selected.innerText)) {{
    return JSON.stringify({{
      ok: true,
      selected_size: selectedSize,
      available_sizes: availableSizes,
    }});
  }}
  return JSON.stringify({{
    ok: false,
    reason: 'size_not_selected',
    selected_size: selectedSize,
    available_sizes: availableSizes,
  }});
}})();
""".strip()


def _step_fill_identity_fields_js(article: str, name: str) -> str:
    return f"""
(() => {{
  const articleVal = String({json.dumps(str(article), ensure_ascii=False)});
  const nameVal = String({json.dumps(str(name), ensure_ascii=False)});
  const visible = (el) => {{
    if (!el) return false;
    const st = window.getComputedStyle(el);
    const r = el.getBoundingClientRect();
    return st.display !== 'none' && st.visibility !== 'hidden' && r.width > 0 && r.height > 0;
	  }};
	  const norm = (v) => String(v || '').trim().toLowerCase();
	  const bodyText = String(document.body?.innerText || '');
	  const href = String(location.href || '');
	  if (!href.toLowerCase().includes('#/add-product/v2')) {{
	    return JSON.stringify({{ok:false, reason:'not_add_product_page', url: href}});
	  }}
	  if (!/информация о товаре/i.test(bodyText)) {{
	    return JSON.stringify({{ok:false, reason:'product_info_context_missing'}});
	  }}
	  if (/ваш товар успешно добавлен/i.test(bodyText)) {{
	    return JSON.stringify({{ok:false, reason:'stale_success_modal_present'}});
	  }}

  const contextText = (el) => {{
    const parts = [el.placeholder, el.name, el.id, el.getAttribute && el.getAttribute('aria-label')];
    let p = el.parentElement;
    for (let i = 0; i < 3 && p; i++) {{
      parts.push(p.innerText || '');
      p = p.parentElement;
    }}
    return norm(parts.join(' '));
  }};

  const inputs = Array.from(document.querySelectorAll('input')).filter(visible)
    .filter(el => el.type !== 'hidden' && !el.disabled && !el.readOnly);
  const articleInput = inputs.find(el => /артикул|article/i.test(contextText(el))) || inputs[0];
  const nameInput = inputs.find(el => /наименование|товар|название/i.test(contextText(el))) || inputs[1];
  if (!articleInput || !nameInput) return JSON.stringify({{ok:false, reason:'article_or_name_input_not_found'}});

  const setVal = (el, val) => {{
    el.focus();
    el.select && el.select();
    el.value = val;
    ['input','change','keyup'].forEach(ev => el.dispatchEvent(new Event(ev, {{ bubbles:true }})));
  }};
  setVal(articleInput, articleVal);
  setVal(nameInput, nameVal);

  if (String(articleInput.value || '').trim() !== articleVal.trim()) {{
    return JSON.stringify({{ok:false, reason:'article_value_not_applied'}});
  }}
  if (String(nameInput.value || '').trim() !== nameVal.trim()) {{
    return JSON.stringify({{ok:false, reason:'name_value_not_applied'}});
  }}

  const buttons = Array.from(document.querySelectorAll('button')).filter(visible);
  const nextBtn = buttons.find(btn => /продолжить/i.test(norm(btn.innerText)));
  if (!nextBtn) return JSON.stringify({{ok:false, reason:'continue_button_not_found'}});
  nextBtn.click();
  return JSON.stringify({{ok:true}});
}})();
""".strip()


def _step_fill_info_js(size_rus: Any, article: str, name: str) -> str:
    return _step_fill_identity_fields_js(article, name)


def _step_fill_price_js(price: Any, pp1: Any, pp2: Any, barcode: str) -> str:
    return f"""
(() => {{
  const priceVal = String({json.dumps(str(price), ensure_ascii=False)});
  const pp1Val = String({json.dumps(str(pp1), ensure_ascii=False)});
  const pp2Val = String({json.dumps(str(pp2), ensure_ascii=False)});
  // Owner rule for this uploader path: Kaspi barcode must stay empty.
  const barcodeVal = '';
  const visible = (el) => {{
    if (!el) return false;
    const st = window.getComputedStyle(el);
    const r = el.getBoundingClientRect();
    return st.display !== 'none' && st.visibility !== 'hidden' && r.width > 0 && r.height > 0;
  }};
  const norm = (v) => String(v || '').trim().toLowerCase();
  const bodyText = String(document.body?.innerText || '');
  if (/ваш товар успешно добавлен/i.test(bodyText)) {{
    return JSON.stringify({{ok:false, reason:'stale_success_modal_present'}});
  }}
  const dialogs = Array.from(document.querySelectorAll('[role=\"dialog\"], .modal, .modal-content, .v-dialog, .dialog'))
    .filter(visible);
  const priceWarning = dialogs.find(el => /вы не указали цену/i.test(norm(el.innerText)));
  if (priceWarning) {{
    const warnButtons = Array.from(priceWarning.querySelectorAll('button')).filter(visible);
    const fixBtn = warnButtons.find(btn => /указать цену/i.test(norm(btn.innerText)));
    if (fixBtn) {{
      fixBtn.click();
      return JSON.stringify({{ok:false, reason:'price_required_popup_closed'}});
    }}
    return JSON.stringify({{ok:false, reason:'price_required_popup_no_button'}});
  }}

  let container = dialogs.find(el => /цена и остатки/i.test(norm(el.innerText)));
  if (!container) {{
    const candidates = Array.from(document.querySelectorAll('section,div')).filter(visible);
    container = candidates.find(el => /цена и остатки/i.test(norm(el.innerText)) && el.querySelectorAll('input').length >= 3) || null;
  }}
  if (!container) return JSON.stringify({{ok:false, reason:'price_modal_not_found'}});

  const inputs = Array.from(container.querySelectorAll('input')).filter(visible)
    .filter(el => el.type !== 'hidden' && !el.disabled && !el.readOnly);
  if (inputs.length < 2) return JSON.stringify({{ok:false, reason:'not_enough_inputs_for_price_stock'}});

  const ctx = (el) => {{
    const parts = [el.placeholder, el.name, el.id, el.getAttribute && el.getAttribute('aria-label')];
    let p = el.parentElement;
    for (let i = 0; i < 3 && p; i++) {{
      parts.push(p.innerText || '');
      p = p.parentElement;
    }}
    return norm(parts.join(' '));
  }};

  const pick = (regexes) => inputs.find(el => regexes.some(rx => rx.test(ctx(el))));
  const stockInputsByTestId = inputs.filter(
    el => norm(el.getAttribute && el.getAttribute('data-testid')) === 'edit-stock-input'
  );
  const stockInputs = stockInputsByTestId.length ? stockInputsByTestId : [];
  let priceInput = inputs.find(el => !stockInputs.includes(el)) || null;
  let barcodeInput = null;
  if (priceInput) {{
    const barcodeCandidate = inputs.find(
      el => el !== priceInput && !stockInputs.includes(el) && /штрих|barcode|bar code/i.test(ctx(el))
    );
    if (barcodeCandidate) barcodeInput = barcodeCandidate;
  }} else {{
    priceInput = pick([/цена/, /price/, /₸/i]) || inputs[0];
    const used = new Set([priceInput]);
    const barcodeCandidate = pick([/штрих|barcode|bar code/i]);
    if (barcodeCandidate && !used.has(barcodeCandidate)) {{
      barcodeInput = barcodeCandidate;
      used.add(barcodeInput);
    }}
    stockInputs.push(...inputs.filter(el => !used.has(el)));
  }}
  if (!priceInput) priceInput = inputs[0] || null;
  if (stockInputs.length < 1) return JSON.stringify({{ok:false, reason:'stock_inputs_not_found'}});

  const setVal = (el, val) => {{
    el.focus();
    el.select && el.select();
    const setter = Object.getOwnPropertyDescriptor(window.HTMLInputElement.prototype, 'value').set;
    setter.call(el, String(val || ''));
    el.dispatchEvent(new InputEvent('input', {{ bubbles:true, inputType:'insertText', data: String(val || '') }}));
    el.dispatchEvent(new Event('change', {{ bubbles:true }}));
    el.dispatchEvent(new KeyboardEvent('keyup', {{ key:'0', bubbles:true }}));
    el.blur && el.blur();
  }};

  setVal(priceInput, priceVal);
  setVal(stockInputs[0], pp1Val);
  if (stockInputs[1]) setVal(stockInputs[1], pp2Val);
  if (barcodeInput) setVal(barcodeInput, barcodeVal);

  const digits = (v) => String(v || '').replace(/[^0-9]/g, '');
  if (digits(priceInput.value) !== digits(priceVal)) return JSON.stringify({{ok:false, reason:'price_value_not_applied'}});
  if (digits(stockInputs[0].value) !== digits(pp1Val)) return JSON.stringify({{ok:false, reason:'pp1_value_not_applied'}});
  if (stockInputs[1] && digits(stockInputs[1].value) !== digits(pp2Val)) return JSON.stringify({{ok:false, reason:'pp2_value_not_applied'}});
  if (barcodeInput && digits(barcodeInput.value)) return JSON.stringify({{ok:false, reason:'barcode_not_empty'}});
  return JSON.stringify({{ok:true, reason:'price_values_applied'}});
}})();
""".strip()


def _step_click_save_price_js() -> str:
    return r"""
(() => {
  const visible = (el) => {
    if (!el) return false;
    const st = window.getComputedStyle(el);
    const r = el.getBoundingClientRect();
    return st.display !== 'none' && st.visibility !== 'hidden' && r.width > 0 && r.height > 0;
  };
  const norm = (v) => String(v || '').trim().toLowerCase();
  const bodyText = String(document.body?.innerText || '');
  if (/ваш товар успешно добавлен/i.test(bodyText)) {
    return JSON.stringify({ok:false, reason:'stale_success_modal_present'});
  }
  const dialogs = Array.from(document.querySelectorAll('[role="dialog"], .modal, .modal-content, .v-dialog, .dialog'))
    .filter(visible);
  const priceWarning = dialogs.find(el => /вы не указали цену/i.test(norm(el.innerText)));
  if (priceWarning) {
    const warnButtons = Array.from(priceWarning.querySelectorAll('button')).filter(visible);
    const fixBtn = warnButtons.find(btn => /указать цену/i.test(norm(btn.innerText)));
    if (fixBtn) {
      fixBtn.click();
      return JSON.stringify({ok:false, reason:'price_required_popup_closed'});
    }
    return JSON.stringify({ok:false, reason:'price_required_popup_no_button'});
  }

  const scope = dialogs.find(el => /цена и остатки/i.test(norm(el.innerText))) || document;
  const buttons = Array.from(scope.querySelectorAll('button')).filter(visible);
  const globalButtons = Array.from(document.querySelectorAll('button')).filter(visible);
  const saveBtn = buttons.find(btn => /сохранить изменения/i.test(norm(btn.innerText)))
    || globalButtons.find(btn => /сохранить изменения/i.test(norm(btn.innerText)));
  if (!saveBtn) return JSON.stringify({ok:false, reason:'save_button_not_found'});
  saveBtn.click();
  return JSON.stringify({ok:true});
})();
""".strip()


def _step_detect_price_warning_js() -> str:
    return r"""
(() => {
  const visible = (el) => {
    if (!el) return false;
    const st = window.getComputedStyle(el);
    const r = el.getBoundingClientRect();
    return st.display !== 'none' && st.visibility !== 'hidden' && r.width > 0 && r.height > 0;
  };
  const norm = (v) => String(v || '').trim().toLowerCase();
  const dialogs = Array.from(document.querySelectorAll('[role="dialog"], .modal, .modal-content, .v-dialog, .dialog'))
    .filter(visible);
  const warning = dialogs.find(el => /вы не указали цену/i.test(norm(el.innerText)));
  return JSON.stringify({ok:true, warning_present: !!warning});
})();
""".strip()


def _step_verify_success_js() -> str:
    return r"""
(() => {
  const txt = String(document.body?.innerText || '');
  if (/ваш товар успешно добавлен/i.test(txt)) return JSON.stringify({ok:true});
  if (/товар.*уже существует|уже есть в вашем магазине|уже добавлен/i.test(txt)) {
    return JSON.stringify({ok:true, mode:'already_exists'});
  }
  return JSON.stringify({ok:false, reason:'success_message_not_found'});
})();
""".strip()


def _json_result(raw: str) -> dict[str, Any]:
    try:
        obj = json.loads(raw or "{}")
        if isinstance(obj, dict):
            return obj
    except Exception:
        pass
    return {"ok": False, "reason": f"invalid_json:{raw}"}


def _write_run_log_header(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as fh:
        wr = csv.writer(fh)
        wr.writerow(
            [
                "timestamp",
                "store_code",
                "merchant_id",
                "row_number",
                "merchant_sku_article",
                "variant_url",
                "status",
                "error",
            ]
        )


def _append_run_log(path: Path, row: list[Any]) -> None:
    with path.open("a", newline="", encoding="utf-8") as fh:
        wr = csv.writer(fh)
        wr.writerow(row)


def _execute_row(
    window_index: int,
    row: dict[str, Any],
    *,
    step_delay_seconds: float,
    step_screenshot_dir: Path | None = None,
) -> tuple[bool, str]:
    store_code = normalize_store_code(row.get("store_code"))
    expected_mid = str(_to_int(row.get("merchant_id")))
    if store_code not in ALLOWED_UPLOAD_STORE_CODES:
        return False, f"forbidden_store:{store_code or '<empty>'}"
    if store_code in FORBIDDEN_UPLOAD_STORE_CODES or expected_mid in FORBIDDEN_UPLOAD_MERCHANT_IDS:
        return False, f"forbidden_upload_surface store={store_code or '<empty>'} merchant_id={expected_mid}"
    expected_mid_for_store = ALLOWED_UPLOAD_MERCHANT_IDS_BY_STORE.get(store_code)
    if expected_mid_for_store and expected_mid != expected_mid_for_store:
        return False, f"merchant_id_store_mismatch store={store_code} merchant_id={expected_mid} expected={expected_mid_for_store}"
    normalized_barcode = _normalize_barcode(row.get("barcode"))
    normalized_size_rus = _normalize_size_rus(row.get("size_rus"))
    offer_code = _product_code_for_search(row)

    def has_size_like_option(values: Any) -> bool:
        for value in values or []:
            text = str(value or "").strip().upper()
            if re.search(r"\d", text):
                return True
            if re.search(r"\b(?:XS|S|M|L|XL|2XL|3XL|4XL|5XL|6XL)\b", text):
                return True
        return False

    def resolve_window() -> int:
        nonlocal window_index
        try:
            candidates = list_store_window_candidates()
            window_index = _choose_window_index(int(window_index), expected_mid, candidates)
        except Exception:
            window_index = int(window_index)
        return int(window_index)

    def snap(stage: str) -> None:
        if not step_screenshot_dir:
            return
        row_no = int(row.get("row_number") or 0)
        article = re.sub(r"[^a-zA-Z0-9._-]+", "_", str(row.get("merchant_sku_article") or "")).strip("_")
        if not article:
            article = "no_article"
        stem = f"row{row_no:04d}_{stage}_{article[:80]}"
        win = resolve_window()
        try:
            state = _json_result(_execute_tab_js(win, _step_page_state_js()))
            state["stage"] = stage
            state["row_number"] = row_no
            state["merchant_sku_article"] = str(row.get("merchant_sku_article") or "")
            state["expected_price_kzt"] = row.get("price_kzt")
            state["expected_stock_pp1"] = row.get("stock_pp1")
            state["expected_stock_pp2"] = row.get("stock_pp2")
            state["expected_product_code"] = offer_code
            (step_screenshot_dir / f"{stem}.json").write_text(
                json.dumps(state, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
        except Exception:
            pass
        _capture_window_screenshot(win, step_screenshot_dir / f"{stem}.png")

    def close_stale_success(stage: str) -> dict[str, Any]:
        try:
            res = _json_result(_execute_tab_js(resolve_window(), _step_close_stale_success_modal_js()))
        except Exception as exc:
            return {"ok": False, "reason": str(exc)}
        if res.get("closed"):
            time.sleep(0.5)
            snap(f"{stage}_closed_stale_success")
        return res

    entry_url = str(row.get("ui_entry_url") or DEFAULT_ENTRY_URL)
    direct_code_mode = ("link-catalog" in entry_url.lower()) and ("code=" in entry_url.lower())

    selected_window = resolve_window()
    pre_mid = ""
    try:
        pre_mid = str(_execute_tab_js(selected_window, _merchant_id_js()) or "").strip()
    except Exception:
        pre_mid = ""
    if pre_mid in FORBIDDEN_UPLOAD_MERCHANT_IDS:
        return False, f"forbidden_merchant_window_before_navigation window={pre_mid}"
    if pre_mid and pre_mid != expected_mid:
        return False, f"merchant_id_mismatch_before_navigation window={pre_mid} expected={expected_mid}"

    _set_tab_url(selected_window, entry_url)
    time.sleep(max(step_delay_seconds, 1.2))
    close_stale_success("01_open")
    snap("01_open")

    mid_raw = _execute_tab_js(resolve_window(), _merchant_id_js())
    actual_mid = str(mid_raw or "").strip()
    if actual_mid in FORBIDDEN_UPLOAD_MERCHANT_IDS:
        snap("err_forbidden_merchant")
        return False, f"forbidden_merchant_window_after_navigation window={actual_mid}"
    if actual_mid and actual_mid != expected_mid:
        snap("err_mid_mismatch")
        return False, f"merchant_id_mismatch window={actual_mid} expected={expected_mid}"

    if not direct_code_mode:
        close_stale_success("pre_search")
        unsafe = _json_result(_execute_tab_js(resolve_window(), _step_assert_not_unsafe_identity_form_js()))
        if not unsafe.get("ok"):
            snap("err_unsafe_identity_form")
            return False, f"unsafe_page_state:{unsafe.get('reason')}"
        search_ready = False
        search_reason = "search_input_not_found"
        for _ in range(12):
            probe = _json_result(_execute_tab_js(resolve_window(), _step_probe_search_input_js()))
            search_reason = str(probe.get("reason") or search_reason)
            if not probe.get("ok") and search_reason in {
                "unsafe_seller_identity_form_not_search",
                "marketplace_url_in_article_field",
            }:
                snap("err_unsafe_identity_form")
                return False, f"unsafe_page_state:{search_reason}"
            if bool(probe.get("ready")):
                search_ready = True
                break
            time.sleep(0.8)
        if not search_ready:
            snap("err_search")
            return False, f"search_failed:{search_reason}"

        res = _json_result(_execute_tab_js(resolve_window(), _step_search_js(offer_code)))
        if not res.get("ok"):
            reason = str(res.get("reason") or "")
            snap("err_unsafe_identity_form" if reason in {
                "unsafe_seller_identity_form_not_search",
                "marketplace_url_in_article_field",
                "refusing_article_field_as_search",
            } else "err_search")
            return False, f"search_failed:{reason}"
        time.sleep(step_delay_seconds)
        close_stale_success("02_search")
        snap("02_search")

        res = _json_result(
            _execute_tab_js(
                resolve_window(),
                _step_choose_card_js(
                    str(row.get("expected_kaspi_heading") or ""),
                    offer_code,
                ),
            )
        )
        if not res.get("ok"):
            snap("err_choose")
            return False, f"choose_card_failed:{res.get('reason')}"
        time.sleep(step_delay_seconds + 0.6)
        close_stale_success("03_choose")
        snap("03_choose")

        identity_check = _json_result(
            _execute_tab_js(
                resolve_window(),
                _step_verify_offer_identity_js(
                    str(row.get("expected_kaspi_heading") or ""),
                    "",
                ),
            )
        )
        if not identity_check.get("ok"):
            snap("err_identity")
            return False, f"identity_failed:{identity_check.get('reason')}"
    else:
        snap("02_direct_code_entry")

    size_ok = False
    size_reason = ""
    for i in range(4):
        res_select = _json_result(
            _execute_tab_js(
                resolve_window(),
                _step_select_size_js(normalized_size_rus),
            )
        )
        if not res_select.get("ok"):
            size_reason = str(res_select.get("reason") or "size_select_failed")
            if (
                direct_code_mode
                and size_reason == "size_option_not_found"
                and not has_size_like_option(res_select.get("available_sizes"))
            ):
                identity_check = _json_result(
                    _execute_tab_js(
                        resolve_window(),
                        _step_verify_offer_identity_js(
                            str(row.get("expected_kaspi_heading") or ""),
                            offer_code,
                        ),
                    )
                )
                if identity_check.get("ok"):
                    size_ok = True
                    size_reason = "direct_code_no_size_options_identity_verified"
                    break
            snap("err_size_select")
            break
        time.sleep(step_delay_seconds + 0.4)
        res_verify = _json_result(
            _execute_tab_js(
                resolve_window(),
                _step_verify_selected_size_js(normalized_size_rus),
            )
        )
        if res_verify.get("ok"):
            size_ok = True
            break
        size_reason = str(res_verify.get("reason") or "size_not_selected")
        snap(f"04_size_retry{i+1}")
        time.sleep(step_delay_seconds + 0.4)
    if not size_ok:
        snap("err_fill_info")
        return False, f"fill_info_failed:{size_reason}"
    close_stale_success("04_size_selected")
    snap("04_size_selected")

    identity_check = _json_result(
        _execute_tab_js(
            resolve_window(),
            _step_verify_offer_identity_js(
                str(row.get("expected_kaspi_heading") or ""),
                offer_code,
            ),
        )
    )
    if not identity_check.get("ok"):
        snap("err_identity")
        return False, f"identity_failed:{identity_check.get('reason')}"

    res: dict[str, Any] = {"ok": False, "reason": "identity_fill_not_attempted"}
    for _ in range(2):
        close_stale_success("pre_fill_info")
        res = _json_result(
            _execute_tab_js(
                resolve_window(),
                _step_fill_identity_fields_js(
                    str(row.get("merchant_sku_article") or ""),
                    str(row.get("merchant_offer_name") or ""),
                ),
            )
        )
        if res.get("ok") or res.get("reason") != "stale_success_modal_present":
            break
        close_stale_success("fill_info")
    if not res.get("ok"):
        snap("err_fill_info")
        return False, f"fill_info_failed:{res.get('reason')}"
    time.sleep(step_delay_seconds + 0.6)
    snap("05_fill_info")

    price_ok = False
    price_reason = ""
    for i in range(8):
        close_stale_success(f"pre_fill_price_retry{i+1}")
        res_fill = _json_result(
            _execute_tab_js(
                resolve_window(),
                _step_fill_price_js(
                    row.get("price_kzt"),
                    row.get("stock_pp1"),
                    row.get("stock_pp2"),
                    normalized_barcode,
                ),
            )
        )
        if not res_fill.get("ok"):
            price_reason = str(res_fill.get("reason") or "unknown")
            snap(f"06_fill_price_retry{i+1}")
            if price_reason == "stale_success_modal_present":
                close_stale_success(f"fill_price_retry{i+1}")
                time.sleep(step_delay_seconds + 0.6)
                continue
            if price_reason in {
                "price_required_popup_closed",
                "price_modal_not_found",
                "not_enough_inputs_for_price_stock",
                "pp_inputs_not_found",
            }:
                time.sleep(step_delay_seconds + 0.8)
                continue
            break

        # Let UI model settle before attempting save; this prevents race with validation popups.
        time.sleep(step_delay_seconds + 0.9)
        res_save = _json_result(_execute_tab_js(resolve_window(), _step_click_save_price_js()))
        if res_save.get("ok"):
            # Server-side validation is async; verify popup does not reappear after save.
            time.sleep(step_delay_seconds + 0.6)
            post_save = _json_result(_execute_tab_js(resolve_window(), _step_detect_price_warning_js()))
            if bool(post_save.get("warning_present")):
                price_reason = "price_required_popup_after_save"
                snap(f"06_fill_price_retry{i+1}")
                # Attempt to close popup and retry fill+save.
                _execute_tab_js(resolve_window(), _step_click_save_price_js())
                time.sleep(step_delay_seconds + 0.8)
                continue
            price_ok = True
            break
        price_reason = str(res_save.get("reason") or "unknown")
        snap(f"05_fill_price_retry{i+1}")
        if price_reason == "stale_success_modal_present":
            close_stale_success(f"save_price_retry{i+1}")
            time.sleep(step_delay_seconds + 0.6)
            continue
        if price_reason in {
            "price_required_popup_closed",
            "price_required_popup_no_button",
            "save_button_not_found",
        }:
            time.sleep(step_delay_seconds + 0.8)
            continue
        break
    if not price_ok:
        snap("err_fill_price")
        return False, f"fill_price_failed:{price_reason}"
    time.sleep(step_delay_seconds + 0.6)
    snap("06_fill_price")

    verify_ok = False
    verify_reason = "success_message_not_found"
    for i in range(5):
        res = _json_result(_execute_tab_js(resolve_window(), _step_verify_success_js()))
        if res.get("ok"):
            verify_ok = True
            break
        verify_reason = str(res.get("reason") or "success_message_not_found")
        snap(f"07_verify_retry{i+1}")
        time.sleep(step_delay_seconds + 0.5)
    if not verify_ok:
        snap("err_verify")
        return False, f"verify_failed:{verify_reason}"
    snap("08_success")

    return True, ""


def run_kaspi_offer_ui_upload(
    *,
    workbook_path: str | Path,
    store_codes: str | list[str],
    dry_run: bool,
    confirm: bool,
    output_root: str | Path = "runs/kaspi_offer_ui_upload",
    window_map: str | None = None,
    step_delay_seconds: float = 1.8,
    retries: int = 1,
    step_screenshots: bool = False,
    require_black_coverage: bool = True,
) -> dict[str, Any]:
    stores = parse_store_codes(store_codes)
    if not stores:
        stores = ["UNIVERSAL", "STOREB"]

    rows = load_offer_upload_rows(workbook_path, store_codes=stores)
    validation = validate_upload_rows(
        rows,
        required_store_codes=stores,
        require_black_coverage=require_black_coverage,
    )

    run_id = datetime.now(ASTANA_TZ).strftime("%Y%m%d_%H%M%S")
    run_dir = Path(output_root) / run_id
    summary: dict[str, Any] = {
        "run_id": run_id,
        "status": "dry_run" if dry_run or not confirm else "success",
        "workbook_path": str(workbook_path),
        "store_codes": stores,
        "rows_total": len(rows),
        "rows_by_store": validation.get("rows_by_store", {}),
        "validation": validation,
        "run_dir": str(run_dir),
        "run_log_csv": str(run_dir / "run_log.csv"),
        "success": 0,
        "failed": 0,
        "errors": [],
        "window_map": {},
        "step_screenshots": bool(step_screenshots),
        "require_black_coverage": bool(require_black_coverage),
    }

    if not validation.get("ok"):
        summary["status"] = "preflight_failed"
        summary["errors"] = list(validation.get("errors", []))
        return summary

    if dry_run or not confirm:
        return summary

    run_dir.mkdir(parents=True, exist_ok=True)
    run_log_path = run_dir / "run_log.csv"
    _write_run_log_header(run_log_path)

    if window_map:
        mapped = parse_window_map(window_map)
    else:
        mapped = _auto_map_windows(rows)
    summary["window_map"] = mapped

    for store in stores:
        if store not in mapped:
            summary["status"] = "preflight_failed"
            summary["errors"].append(f"window mapping missing for store {store}")
    if summary["status"] == "preflight_failed":
        return summary

    max_attempts = max(1, int(retries) + 1)
    for row in rows:
        store = normalize_store_code(row.get("store_code"))
        candidates = list_store_window_candidates()
        expected_mid = str(_to_int(row.get("merchant_id")))
        mapped_win = int(mapped[store])
        win = _choose_window_index(mapped_win, expected_mid, candidates)
        mapped[store] = int(win)
        summary["window_map"] = dict(mapped)
        ok = False
        err = ""
        row_shots_dir = (run_dir / "screenshots") if step_screenshots else None
        for _ in range(max_attempts):
            try:
                ok, err = _execute_row(
                    win,
                    row,
                    step_delay_seconds=step_delay_seconds,
                    step_screenshot_dir=row_shots_dir,
                )
            except Exception as exc:
                ok, err = False, f"exception:{exc}"
            if ok:
                break
            time.sleep(0.6)

        if ok:
            summary["success"] += 1
            status = "ok"
        else:
            summary["failed"] += 1
            status = "failed"
            summary["errors"].append(
                f"row {row.get('row_number')} store {store}: {err}"
            )

        _append_run_log(
            run_log_path,
            [
                datetime.now(ASTANA_TZ).isoformat(),
                store,
                row.get("merchant_id"),
                row.get("row_number"),
                row.get("merchant_sku_article"),
                row.get("variant_url"),
                status,
                err,
            ],
        )

    if summary["failed"] > 0:
        summary["status"] = "partial" if summary["success"] > 0 else "failed"

    (run_dir / "summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return summary


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Kaspi offer UI upload runner")
    parser.add_argument("--workbook", help="Workbook path with 'offer upload' sheet")
    parser.add_argument("--stores", default="UNIVERSAL,STOREB", help="Comma-separated store codes")
    parser.add_argument("--dry-run", action="store_true", help="Only validate and build summary")
    parser.add_argument("--confirm", action="store_true", help="Execute real UI writes")
    parser.add_argument("--output-root", default="runs/kaspi_offer_ui_upload", help="Run artifacts root")
    parser.add_argument("--window-map", help="Manual mapping STORE:WINDOW_INDEX,...")
    parser.add_argument("--step-delay-seconds", type=float, default=1.8, help="Delay between UI steps")
    parser.add_argument("--retries", type=int, default=1, help="Retries per row")
    parser.add_argument("--step-screenshots", action="store_true", help="Capture window screenshots for each row step")
    parser.add_argument(
        "--allow-partial-color-batch",
        action="store_true",
        help="Allow scoped white/grey-only retry batches without requiring a black row in each target store",
    )
    parser.add_argument("--list-windows", action="store_true", help="Print Chrome window candidates with detected merchant IDs")
    parser.add_argument("--json", action="store_true", help="Print JSON summary")
    args = parser.parse_args(argv)

    if args.list_windows:
        windows = list_store_window_candidates()
        if args.json:
            print(json.dumps({"windows": windows}, ensure_ascii=False, indent=2))
        else:
            for w in windows:
                print(
                    f"window={w['window_index']} merchant_id={w['merchant_id'] or '-'} "
                    f"title={w['title'][:80]} url={w['url'][:120]}"
                )
        return 0

    if not args.workbook:
        parser.error("--workbook is required unless --list-windows is used")

    summary = run_kaspi_offer_ui_upload(
        workbook_path=args.workbook,
        store_codes=args.stores,
        dry_run=args.dry_run,
        confirm=args.confirm,
        output_root=args.output_root,
        window_map=args.window_map,
        step_delay_seconds=args.step_delay_seconds,
        retries=args.retries,
        step_screenshots=args.step_screenshots,
        require_black_coverage=not args.allow_partial_color_batch,
    )
    if args.json:
        print(json.dumps(summary, ensure_ascii=False, indent=2))
    else:
        print(
            f"status={summary.get('status')} rows={summary.get('rows_total')} "
            f"success={summary.get('success')} failed={summary.get('failed')}"
        )
    return 0 if summary.get("status") in {"success", "partial", "dry_run"} else 4


if __name__ == "__main__":
    raise SystemExit(main())
