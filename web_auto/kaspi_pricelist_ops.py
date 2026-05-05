from __future__ import annotations

import json
import shutil
import tempfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

from .kaspi_merchant_common import normalize_store_name


TEMPLATE_COLUMNS = ["SKU", "model", "brand", "price", "PP1", "PP2", "PP3", "PP4", "PP5", "preorder"]
WAREHOUSE_COLUMNS = ["PP1", "PP2", "PP3", "PP4", "PP5"]
DEFAULT_OFFERS_BOOK = Path("exports/offers_book.xlsx")
DEFAULT_TRUTH_XLSX = Path("exports/repricer_unified_truth.xlsx")
DEFAULT_REPAIR_STOCK_UNITS = 500


def parse_int(value: object) -> int:
    text = str(value or "").strip().replace(" ", "").replace(",", ".")
    if not text:
        return 0
    try:
        return int(float(text))
    except Exception:
        return 0


def is_no_like(value: object) -> bool:
    return str(value or "").strip().lower() in {"", "no"}


def normalize_url(value: object) -> str:
    return str(value or "").strip().lower().rstrip("/")


def coerce_whole_number(value: object) -> int | None:
    text = str(value or "").strip().replace(" ", "").replace(",", ".")
    if not text:
        return None
    try:
        return int(float(text))
    except Exception:
        return None


def merchant_sku_candidates(merchant_sku: object) -> list[str]:
    raw = str(merchant_sku or "").strip()
    if not raw:
        return []
    candidates: list[str] = []
    parts = [p for p in raw.replace("\t", " ").split() if p]
    for token in [raw, raw.replace("\t", " "), *parts]:
        token = str(token or "").strip()
        if token and token not in candidates:
            candidates.append(token)
    return candidates


def row_active_warehouse_pattern(row: dict[str, Any] | pd.Series) -> tuple[str, ...]:
    return tuple(col for col in WAREHOUSE_COLUMNS if not is_no_like(row.get(col, "")))


def row_has_positive_stock(row: dict[str, Any] | pd.Series) -> bool:
    return any(parse_int(row.get(col, "")) > 0 for col in WAREHOUSE_COLUMNS)


def detect_active_stock_units(row: dict[str, Any] | pd.Series) -> dict[str, Any]:
    pattern = row_active_warehouse_pattern(row)
    if not pattern:
        return {
            "pattern": tuple(),
            "status": "no_active_warehouse_pattern",
            "issue": True,
            "issue_columns": [],
        }
    issue_columns: list[str] = []
    issue_kinds: set[str] = set()
    for col in pattern:
        raw_value = row.get(col, "")
        numeric = coerce_whole_number(raw_value)
        if numeric is None:
            issue_columns.append(col)
            issue_kinds.add("non_numeric")
            continue
        if numeric <= 0:
            issue_columns.append(col)
            issue_kinds.add("non_positive")
    if not issue_columns:
        return {
            "pattern": pattern,
            "status": "ok",
            "issue": False,
            "issue_columns": [],
        }
    if issue_kinds == {"non_numeric"}:
        status = "non_numeric_active_units"
    elif issue_kinds == {"non_positive"}:
        status = "non_positive_active_units"
    else:
        status = "mixed_invalid_active_units"
    return {
        "pattern": pattern,
        "status": status,
        "issue": True,
        "issue_columns": issue_columns,
    }


def _load_template_sheet(path: Path, sheet_name: str) -> pd.DataFrame:
    df = pd.read_excel(path, sheet_name=sheet_name, dtype=str).fillna("")
    for col in TEMPLATE_COLUMNS:
        if col not in df.columns:
            raise ValueError(f"missing required column `{col}` in {path}:{sheet_name}")
    return df


def _load_l2(path: Path) -> pd.DataFrame:
    return pd.read_excel(path, sheet_name="Лист2", dtype=str).fillna("")


def _size_lookup_from_offers_book(offers_book_path: Path) -> dict[tuple[str, str], str]:
    size_df = pd.read_excel(offers_book_path, sheet_name="size_level", dtype=str).fillna("")
    by_store_url: dict[tuple[str, str], set[str]] = {}
    by_url: dict[str, set[str]] = {}
    for _, row in size_df.iterrows():
        store_name = normalize_store_name(row.get("store_name", ""))
        url = str(row.get("resolved_url", "") or "").strip().lower().rstrip("/")
        size = str(row.get("final_attached_size", "") or "").strip().upper()
        if not url or not size:
            continue
        if store_name:
            by_store_url.setdefault((store_name, url), set()).add(size)
        by_url.setdefault(url, set()).add(size)
    resolved: dict[tuple[str, str], str] = {}
    for key, sizes in by_store_url.items():
        if len(sizes) == 1:
            resolved[key] = next(iter(sizes))
    for url, sizes in by_url.items():
        if len(sizes) == 1:
            resolved[("", url)] = next(iter(sizes))
    return resolved


def _build_offer_index(offers_book_path: Path) -> dict[tuple[str, str], list[dict[str, str]]]:
    offers_df = pd.read_excel(offers_book_path, sheet_name="offers_merged", dtype=str).fillna("")
    size_lookup = _size_lookup_from_offers_book(offers_book_path)
    index: dict[tuple[str, str], list[dict[str, str]]] = {}
    for _, row in offers_df.iterrows():
        store_name = normalize_store_name(row.get("store_name", ""))
        sku_raw = str(row.get("sku_raw", "") or "").strip()
        if not store_name or not sku_raw:
            continue
        resolved_url = str(row.get("resolved_url", "") or "").strip().lower().rstrip("/")
        final_size = size_lookup.get((store_name, resolved_url), "") or size_lookup.get(("", resolved_url), "")
        payload = {
            "sku_raw": sku_raw,
            "resolved_url": str(row.get("resolved_url", "") or "").strip(),
            "resolved_sku_key": str(row.get("resolved_sku_key", "") or "").strip(),
            "resolved_sku_id": str(row.get("resolved_sku_id", "") or "").strip(),
            "resolved_kaspi_offer_name": str(row.get("resolved_kaspi_offer_name", "") or "").strip(),
            "final_attached_size": final_size,
            "mapping_status": str(row.get("mapping_status", "") or "").strip(),
        }
        keys = set()
        for source in (row.get("sku_raw", ""), row.get("sku_id_ksp_input", ""), row.get("resolved_sku_id", "")):
            for candidate in merchant_sku_candidates(source):
                keys.add(candidate)
        for candidate in keys:
            index.setdefault((store_name, candidate), []).append(payload)
    return index


def _collapse_offer_matches(matches: list[dict[str, str]]) -> tuple[dict[str, str], str]:
    if not matches:
        return {}, "missing"
    unique = {
        (
            item.get("resolved_url", ""),
            item.get("resolved_sku_key", ""),
            item.get("resolved_sku_id", ""),
            item.get("resolved_kaspi_offer_name", ""),
            item.get("final_attached_size", ""),
        ): item
        for item in matches
    }
    if len(unique) == 1:
        return next(iter(unique.values())), "unique"
    return {}, "ambiguous"


def _build_truth_index(truth_xlsx_path: Path) -> dict[tuple[str, str], list[dict[str, str]]]:
    truth_df = pd.read_excel(truth_xlsx_path, dtype=str).fillna("")
    truth_df["store_name"] = truth_df["store_name"].map(normalize_store_name)
    truth_df["price_num"] = truth_df["price"].map(parse_int)
    truth_df["min_price_num"] = truth_df["min_price"].map(parse_int)
    truth_df["max_price_num"] = truth_df["max_price"].map(parse_int)
    truth_df = truth_df.sort_values(["store_name", "merchant_sku", "fetched_at"], ascending=[True, True, False])
    index: dict[tuple[str, str], list[dict[str, str]]] = {}
    for _, row in truth_df.iterrows():
        store_name = normalize_store_name(row.get("store_name", ""))
        merchant_sku = str(row.get("merchant_sku", "") or "").strip()
        if not store_name or not merchant_sku:
            continue
        payload = {
            "price": str(row.get("price_num", "") or ""),
            "min_price": str(row.get("min_price_num", "") or ""),
            "max_price": str(row.get("max_price_num", "") or ""),
            "link": str(row.get("link", "") or "").strip(),
            "merchant_title": str(row.get("merchant_title", "") or "").strip(),
            "dumping": str(row.get("dumping", "") or "").strip(),
            "fetched_at": str(row.get("fetched_at", "") or "").strip(),
        }
        for candidate in merchant_sku_candidates(merchant_sku):
            index.setdefault((store_name, candidate), []).append(payload)
    return index


def _collapse_truth_matches(matches: list[dict[str, str]]) -> dict[str, str]:
    if not matches:
        return {}
    return matches[0]


@dataclass(frozen=True)
class PricelistSnapshot:
    snapshot_df: pd.DataFrame
    active_l2: pd.DataFrame
    archive_l2: pd.DataFrame
    source_active_path: Path
    source_archive_path: Path
    store_name: str


def build_store_snapshot(
    *,
    active_path: Path,
    archive_path: Path,
    store_name: str,
    offers_book_path: Path = DEFAULT_OFFERS_BOOK,
    truth_xlsx_path: Path = DEFAULT_TRUTH_XLSX,
) -> PricelistSnapshot:
    active_df = _load_template_sheet(active_path, "Лист1").copy()
    archive_df = _load_template_sheet(archive_path, "Лист1").copy()
    active_l2 = _load_l2(active_path)
    archive_l2 = _load_l2(archive_path)
    active_df["source_state"] = "ACTIVE"
    archive_df["source_state"] = "ARCHIVE"
    merged = pd.concat([active_df, archive_df], ignore_index=True)
    store_norm = normalize_store_name(store_name)
    offer_index = _build_offer_index(offers_book_path)
    truth_index = _build_truth_index(truth_xlsx_path)

    records: list[dict[str, Any]] = []
    for _, row in merged.iterrows():
        record = {col: str(row.get(col, "") or "") for col in merged.columns}
        record["store_name"] = store_norm
        sku = record["SKU"]
        offer_matches: list[dict[str, str]] = []
        truth_matches: list[dict[str, str]] = []
        for candidate in merchant_sku_candidates(sku):
            offer_matches.extend(offer_index.get((store_norm, candidate), []))
            truth_matches.extend(truth_index.get((store_norm, candidate), []))
        offer_match, match_confidence = _collapse_offer_matches(offer_matches)
        truth_match = _collapse_truth_matches(truth_matches)
        record["resolved_url"] = offer_match.get("resolved_url", "")
        record["resolved_sku_key"] = offer_match.get("resolved_sku_key", "")
        record["resolved_sku_id"] = offer_match.get("resolved_sku_id", "")
        record["resolved_kaspi_offer_name"] = offer_match.get("resolved_kaspi_offer_name", "")
        record["final_attached_size"] = offer_match.get("final_attached_size", "")
        record["mapping_status"] = offer_match.get("mapping_status", "")
        record["match_confidence"] = match_confidence
        record["live_price"] = truth_match.get("price", "")
        record["live_min_price"] = truth_match.get("min_price", "")
        record["live_max_price"] = truth_match.get("max_price", "")
        record["live_link"] = truth_match.get("link", "")
        record["live_title"] = truth_match.get("merchant_title", "")
        record["live_dumping"] = truth_match.get("dumping", "")
        active_stock_units = detect_active_stock_units(record) if record["source_state"] == "ACTIVE" else {
            "pattern": row_active_warehouse_pattern(record),
            "status": "out_of_scope",
            "issue": False,
            "issue_columns": [],
        }
        record["current_active_pattern"] = ",".join(active_stock_units["pattern"])
        record["active_stock_unit_status"] = active_stock_units["status"]
        record["active_stock_unit_issue"] = "yes" if active_stock_units["issue"] else "no"
        record["active_stock_unit_issue_columns"] = ",".join(active_stock_units["issue_columns"])
        record["stock_positive_current"] = "yes" if row_has_positive_stock(record) else "no"
        price_num = parse_int(record.get("price", ""))
        live_price_num = parse_int(record.get("live_price", ""))
        record["current_price_num"] = str(price_num)
        record["live_price_num"] = str(live_price_num)
        record["suspicious_off"] = "yes" if record["source_state"] == "ARCHIVE" and record["stock_positive_current"] == "yes" else "no"
        record["suspicious_on"] = "yes" if record["source_state"] == "ACTIVE" and record["stock_positive_current"] == "no" else "no"
        record["price_mismatch"] = "yes" if live_price_num and live_price_num != price_num else "no"
        record["intended_state"] = record["source_state"]
        record["intended_price"] = record["price"]
        record["intent_action"] = ""
        records.append(record)

    return PricelistSnapshot(
        snapshot_df=pd.DataFrame(records),
        active_l2=active_l2,
        archive_l2=archive_l2,
        source_active_path=active_path,
        source_archive_path=archive_path,
        store_name=store_norm,
    )


def apply_intent(
    snapshot_df: pd.DataFrame,
    *,
    intent: str,
    sku: str = "",
    sku_key: str = "",
    group_url: str = "",
    target_price: str = "",
) -> pd.DataFrame:
    out = snapshot_df.copy()
    group_url_norm = str(group_url or "").strip().lower().rstrip("/")
    selection = pd.Series([True] * len(out), index=out.index)
    if sku:
        selection &= out["SKU"].astype(str).str.strip().eq(str(sku).strip())
    if sku_key:
        selection &= out["resolved_sku_key"].astype(str).str.strip().eq(str(sku_key).strip())
    if group_url_norm:
        selection &= out["resolved_url"].astype(str).str.lower().str.rstrip("/").eq(group_url_norm)
    if intent == "repair-active-stock-units":
        selection &= out["source_state"].astype(str).str.strip().eq("ACTIVE")
    out["selected_for_intent"] = selection.map(lambda v: "yes" if v else "no")
    out["intent_name"] = intent
    out["intent_review_reason"] = ""
    out["intended_active_pattern"] = out.get("current_active_pattern", pd.Series([""] * len(out), index=out.index))
    out["intended_stock_value"] = ""

    for idx, row in out.iterrows():
        if out.at[idx, "selected_for_intent"] != "yes":
            continue
        if intent in {"inspect-group-status", "inspect"}:
            out.at[idx, "intent_action"] = "inspect"
            continue
        if intent == "repair-active-stock-units":
            active_stock_units = detect_active_stock_units(row)
            out.at[idx, "current_active_pattern"] = ",".join(active_stock_units["pattern"])
            out.at[idx, "active_stock_unit_status"] = active_stock_units["status"]
            out.at[idx, "active_stock_unit_issue"] = "yes" if active_stock_units["issue"] else "no"
            out.at[idx, "active_stock_unit_issue_columns"] = ",".join(active_stock_units["issue_columns"])
            out.at[idx, "intended_state"] = "ACTIVE"
            out.at[idx, "intended_active_pattern"] = ",".join(active_stock_units["pattern"])
            if active_stock_units["status"] == "no_active_warehouse_pattern":
                out.at[idx, "intent_action"] = "blocked_no_active_warehouse_pattern"
                out.at[idx, "intent_review_reason"] = "no_active_warehouse_pattern"
                continue
            if active_stock_units["issue"]:
                out.at[idx, "intent_action"] = "repair_active_stock_units"
                out.at[idx, "intended_stock_value"] = str(DEFAULT_REPAIR_STOCK_UNITS)
                continue
            out.at[idx, "intent_action"] = "active_stock_units_already_numeric"
            continue
        if intent in {"turn-on-in-stock", "repair-suspicious-off"}:
            if row["source_state"] == "ARCHIVE" and row["stock_positive_current"] == "yes":
                out.at[idx, "intended_state"] = "ACTIVE"
                out.at[idx, "intent_action"] = "archive_to_active"
                if parse_int(row.get("live_price", "")) > 0:
                    out.at[idx, "intended_price"] = str(parse_int(row.get("live_price", "")))
            continue
        if intent in {"turn-off-oos", "repair-suspicious-on"}:
            if row["source_state"] == "ACTIVE" and row["stock_positive_current"] == "no":
                out.at[idx, "intended_state"] = "ARCHIVE"
                out.at[idx, "intent_action"] = "active_to_archive"
            continue
        if intent == "set-upload-prices":
            resolved_target = str(target_price or "").strip()
            if not resolved_target and parse_int(row.get("live_price", "")) > 0:
                resolved_target = str(parse_int(row.get("live_price", "")))
            if resolved_target:
                out.at[idx, "intended_price"] = resolved_target
                out.at[idx, "intent_action"] = "price_refresh"
            continue
        raise ValueError(f"unsupported intent: {intent}")
    return out


def _review_bucket(row: pd.Series) -> str:
    intent_name = str(row.get("intent_name", "") or "").strip()
    if str(row.get("intent_review_reason", "") or "").strip():
        return str(row.get("intent_review_reason", "") or "").strip()
    if intent_name == "repair-active-stock-units":
        return ""
    if row.get("match_confidence", "") == "ambiguous":
        return "ambiguous_offer_match"
    if row.get("match_confidence", "") == "missing":
        return "missing_offer_match"
    if row.get("selected_for_intent", "") == "yes" and not row.get("intent_action", ""):
        return "selected_but_not_actionable"
    if row.get("suspicious_off", "") == "yes":
        return "suspicious_off"
    if row.get("suspicious_on", "") == "yes":
        return "suspicious_on"
    return ""


def _normalize_output_df(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    for col in TEMPLATE_COLUMNS:
        if col not in out.columns:
            out[col] = ""
    return out[TEMPLATE_COLUMNS].fillna("")


def _apply_row_write_rules(row: pd.Series) -> dict[str, Any]:
    out: dict[str, Any] = {col: str(row.get(col, "") or "") for col in TEMPLATE_COLUMNS}
    if str(row.get("intended_price", "") or "").strip():
        out["price"] = str(row.get("intended_price", "") or "").strip()
    if str(row.get("intent_action", "") or "").strip() == "repair_active_stock_units":
        active_pattern = tuple(
            part.strip() for part in str(row.get("intended_active_pattern", "") or "").split(",") if part.strip()
        )
        stock_value = coerce_whole_number(row.get("intended_stock_value", ""))
        if active_pattern and stock_value is not None and stock_value > 0:
            for col in WAREHOUSE_COLUMNS:
                out[col] = int(stock_value) if col in active_pattern else "no"
    if str(row.get("intended_state", "") or "") == "ARCHIVE":
        for col in WAREHOUSE_COLUMNS:
            out[col] = "no"
    return out


def _write_workbook_atomic(path: Path, l1_df: pd.DataFrame, l2_df: pd.DataFrame) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp_name = tempfile.mkstemp(prefix=f"{path.stem}.", suffix=path.suffix, dir=str(path.parent))
    tmp_path = Path(tmp_name)
    try:
        with pd.ExcelWriter(tmp_path, engine="openpyxl") as writer:
            l1_df.to_excel(writer, sheet_name="Лист1", index=False)
            l2_df.to_excel(writer, sheet_name="Лист2", index=False)
        wb = load_workbook(tmp_path)
        ws = wb["Лист1"]
        header_map = {str(cell.value or ""): idx + 1 for idx, cell in enumerate(ws[1])}
        numeric_columns = {"price", "PP1", "PP2", "PP3", "PP4", "PP5", "preorder"}
        for row_idx in range(2, ws.max_row + 1):
            for column_name in numeric_columns:
                column_index = header_map.get(column_name)
                if not column_index:
                    continue
                cell = ws.cell(row=row_idx, column=column_index)
                value = cell.value
                if str(value or "").strip().lower() == "no":
                    cell.value = "no"
                    continue
                coerced = coerce_whole_number(value)
                if coerced is None:
                    continue
                cell.value = coerced
                cell.number_format = "0"
        wb.save(tmp_path)
        shutil.move(str(tmp_path), str(path))
    finally:
        try:
            Path(tmp_name).unlink(missing_ok=True)
        except Exception:
            pass


def emit_outputs(
    snapshot: PricelistSnapshot,
    *,
    mutated_df: pd.DataFrame,
    output_dir: Path,
    prefix: str | None = None,
) -> dict[str, Any]:
    output_dir.mkdir(parents=True, exist_ok=True)
    prefix_value = prefix or snapshot.store_name.replace("-", "").lower()
    active_rows = mutated_df[mutated_df["intended_state"] == "ACTIVE"].copy()
    archive_rows = mutated_df[mutated_df["intended_state"] == "ARCHIVE"].copy()
    review_bucket_series = mutated_df.apply(_review_bucket, axis=1)
    review_rows = mutated_df[review_bucket_series != ""].copy()
    review_rows["review_bucket"] = review_bucket_series.loc[review_rows.index]

    final_active = _normalize_output_df(pd.DataFrame([_apply_row_write_rules(row) for _, row in active_rows.iterrows()]))
    final_archive = _normalize_output_df(pd.DataFrame([_apply_row_write_rules(row) for _, row in archive_rows.iterrows()]))
    if final_active.empty:
        final_active = pd.DataFrame(columns=TEMPLATE_COLUMNS)
    if final_archive.empty:
        final_archive = pd.DataFrame(columns=TEMPLATE_COLUMNS)

    active_out = output_dir / f"{prefix_value}_ACTIVE.xlsx"
    archive_out = output_dir / f"{prefix_value}_ARCHIVE.xlsx"
    review_out = output_dir / f"{prefix_value}_review.xlsx"
    snapshot_out = output_dir / f"{prefix_value}_snapshot.xlsx"
    summary_out = output_dir / f"{prefix_value}_summary.json"

    _write_workbook_atomic(active_out, final_active, snapshot.active_l2)
    _write_workbook_atomic(archive_out, final_archive, snapshot.archive_l2)
    with pd.ExcelWriter(review_out, engine="openpyxl") as writer:
        pd.DataFrame(columns=TEMPLATE_COLUMNS).to_excel(writer, sheet_name="Лист1", index=False)
        review_rows.to_excel(writer, sheet_name="review_rows", index=False)
        snapshot.archive_l2.to_excel(writer, sheet_name="Лист2", index=False)
    with pd.ExcelWriter(snapshot_out, engine="openpyxl") as writer:
        mutated_df.to_excel(writer, sheet_name="snapshot", index=False)

    summary = {
        "status": "ready" if len(review_rows) == 0 else "blocked",
        "store_name": snapshot.store_name,
        "snapshot_rows": int(len(mutated_df)),
        "selected_rows": int((mutated_df["selected_for_intent"] == "yes").sum()) if "selected_for_intent" in mutated_df.columns else 0,
        "active_rows": int(len(final_active)),
        "archive_rows": int(len(final_archive)),
        "review_rows": int(len(review_rows)),
        "action_counts": mutated_df["intent_action"].value_counts().to_dict(),
        "repair_rows": int((mutated_df["intent_action"] == "repair_active_stock_units").sum()),
        "already_numeric_rows": int((mutated_df["intent_action"] == "active_stock_units_already_numeric").sum()),
        "blocked_rows": int((mutated_df["intent_action"] == "blocked_no_active_warehouse_pattern").sum()),
        "active_output": str(active_out),
        "archive_output": str(archive_out),
        "review_output": str(review_out),
        "snapshot_output": str(snapshot_out),
    }
    summary_out.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return summary


def verify_uploaded_state(
    *,
    before_df: pd.DataFrame,
    after_snapshot: PricelistSnapshot,
) -> dict[str, Any]:
    after_df = after_snapshot.snapshot_df.copy()
    after_index = {
        str(row.get("SKU", "")).strip(): row
        for _, row in after_df.iterrows()
        if str(row.get("SKU", "")).strip()
    }
    mismatches: list[dict[str, Any]] = []
    selected = before_df[before_df["selected_for_intent"] == "yes"].copy()
    for _, row in selected.iterrows():
        sku = str(row.get("SKU", "")).strip()
        if not sku:
            continue
        after_row = after_index.get(sku)
        if after_row is None:
            mismatches.append({"SKU": sku, "reason": "missing_after_upload"})
            continue
        expected_state = str(row.get("intended_state", "") or "")
        actual_state = str(after_row.get("source_state", "") or "")
        if expected_state and expected_state != actual_state:
            mismatches.append({"SKU": sku, "reason": "state_mismatch", "expected": expected_state, "actual": actual_state})
        expected_price = parse_int(row.get("intended_price", ""))
        actual_price = parse_int(after_row.get("price", ""))
        if expected_price and expected_price != actual_price:
            mismatches.append({"SKU": sku, "reason": "price_mismatch", "expected": expected_price, "actual": actual_price})
    return {
        "selected_rows": int(len(selected)),
        "mismatch_count": int(len(mismatches)),
        "mismatches": mismatches,
    }


def plan_append_missing_active_rows(
    *,
    target_snapshot: PricelistSnapshot,
    source_snapshot: PricelistSnapshot,
) -> dict[str, Any]:
    target_df = target_snapshot.snapshot_df.copy()
    source_df = source_snapshot.snapshot_df.copy()

    target_all = target_df.copy()
    target_all["url_norm"] = target_all["resolved_url"].map(normalize_url)
    target_all["sku_norm"] = target_all["SKU"].astype(str).str.strip()
    target_urls = {value for value in target_all["url_norm"] if value}
    target_skus = {value for value in target_all["sku_norm"] if value}

    source_active = source_df[source_df["source_state"] == "ACTIVE"].copy()
    source_active["url_norm"] = source_active["resolved_url"].map(normalize_url)
    source_active["sku_norm"] = source_active["SKU"].astype(str).str.strip()

    append_rows: list[dict[str, Any]] = []
    review_rows: list[dict[str, Any]] = []
    planned_url_keys: set[str] = set()
    planned_sku_keys: set[str] = set()

    for _, row in source_active.iterrows():
        if str(row.get("stock_positive_current", "") or "").strip().lower() != "yes":
            continue
        url_norm = str(row.get("url_norm", "") or "").strip()
        sku_norm = str(row.get("sku_norm", "") or "").strip()
        if url_norm and url_norm in target_urls:
            continue
        if sku_norm and sku_norm in target_skus:
            continue
        if url_norm and url_norm in planned_url_keys:
            continue
        if sku_norm and sku_norm in planned_sku_keys:
            continue
        if str(row.get("match_confidence", "") or "").strip() != "unique":
            review = row.to_dict()
            review["review_bucket"] = "non_unique_mapping"
            review_rows.append(review)
            continue
        if not url_norm:
            review = row.to_dict()
            review["review_bucket"] = "missing_resolved_url"
            review_rows.append(review)
            continue
        append_rows.append(row.to_dict())
        planned_url_keys.add(url_norm)
        planned_sku_keys.add(sku_norm)

    append_df = pd.DataFrame(append_rows)
    review_df = pd.DataFrame(review_rows)

    target_active = target_df[target_df["source_state"] == "ACTIVE"].copy()
    target_archive = target_df[target_df["source_state"] == "ARCHIVE"].copy()

    final_active_rows = target_active.copy()
    if not append_df.empty:
        final_active_rows = pd.concat([final_active_rows, append_df], ignore_index=True)

    summary = {
        "target_store": target_snapshot.store_name,
        "source_store": source_snapshot.store_name,
        "target_active_count": int(len(target_active)),
        "target_archive_count": int(len(target_archive)),
        "source_active_count": int(len(source_active)),
        "append_count": int(len(append_df)),
        "review_count": int(len(review_df)),
        "final_active_count": int(len(final_active_rows)),
        "final_archive_count": int(len(target_archive)),
    }
    return {
        "append_rows": append_df,
        "review_rows": review_df,
        "final_active_rows": final_active_rows,
        "final_archive_rows": target_archive,
        "summary": summary,
    }


def emit_append_missing_active_outputs(
    *,
    target_snapshot: PricelistSnapshot,
    source_snapshot: PricelistSnapshot,
    output_dir: Path,
    prefix: str | None = None,
) -> dict[str, Any]:
    plan = plan_append_missing_active_rows(target_snapshot=target_snapshot, source_snapshot=source_snapshot)
    output_dir.mkdir(parents=True, exist_ok=True)
    prefix_value = prefix or f"{target_snapshot.store_name.replace('-', '').lower()}_append"

    active_out = output_dir / f"{prefix_value}_ACTIVE.xlsx"
    archive_out = output_dir / f"{prefix_value}_ARCHIVE.xlsx"
    review_out = output_dir / f"{prefix_value}_review.xlsx"
    summary_out = output_dir / f"{prefix_value}_summary.json"

    final_active = _normalize_output_df(plan["final_active_rows"])
    final_archive = _normalize_output_df(plan["final_archive_rows"])
    if final_active.empty:
        final_active = pd.DataFrame(columns=TEMPLATE_COLUMNS)
    if final_archive.empty:
        final_archive = pd.DataFrame(columns=TEMPLATE_COLUMNS)

    _write_workbook_atomic(active_out, final_active, target_snapshot.active_l2)
    _write_workbook_atomic(archive_out, final_archive, target_snapshot.archive_l2)
    with pd.ExcelWriter(review_out, engine="openpyxl") as writer:
        _normalize_output_df(plan["append_rows"]).to_excel(writer, sheet_name="append_candidates", index=False)
        plan["review_rows"].to_excel(writer, sheet_name="review_rows", index=False)
        target_snapshot.archive_l2.to_excel(writer, sheet_name="Лист2", index=False)

    summary = dict(plan["summary"])
    summary.update(
        {
            "active_output": str(active_out),
            "archive_output": str(archive_out),
            "review_output": str(review_out),
        }
    )
    summary_out.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return summary
