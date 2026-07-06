from __future__ import annotations

import csv
import json
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

from .kaspi_forbidden_cards import OWNER_DECISION_LABEL, forbidden_saleable_row_details
from .kaspi_merchant_common import normalize_store_name
from .kaspi_price_floors import clamp_rows_to_price_floors, write_price_floor_clamp_report
from .kaspi_pricelist_ops import (
    TEMPLATE_COLUMNS,
    _load_l2,
    _load_template_sheet,
    _write_workbook_atomic,
)


SAFE_ACTIVE_CONFIRM_PHRASE = "FULL_ACTIVE_STATE"
UPDATE_COLUMNS = {"price", "PP1", "PP2", "PP3", "PP4", "PP5", "preorder"}
CONTROL_COLUMNS = {"SKU", "sku", "activate_from_archive", "note", "owner_note"}
RESTRICTED_CLASSIFICATIONS = {
    "direct_platform_restricted",
    "family_platform_restriction_risk",
    "platform_restricted",
}


class SafeActivePatchError(ValueError):
    pass


@dataclass(frozen=True)
class SafeActivePatchUpdate:
    sku: str
    values: dict[str, str]
    activate_from_archive: bool = False


def _truthy(value: object) -> bool:
    return str(value or "").strip().lower() in {"1", "true", "yes", "y", "да", "activate", "active"}


def _clean_cell(value: object) -> str:
    text = str(value or "").strip()
    if text.endswith(".0"):
        head = text[:-2]
        if head.isdigit():
            return head
    return text


def load_safe_patch_updates(path: Path) -> list[SafeActivePatchUpdate]:
    with path.open(newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.DictReader(fh))
    if not rows:
        raise SafeActivePatchError("updates CSV has no rows")

    header = set(rows[0].keys())
    unknown = sorted(col for col in header if col not in CONTROL_COLUMNS and col not in UPDATE_COLUMNS)
    if unknown:
        raise SafeActivePatchError(f"updates CSV has unsupported columns: {', '.join(unknown)}")

    updates: list[SafeActivePatchUpdate] = []
    seen: set[str] = set()
    for row_index, row in enumerate(rows, start=2):
        sku = _clean_cell(row.get("SKU") or row.get("sku"))
        if not sku:
            raise SafeActivePatchError(f"updates CSV row {row_index} has empty SKU")
        if sku in seen:
            raise SafeActivePatchError(f"updates CSV has duplicate SKU: {sku}")
        seen.add(sku)

        values = {
            col: _clean_cell(row.get(col))
            for col in UPDATE_COLUMNS
            if _clean_cell(row.get(col)) != ""
        }
        activate_from_archive = _truthy(row.get("activate_from_archive"))
        if not values and not activate_from_archive:
            raise SafeActivePatchError(f"updates CSV row {row_index} has no mutation fields")
        updates.append(SafeActivePatchUpdate(sku=sku, values=values, activate_from_archive=activate_from_archive))
    return updates


def load_restricted_skus_from_ledger(path: Path | None) -> dict[str, dict[str, str]]:
    if path is None:
        return {}
    with path.open(newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.DictReader(fh))
    restricted: dict[str, dict[str, str]] = {}
    for row_index, row in enumerate(rows, start=2):
        sku = _clean_cell(row.get("merchant_sku") or row.get("SKU") or row.get("sku"))
        classification = _clean_cell(row.get("classification")).lower()
        if not sku:
            continue
        if classification in RESTRICTED_CLASSIFICATIONS:
            restricted[sku] = {
                "classification": classification,
                "row_index": str(row_index),
                "evidence_path": _clean_cell(row.get("evidence_path") or row.get("source_artifact") or row.get("direct_restriction_evidence") or row.get("family_restriction_evidence")),
            }
    return restricted


def load_restriction_probe_approval(path: Path | None, *, store_name: str) -> dict[str, str]:
    if path is None:
        return {}
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise SafeActivePatchError("restriction probe approval must be a JSON object")
    if _clean_cell(data.get("approval_type")) != "restriction_probe":
        raise SafeActivePatchError("restriction probe approval approval_type must equal restriction_probe")
    if data.get("owner_approved") is not True:
        raise SafeActivePatchError("restriction probe approval owner_approved must be true")
    store_norm = normalize_store_name(store_name)
    approved_store = _clean_cell(data.get("store"))
    if approved_store and normalize_store_name(approved_store) != store_norm:
        raise SafeActivePatchError(f"restriction probe approval store {approved_store} does not match {store_norm}")
    allowed_skus = data.get("allowed_skus")
    if not isinstance(allowed_skus, list) or not allowed_skus:
        raise SafeActivePatchError("restriction probe approval must include non-empty allowed_skus list")
    approval_id = _clean_cell(data.get("approval_id") or data.get("created_at") or path.name)
    return {_clean_cell(sku): approval_id for sku in allowed_skus if _clean_cell(sku)}


def _load_pricelist_l1(path: Path) -> pd.DataFrame:
    df = _load_template_sheet(path, "Лист1").copy()
    df = df[df["SKU"].astype(str).str.strip() != ""].copy()
    for col in TEMPLATE_COLUMNS:
        df[col] = df[col].map(_clean_cell)
    return df.reset_index(drop=True)


def _unique_index(df: pd.DataFrame) -> tuple[dict[str, pd.Series], list[str]]:
    index: dict[str, pd.Series] = {}
    duplicates: list[str] = []
    for _, row in df.iterrows():
        sku = _clean_cell(row.get("SKU"))
        if not sku:
            continue
        if sku in index and sku not in duplicates:
            duplicates.append(sku)
        index[sku] = row
    return index, sorted(duplicates)


def _row_dict(row: pd.Series | dict[str, Any]) -> dict[str, str]:
    return {col: _clean_cell(row.get(col, "")) for col in TEMPLATE_COLUMNS}


def _changed_fields(before: dict[str, str], after: dict[str, str]) -> dict[str, dict[str, str]]:
    return {
        col: {"before": before.get(col, ""), "after": after.get(col, "")}
        for col in TEMPLATE_COLUMNS
        if before.get(col, "") != after.get(col, "")
    }


def _write_preview_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    fieldnames = [
        "SKU",
        "source_state_before",
        "source_state_after",
        "price",
        "PP1",
        "PP2",
        "PP3",
        "PP4",
        "PP5",
        "preorder",
        "changed_fields_json",
    ]
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            writer.writerow({key: row.get(key, "") for key in fieldnames})


def validate_excel_workbook(path: Path) -> dict[str, Any]:
    try:
        with zipfile.ZipFile(path) as zf:
            bad_member = zf.testzip()
        wb = load_workbook(path, read_only=True, data_only=False)
        sheets = list(wb.sheetnames)
        wb.close()
        return {"status": "ok", "bad_zip_member": bad_member or "", "sheets": sheets}
    except Exception as exc:
        return {"status": "failed", "error": str(exc)}


def build_safe_active_patch(
    *,
    active_path: Path,
    archive_path: Path,
    updates_path: Path,
    output_dir: Path,
    store_name: str,
    expected_active_before: int | None = None,
    expected_active_after: int | None = None,
    allow_activate_from_archive: bool = False,
    restriction_ledger_path: Path | None = None,
    restriction_probe_approval_path: Path | None = None,
    prefix: str | None = None,
) -> dict[str, Any]:
    output_dir.mkdir(parents=True, exist_ok=True)
    store_norm = normalize_store_name(store_name)
    prefix_value = prefix or store_norm.replace("-", "").lower()
    summary_path = output_dir / f"{prefix_value}_safe_active_patch_summary.json"

    active_df = _load_pricelist_l1(active_path)
    archive_df = _load_pricelist_l1(archive_path)
    active_l2 = _load_l2(active_path)
    updates = load_safe_patch_updates(updates_path)
    restricted_skus = load_restricted_skus_from_ledger(restriction_ledger_path)
    restriction_probe_approvals = load_restriction_probe_approval(
        restriction_probe_approval_path,
        store_name=store_name,
    )

    active_index, active_duplicates = _unique_index(active_df)
    archive_index, archive_duplicates = _unique_index(archive_df)
    active_skus = set(active_index)
    archive_skus = set(archive_index)
    active_archive_overlap = sorted(active_skus & archive_skus)

    errors: list[str] = []
    if active_duplicates:
        errors.append(f"duplicate ACTIVE SKU rows: {', '.join(active_duplicates)}")
    if archive_duplicates:
        errors.append(f"duplicate ARCHIVE SKU rows: {', '.join(archive_duplicates)}")
    if active_archive_overlap:
        errors.append(f"ACTIVE/ARCHIVE overlap SKU rows: {', '.join(active_archive_overlap[:20])}")
    if expected_active_before is not None and len(active_df) != expected_active_before:
        errors.append(f"expected ACTIVE before {expected_active_before}, got {len(active_df)}")

    final_rows = [_row_dict(row) for _, row in active_df.iterrows()]
    final_by_sku = {row["SKU"]: row for row in final_rows}
    preview_rows: list[dict[str, Any]] = []
    activated_from_archive: list[str] = []
    updated_active: list[str] = []
    unknown_skus: list[str] = []
    archive_activation_blocked: list[str] = []
    restricted_update_skus: list[str] = []
    restricted_update_details: list[dict[str, str]] = []
    restriction_probe_override_skus: list[str] = []

    for update in updates:
        restricted_detail = restricted_skus.get(update.sku)
        if restricted_detail:
            approval_id = restriction_probe_approvals.get(update.sku)
            if not approval_id:
                restricted_update_skus.append(update.sku)
                restricted_update_details.append({"SKU": update.sku, **restricted_detail})
                continue
            restriction_probe_override_skus.append(update.sku)

        if update.sku in final_by_sku:
            before = dict(final_by_sku[update.sku])
            after = dict(before)
            for col, value in update.values.items():
                after[col] = value
            final_by_sku[update.sku].update(after)
            updated_active.append(update.sku)
            preview_rows.append(
                {
                    "SKU": update.sku,
                    "source_state_before": "ACTIVE",
                    "source_state_after": "ACTIVE",
                    **{col: after.get(col, "") for col in TEMPLATE_COLUMNS if col != "SKU"},
                    "changed_fields_json": json.dumps(_changed_fields(before, after), ensure_ascii=False, sort_keys=True),
                }
            )
            continue

        if update.sku in archive_index:
            if not (allow_activate_from_archive or update.activate_from_archive):
                archive_activation_blocked.append(update.sku)
                continue
            before = _row_dict(archive_index[update.sku])
            after = dict(before)
            for col, value in update.values.items():
                after[col] = value
            final_rows.append(after)
            final_by_sku[update.sku] = after
            activated_from_archive.append(update.sku)
            preview_rows.append(
                {
                    "SKU": update.sku,
                    "source_state_before": "ARCHIVE",
                    "source_state_after": "ACTIVE",
                    **{col: after.get(col, "") for col in TEMPLATE_COLUMNS if col != "SKU"},
                    "changed_fields_json": json.dumps(_changed_fields(before, after), ensure_ascii=False, sort_keys=True),
                }
            )
            continue

        unknown_skus.append(update.sku)

    if archive_activation_blocked:
        errors.append(
            "updates target ARCHIVE rows without activation approval: "
            + ", ".join(sorted(archive_activation_blocked))
        )
    if restricted_update_skus:
        errors.append(
            "updates target platform-restricted/risk SKU rows: "
            + ", ".join(sorted(restricted_update_skus))
        )
    if unknown_skus:
        errors.append(f"updates target SKU rows missing from ACTIVE and ARCHIVE: {', '.join(sorted(unknown_skus))}")

    final_active_skus = {row["SKU"] for row in final_rows if row.get("SKU")}
    missing_original_active = sorted(active_skus - final_active_skus)
    if missing_original_active:
        errors.append(f"output would remove baseline ACTIVE SKU rows: {', '.join(missing_original_active[:20])}")
    if expected_active_after is not None and len(final_rows) != expected_active_after:
        errors.append(f"expected ACTIVE after {expected_active_after}, got {len(final_rows)}")

    active_floor_result = clamp_rows_to_price_floors(final_rows, store_name=store_norm)
    final_rows = active_floor_result["rows"]
    restore_floor_result = clamp_rows_to_price_floors(
        [_row_dict(row) for _, row in active_df.iterrows()],
        store_name=store_norm,
    )
    price_floor_remaining = active_floor_result["remaining_violations"] + restore_floor_result["remaining_violations"]
    if price_floor_remaining:
        errors.append(
            "price floor guard failed after clamp: "
            + ", ".join(row["row_label"] for row in price_floor_remaining[:20])
        )
    forbidden_saleable_rows = forbidden_saleable_row_details(final_rows, store_name=store_norm)
    if forbidden_saleable_rows:
        owner_labels = sorted(
            {
                row.get("owner_decision_label") or OWNER_DECISION_LABEL
                for row in forbidden_saleable_rows
            }
        )
        errors.append(
            f"{'; '.join(owner_labels)}: output would set forbidden Kaspi offer cards saleable: "
            + ", ".join(row["row_label"] for row in forbidden_saleable_rows[:20])
        )

    active_output = output_dir / f"{prefix_value}_FULL_ACTIVE_UPLOAD.xlsx"
    restore_output = output_dir / f"{prefix_value}_RESTORE_BASELINE_ACTIVE.xlsx"
    preview_output = output_dir / f"{prefix_value}_safe_active_patch_preview.csv"
    price_floor_report_output = output_dir / f"{prefix_value}_price_floor_clamps.csv"
    status = "ready" if not errors else "blocked"

    price_floor_clamps = [
        {"surface": "active_upload", **row}
        for row in active_floor_result["clamps"]
    ] + [
        {"surface": "restore_baseline_active", **row}
        for row in restore_floor_result["clamps"]
    ]
    write_price_floor_clamp_report(price_floor_report_output, price_floor_clamps)

    restore_df = pd.DataFrame(restore_floor_result["rows"], columns=TEMPLATE_COLUMNS)
    _write_workbook_atomic(restore_output, restore_df, active_l2)
    workbook_validation: dict[str, Any] = {"restore": validate_excel_workbook(restore_output)}
    if status == "ready":
        final_df = pd.DataFrame(final_rows, columns=TEMPLATE_COLUMNS)
        _write_workbook_atomic(active_output, final_df, active_l2)
        workbook_validation["active_upload"] = validate_excel_workbook(active_output)
        if workbook_validation["active_upload"].get("status") != "ok":
            status = "blocked"
            errors.append("active upload workbook failed Excel/ZIP validation")
    _write_preview_csv(preview_output, preview_rows)

    summary = {
        "status": status,
        "store_name": store_norm,
        "source_active_path": str(active_path),
        "source_archive_path": str(archive_path),
        "updates_path": str(updates_path),
        "restriction_ledger_path": str(restriction_ledger_path) if restriction_ledger_path else "",
        "restriction_probe_approval_path": str(restriction_probe_approval_path) if restriction_probe_approval_path else "",
        "safe_patch_contract": "FULL_ACTIVE_STATE_REPLACEMENT",
        "required_confirm_phrase_for_apply": SAFE_ACTIVE_CONFIRM_PHRASE,
        "archive_upload_allowed": False,
        "active_before_count": int(len(active_df)),
        "archive_before_count": int(len(archive_df)),
        "update_count": int(len(updates)),
        "updated_active_count": int(len(updated_active)),
        "activated_from_archive_count": int(len(activated_from_archive)),
        "active_after_count": int(len(final_rows)),
        "expected_active_before": expected_active_before,
        "expected_active_after": expected_active_after,
        "missing_original_active_count": int(len(missing_original_active)),
        "missing_original_active_skus": missing_original_active,
        "updated_active_skus": sorted(updated_active),
        "activated_from_archive_skus": sorted(activated_from_archive),
        "unknown_update_skus": sorted(unknown_skus),
        "restricted_update_skus": sorted(restricted_update_skus),
        "restricted_update_details": sorted(restricted_update_details, key=lambda row: row["SKU"]),
        "restriction_probe_override_skus": sorted(restriction_probe_override_skus),
        "price_floor_clamp_rows_count": int(len(price_floor_clamps)),
        "price_floor_clamp_rows": price_floor_clamps,
        "price_floor_remaining_violations_count": int(len(price_floor_remaining)),
        "price_floor_remaining_violations": price_floor_remaining,
        "price_floor_clamp_report": str(price_floor_report_output),
        "forbidden_saleable_rows_count": int(len(forbidden_saleable_rows)),
        "forbidden_saleable_rows": forbidden_saleable_rows,
        "errors": errors,
        "active_output": str(active_output) if status == "ready" else "",
        "restore_output": str(restore_output),
        "preview_output": str(preview_output),
        "summary_output": str(summary_path),
        "workbook_validation": workbook_validation,
    }
    summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return summary


def verify_safe_active_upload(
    *,
    intended_active_path: Path,
    redownloaded_active_path: Path,
    updates_path: Path,
) -> dict[str, Any]:
    intended = _load_pricelist_l1(intended_active_path)
    actual = _load_pricelist_l1(redownloaded_active_path)
    updates = load_safe_patch_updates(updates_path)
    intended_index, _ = _unique_index(intended)
    actual_index, _ = _unique_index(actual)

    missing = sorted(set(intended_index) - set(actual_index))
    extra = sorted(set(actual_index) - set(intended_index))
    mismatches: list[dict[str, Any]] = []
    for update in updates:
        intended_row = intended_index.get(update.sku)
        actual_row = actual_index.get(update.sku)
        if intended_row is None or actual_row is None:
            continue
        for col in update.values:
            expected = _clean_cell(intended_row.get(col, ""))
            actual_value = _clean_cell(actual_row.get(col, ""))
            if expected != actual_value:
                mismatches.append({"SKU": update.sku, "field": col, "expected": expected, "actual": actual_value})

    return {
        "status": "ok" if not missing and not extra and not mismatches else "mismatch",
        "intended_active_count": int(len(intended)),
        "actual_active_count": int(len(actual)),
        "missing_count": int(len(missing)),
        "extra_count": int(len(extra)),
        "field_mismatch_count": int(len(mismatches)),
        "missing_skus": missing,
        "extra_skus": extra,
        "field_mismatches": mismatches,
    }
