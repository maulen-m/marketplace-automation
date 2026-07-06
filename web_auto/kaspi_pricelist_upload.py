from __future__ import annotations

import argparse
import json
import re
import time
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from playwright.sync_api import Page, TimeoutError as PlaywrightTimeoutError, sync_playwright

from .kaspi_forbidden_cards import OWNER_DECISION_LABEL, forbidden_saleable_row_details
from .kaspi_merchant_common import (
    login_kaspi_merchant,
    merchant_browser_launch_kwargs,
    normalize_store_name,
    resolve_store_credentials,
)
from .kaspi_price_floors import below_floor_details


PRICE_LIST_URL = "https://kaspi.kz/mc/#/price-list"
HISTORY_URL = "https://kaspi.kz/mc/#/history?page=1"


@dataclass(frozen=True)
class HistoryRow:
    filename: str
    status: str
    processed: int
    total: int
    uploaded_at: datetime
    detail_href: str


def parse_kaspi_history_datetime(value: str) -> datetime:
    return datetime.strptime(str(value or "").strip(), "%d.%m.%Y %H:%M")


def history_row_is_success(row: HistoryRow) -> bool:
    return row.status.strip().lower() == "файл загружен" and row.total > 0 and row.processed == row.total


def history_row_is_terminal_failure(row: HistoryRow) -> bool:
    return "ошибка загрузки файла" in row.status.strip().lower()


def history_row_is_processing(row: HistoryRow) -> bool:
    status = row.status.strip().lower()
    return "идет загрузка" in status or "идёт загрузка" in status or "обрабатывается" in status


def history_row_to_payload(row: HistoryRow) -> dict[str, Any]:
    payload = asdict(row)
    payload["uploaded_at"] = row.uploaded_at.isoformat()
    return payload


def detail_url_from_history_row(row: HistoryRow) -> str:
    return detail_url_from_history_ref(row.detail_href)


def detail_url_from_history_ref(value: str) -> str:
    ref = str(value or "").strip()
    if not ref:
        return ""
    if ref.startswith("http://") or ref.startswith("https://"):
        return ref
    if ref.startswith("#/"):
        return f"https://kaspi.kz/mc/{ref}"
    if "/history/detail/" in ref:
        return f"https://kaspi.kz/mc/{ref.lstrip('/')}"
    if re.fullmatch(r"[A-Za-z0-9_-]+", ref):
        return f"https://kaspi.kz/mc/#/history/detail/{ref}"
    return ref


def resolve_upload_file_paths(
    *,
    archive: str | Path | None = None,
    active: str | Path | None = None,
    files: list[str | Path] | None = None,
) -> list[Path]:
    ordered: list[Path] = []
    if files:
        for item in files:
            path = Path(item)
            if path not in ordered:
                ordered.append(path)
        if ordered:
            return ordered
    if archive:
        ordered.append(Path(archive))
    if active:
        active_path = Path(active)
        if active_path not in ordered:
            ordered.append(active_path)
    if not ordered:
        raise ValueError("at least one upload file is required")
    return ordered


def _load_pricelist_rows_for_guard(path: Path) -> list[dict[str, Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Лист1" not in wb.sheetnames:
            raise ValueError(f"missing sheet Лист1 in {path}")
        ws = wb["Лист1"]
        headers = [str(cell.value or "").strip() for cell in ws[1]]
        rows: list[dict[str, Any]] = []
        for row_number, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            record = {header: values[idx] if idx < len(values) else "" for idx, header in enumerate(headers) if header}
            if not any(str(value or "").strip() for value in record.values()):
                continue
            record["_row_number"] = row_number
            rows.append(record)
        return rows
    finally:
        wb.close()


def validate_pricelist_upload_files(file_paths: list[Path], *, store_name: str = "") -> dict[str, Any]:
    violations: list[dict[str, str]] = []
    floor_violations: list[dict[str, str]] = []
    scanned_files: list[str] = []
    for file_path in file_paths:
        path = Path(file_path)
        scanned_files.append(str(path))
        rows = _load_pricelist_rows_for_guard(path)
        for detail in forbidden_saleable_row_details(rows, store_name=store_name):
            detail = dict(detail)
            detail["file"] = str(path)
            violations.append(detail)
        for detail in below_floor_details(rows, store_name=store_name):
            detail = dict(detail)
            detail["file"] = str(path)
            floor_violations.append(detail)
    status_ok = not violations and not floor_violations
    errors = []
    if violations:
        owner_labels = sorted({row.get("owner_decision_label") or OWNER_DECISION_LABEL for row in violations})
        errors.append(f"{'; '.join(owner_labels)}: upload file would set forbidden Kaspi offer cards saleable")
    if floor_violations:
        errors.append("kaspi price floor guard: upload file contains mapped article prices below v7 floor")
    return {
        "status": "ok" if status_ok else "blocked",
        "owner_decision": "; ".join(sorted({row.get("owner_decision_label") or OWNER_DECISION_LABEL for row in violations}))
        if violations
        else "",
        "store_name": normalize_store_name(store_name),
        "scanned_files": scanned_files,
        "forbidden_saleable_rows_count": int(len(violations)),
        "forbidden_saleable_rows": violations,
        "price_floor_violations_count": int(len(floor_violations)),
        "price_floor_violations": floor_violations,
        "error": "; ".join(errors),
    }


def should_extend_history_poll_deadline(row: HistoryRow | None, *, already_extended: bool) -> bool:
    if already_extended or row is None:
        return False
    return history_row_is_processing(row)


def detect_upload_block_reason(body_text: str) -> str:
    text = " ".join(str(body_text or "").split())
    if not text:
        return ""
    match = re.search(r"Превышен лимит на выставление товара\.\s*Можно изменить в\s*(\d{1,2}:\d{2})", text, re.I)
    if match:
        return f"listing_limit_until_{match.group(1)}"
    if "неверный формат файла" in text.lower():
        return "invalid_file_format"
    return ""


def select_latest_matching_history_row(rows: list[HistoryRow], filename: str, *, after_dt: datetime) -> HistoryRow | None:
    basename = Path(filename).name
    matches = [row for row in rows if row.filename == basename and row.uploaded_at >= after_dt]
    if not matches:
        return None
    matches.sort(key=lambda row: row.uploaded_at, reverse=True)
    return matches[0]


def extract_detail_error_text_from_text(text: str) -> str:
    for line in text.splitlines():
        line = line.strip()
        if line.startswith("В ячейке "):
            return line
    return text[:3000].strip()


def _extract_detail_error_text(page: Page) -> str:
    return extract_detail_error_text_from_text(page.locator("body").inner_text())


def extract_detail_metrics_from_text(text: str) -> dict[str, int]:
    patterns = {
        "total": r"Всего товаров:\s*(\d+)",
        "unrecognized": r"Нераспознанные товары:\s*(\d+)",
        "restricted": r"Ограниченные товары:\s*(\d+)",
        "errors": r"Товары с ошибками:\s*(\d+)",
        "warnings": r"Товары с предупреждениями:\s*(\d+)",
        "unchanged": r"Товары без изменений:\s*(\d+)",
    }
    out: dict[str, int] = {}
    for key, pattern in patterns.items():
        match = re.search(pattern, text)
        out[key] = int(match.group(1)) if match else 0
    return out


def _extract_detail_metrics(page: Page) -> dict[str, int]:
    return extract_detail_metrics_from_text(page.locator("body").inner_text())


def _parse_history_rows(page: Page) -> list[HistoryRow]:
    rows: list[HistoryRow] = []
    table_rows = page.locator("table tbody tr")
    for i in range(table_rows.count()):
        tr = table_rows.nth(i)
        cells = [tr.locator("td").nth(j).inner_text().strip() for j in range(tr.locator("td").count())]
        if len(cells) < 4:
            continue
        filename = cells[0].strip()
        status = cells[1].strip()
        processed_match = re.search(r"(\d+)\s+из\s+(\d+)", cells[2])
        processed = int(processed_match.group(1)) if processed_match else 0
        total = int(processed_match.group(2)) if processed_match else 0
        uploaded_at = parse_kaspi_history_datetime(cells[3].strip())
        detail_href = ""
        for j in range(tr.locator("a").count()):
            href = tr.locator("a").nth(j).get_attribute("href") or ""
            if "/history/detail/" in href:
                detail_href = href
                break
        rows.append(
            HistoryRow(
                filename=filename,
                status=status,
                processed=processed,
                total=total,
                uploaded_at=uploaded_at,
                detail_href=detail_href,
            )
        )
    return rows


def _ensure_upload_page(page: Page) -> None:
    page.goto(PRICE_LIST_URL, wait_until="domcontentloaded", timeout=60000)
    page.wait_for_timeout(3000)
    if page.locator('input[type="file"]').count() == 0:
        nav_link = page.get_by_text("Загрузить прайс-лист", exact=True)
        if nav_link.count() > 0:
            nav_link.first.click()
            page.wait_for_timeout(3000)
    page.locator('input[type="file"]').wait_for(timeout=15000)
    page.get_by_text("Загрузить файл вручную").wait_for(timeout=15000)


def _snapshot_upload_page(page: Page) -> dict[str, Any]:
    body = page.locator("body").inner_text()
    return {
        "url": page.url,
        "body_excerpt": body[:3000],
        "input_count": page.locator('input[type="file"]').count(),
        "button_count": page.locator("button").count(),
    }


def _upload_one_file(
    page: Page,
    *,
    file_path: Path,
    run_dir: Path,
    timeout_seconds: int,
    processing_grace_seconds: int,
) -> dict[str, Any]:
    start_dt = datetime.now().replace(second=0, microsecond=0)
    _ensure_upload_page(page)
    before = _snapshot_upload_page(page)
    page.screenshot(path=str(run_dir / f"before_upload_{file_path.name}.png"), full_page=True)
    page.locator('input[type="file"]').set_input_files(str(file_path))
    page.wait_for_timeout(1000)
    page.screenshot(path=str(run_dir / f"after_attach_{file_path.name}.png"), full_page=True)
    page.get_by_text("Загрузить", exact=True).click()
    page.wait_for_timeout(2000)
    page.screenshot(path=str(run_dir / f"after_click_{file_path.name}.png"), full_page=True)
    block_reason = detect_upload_block_reason(page.locator("body").inner_text())
    if block_reason:
        return {
            "file": str(file_path),
            "started_at": start_dt.isoformat(timespec="minutes"),
            "before_page": before,
            "status": "blocked",
            "block_reason": block_reason,
        }
    try:
        close_button = page.get_by_text("Закрыть", exact=True)
        if close_button.count():
            close_button.first.click(timeout=5000)
            page.wait_for_timeout(1000)
            page.screenshot(path=str(run_dir / f"after_close_success_modal_{file_path.name}.png"), full_page=True)
    except PlaywrightTimeoutError:
        pass

    deadline = time.time() + max(timeout_seconds, 30)
    found: HistoryRow | None = None
    deadline_extended = False
    while time.time() < deadline:
        page.goto(HISTORY_URL, wait_until="domcontentloaded", timeout=60000)
        page.wait_for_timeout(3000)
        rows = _parse_history_rows(page)
        found = select_latest_matching_history_row(rows, file_path.name, after_dt=start_dt)
        if found and history_row_is_success(found):
            detail_metrics: dict[str, int] = {}
            detail_url = detail_url_from_history_row(found)
            if detail_url:
                try:
                    page.goto(detail_url, wait_until="domcontentloaded", timeout=60000)
                    page.wait_for_timeout(3000)
                    page.screenshot(path=str(run_dir / f"history_success_{file_path.name}.png"), full_page=True)
                    detail_metrics = _extract_detail_metrics(page)
                except PlaywrightTimeoutError:
                    detail_metrics = {}
            return {
                "file": str(file_path),
                "started_at": start_dt.isoformat(timespec="minutes"),
                "before_page": before,
                "history_row": history_row_to_payload(found),
                "detail_metrics": detail_metrics,
                "status": "success",
            }
        if found and history_row_is_terminal_failure(found):
            detail_error = ""
            detail_metrics: dict[str, int] = {}
            detail_url = detail_url_from_history_row(found)
            if detail_url:
                try:
                    page.goto(detail_url, wait_until="domcontentloaded", timeout=60000)
                    page.wait_for_timeout(3000)
                    page.screenshot(path=str(run_dir / f"history_error_{file_path.name}.png"), full_page=True)
                    detail_error = _extract_detail_error_text(page)
                    detail_metrics = _extract_detail_metrics(page)
                except PlaywrightTimeoutError:
                    detail_error = "failed_to_open_history_detail"
            return {
                "file": str(file_path),
                "started_at": start_dt.isoformat(timespec="minutes"),
                "before_page": before,
                "history_row": history_row_to_payload(found),
                "detail_error": detail_error,
                "detail_metrics": detail_metrics,
                "status": "failed",
            }
        if found:
            page.screenshot(path=str(run_dir / f"history_seen_{file_path.name}.png"), full_page=True)
            if should_extend_history_poll_deadline(found, already_extended=deadline_extended):
                deadline = max(deadline, time.time() + max(processing_grace_seconds, 60))
                deadline_extended = True
        time.sleep(5)

    return {
        "file": str(file_path),
        "started_at": start_dt.isoformat(timespec="minutes"),
        "before_page": before,
        "history_row": history_row_to_payload(found) if found else {},
        "status": "timeout",
    }


def run_kaspi_pricelist_upload(
    *,
    store_name: str,
    email: str,
    password: str,
    file_paths: list[Path],
    run_dir: Path,
    headless: bool,
    timeout_seconds: int = 900,
    processing_grace_seconds: int = 900,
) -> dict[str, Any]:
    run_dir.mkdir(parents=True, exist_ok=True)
    guard = validate_pricelist_upload_files(file_paths, store_name=store_name)
    summary: dict[str, Any] = {
        "run_dir": str(run_dir),
        "uploads": [],
        "status": "unknown",
        "forbidden_card_guard": guard,
    }
    if guard.get("status") != "ok":
        summary["status"] = "blocked"
        summary["store_name"] = normalize_store_name(store_name)
        summary["error"] = guard.get("error", "")
        (run_dir / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
        return summary
    with sync_playwright() as p:
        browser = p.chromium.launch(**merchant_browser_launch_kwargs(headless=headless))
        page = browser.new_page(viewport={"width": 1440, "height": 1400})
        login_kaspi_merchant(page, email=email, password=password)
        page.screenshot(path=str(run_dir / "logged_in_home.png"), full_page=True)
        for file_path in file_paths:
            result = _upload_one_file(
                page,
                file_path=file_path,
                run_dir=run_dir,
                timeout_seconds=timeout_seconds,
                processing_grace_seconds=processing_grace_seconds,
            )
            summary["uploads"].append(result)
            if result["status"] != "success":
                summary["status"] = "failed"
                summary["store_name"] = normalize_store_name(store_name)
                browser.close()
                return summary
        browser.close()
    summary["status"] = "success"
    summary["store_name"] = normalize_store_name(store_name)
    return summary


def run_kaspi_pricelist_history_detail(
    *,
    store_name: str,
    email: str,
    password: str,
    detail_ref: str,
    run_dir: Path,
    headless: bool,
    detail_filter: str = "",
    download_result_excel: bool = False,
) -> dict[str, Any]:
    run_dir.mkdir(parents=True, exist_ok=True)
    detail_url = detail_url_from_history_ref(detail_ref)
    if not detail_url:
        return {
            "store_name": normalize_store_name(store_name),
            "run_dir": str(run_dir),
            "status": "invalid_args",
            "error": "empty_history_detail_ref",
        }
    summary: dict[str, Any] = {
        "store_name": normalize_store_name(store_name),
        "run_dir": str(run_dir),
        "detail_ref": detail_ref,
        "detail_url": detail_url,
        "detail_filter": detail_filter or "all",
        "download_result_excel": download_result_excel,
        "status": "unknown",
        "production_write_action_executed": False,
    }
    with sync_playwright() as p:
        browser = p.chromium.launch(**merchant_browser_launch_kwargs(headless=headless))
        page = browser.new_page(viewport={"width": 1440, "height": 1400})
        login_kaspi_merchant(page, email=email, password=password)
        page.screenshot(path=str(run_dir / "logged_in_home.png"), full_page=True)
        page.goto(detail_url, wait_until="domcontentloaded", timeout=60000)
        page.wait_for_timeout(3000)
        filter_labels = {
            "unrecognized": "Нераспознанные товары",
            "restricted": "Ограниченные товары",
            "errors": "Товары с ошибками",
            "warnings": "Товары с предупреждениями",
            "all": "Всего товаров",
        }
        if detail_filter:
            filter_label = filter_labels.get(detail_filter)
            if not filter_label:
                browser.close()
                return {
                    **summary,
                    "status": "invalid_args",
                    "error": f"unknown_detail_filter:{detail_filter}",
                }
            filter_applied = page.evaluate(
                """
                (label) => {
                  for (const select of Array.from(document.querySelectorAll('select'))) {
                    for (const option of Array.from(select.options || [])) {
                      if ((option.textContent || '').includes(label)) {
                        select.value = option.value;
                        select.dispatchEvent(new Event('input', {bubbles: true}));
                        select.dispatchEvent(new Event('change', {bubbles: true}));
                        return true;
                      }
                    }
                  }
                  return false;
                }
                """,
                filter_label,
            )
            summary["filter_applied"] = bool(filter_applied)
            if not filter_applied:
                summary["filter_warning"] = f"filter_control_not_found:{filter_label}"
            page.wait_for_timeout(3000)
        page.screenshot(path=str(run_dir / "history_detail.png"), full_page=True)
        body_text = page.locator("body").inner_text()
        (run_dir / "history_detail_text.txt").write_text(body_text, encoding="utf-8")
        summary.update(
            {
                "status": "success",
                "final_url": page.url,
                "detail_metrics": extract_detail_metrics_from_text(body_text),
                "detail_error": extract_detail_error_text_from_text(body_text),
                "body_excerpt": body_text[:3000],
            }
        )
        if download_result_excel:
            try:
                with page.expect_download(timeout=60000) as download_info:
                    page.get_by_text("Выгрузить в EXCEL", exact=True).click()
                download = download_info.value
                suggested_name = download.suggested_filename or "history_detail_result.xlsx"
                export_path = run_dir / suggested_name
                download.save_as(str(export_path))
                summary["detail_export_path"] = str(export_path)
            except Exception as exc:  # pragma: no cover - transport/runtime guard
                summary["detail_export_error"] = f"{type(exc).__name__}: {exc}"
        browser.close()
    (run_dir / "history_detail_summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return summary


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--env-file", type=Path, default=Path("~/Docs/Autonomous_business/.env"))
    parser.add_argument("--store", default="STORE-B")
    parser.add_argument("--run-dir", type=Path, required=True)
    parser.add_argument("--file", action="append", dest="files", required=True)
    parser.add_argument("--headless", action="store_true")
    parser.add_argument("--timeout-seconds", type=int, default=900)
    parser.add_argument("--processing-grace-seconds", type=int, default=900)
    args = parser.parse_args()

    creds = resolve_store_credentials(args.store, args.env_file)

    summary = run_kaspi_pricelist_upload(
        store_name=creds["store_name"],
        email=creds["email"],
        password=creds["password"],
        file_paths=resolve_upload_file_paths(files=args.files),
        run_dir=args.run_dir,
        headless=args.headless,
        timeout_seconds=args.timeout_seconds,
        processing_grace_seconds=args.processing_grace_seconds,
    )
    summary_path = args.run_dir / "summary.json"
    summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0 if summary["status"] == "success" else 4


if __name__ == "__main__":
    raise SystemExit(main())
