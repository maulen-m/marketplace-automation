from __future__ import annotations

import csv
import json
import shutil
import re
import sqlite3
import tempfile
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from urllib.parse import urlsplit, urlunsplit
from zoneinfo import ZoneInfo

import yaml
from openpyxl import Workbook

from .repricer_items_export import export_repricer_items_to_sqlite


ALMATY_TZ = ZoneInfo("Asia/Almaty")
EXPECTED_STORE_CODES = {"UNIVERSAL", "ACMEWEAR", "MELVIS", "STOREB", "11KZ"}
DISPLAY_BY_CODE = {
    "UNIVERSAL": "UNIVERSAL",
    "ACMEWEAR": "ACMEWEAR",
    "MELVIS": "MELVIS",
    "STOREB": "STORE-B",
    "11KZ": "11KZ",
}
DEFAULT_LINKS_BASE_XLSX = (
    "~/Documents/useful tables/Main crm spreadsheets/main tables/links_database/SKU_link_base.xlsx"
)
DEFAULT_SCRAPE_CATALOG_XLSX = (
    "~/Documents/useful tables/Main crm spreadsheets/main tables/scrape/SKU_CATALOG.xlsx"
)
DEFAULT_SCRAPE_PRICEWARS_DIR = (
    "~/Documents/useful tables/Main crm spreadsheets/main tables/scrape/result"
)


def _now_pair() -> tuple[str, str]:
    now = datetime.now(ALMATY_TZ)
    return now.isoformat(), now.astimezone(UTC).isoformat()


def _quote_ident(name: str) -> str:
    return '"' + name.replace('"', '""') + '"'


def _sanitize_table_name(stem: str) -> str:
    normalized = re.sub(r"[^0-9a-zA-Z_]+", "_", stem).strip("_")
    if not normalized:
        normalized = "table"
    return f"raw_external_{normalized.lower()}"


def _normalize_store_token(value: Any) -> str | None:
    if value is None:
        return None
    token = re.sub(r"[^0-9A-Z]+", "", str(value).upper())
    if not token:
        return None
    if token in {"STOREB", "MGROU"}:
        return "STOREB"
    if token == "MELVIS":
        return "MELVIS"
    if token == "ACMEWEAR":
        return "ACMEWEAR"
    if token == "UNIVERSAL":
        return "UNIVERSAL"
    if token == "11KZ":
        return "11KZ"
    return None


def _normalize_store_code(
    *,
    store_code: Any,
    store_name: Any,
    store_id: Any,
    store_id_to_code: dict[int, str],
) -> str | None:
    code = _normalize_store_token(store_code)
    if code:
        return code
    code = _normalize_store_token(store_name)
    if code:
        return code
    try:
        sid = int(store_id) if store_id is not None and str(store_id).strip() != "" else None
    except ValueError:
        sid = None
    if sid is None:
        return None
    return store_id_to_code.get(sid)


def _normalize_link(value: Any) -> str | None:
    if value is None:
        return None
    raw = str(value).strip()
    if not raw:
        return None
    try:
        parsed = urlsplit(raw)
    except Exception:
        return raw
    if not parsed.scheme or not parsed.netloc:
        return raw
    normalized_path = parsed.path.rstrip("/") or parsed.path
    normalized = urlunsplit(
        (
            parsed.scheme.lower(),
            parsed.netloc.lower(),
            normalized_path,
            "",
            "",
        )
    )
    return normalized


def _normalize_identity_token(value: Any) -> str | None:
    if value is None:
        return None
    token = str(value).strip()
    if not token:
        return None
    return token.upper()


def _as_float(value: Any) -> float | None:
    if value is None:
        return None
    raw = str(value).strip()
    if raw == "":
        return None
    try:
        return float(raw)
    except ValueError:
        return None


def _as_int(value: Any) -> int | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return int(value)
    raw = str(value).strip()
    if raw == "":
        return None
    try:
        return int(float(raw))
    except ValueError:
        return None


def _iso_pair_from_value(value: Any, fallback_almaty: str, fallback_utc: str) -> tuple[str, str]:
    if value is None:
        return fallback_almaty, fallback_utc
    raw = str(value).strip()
    if not raw:
        return fallback_almaty, fallback_utc
    try:
        dt = datetime.fromisoformat(raw)
    except ValueError:
        return raw, fallback_utc
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=ALMATY_TZ)
    return dt.astimezone(ALMATY_TZ).isoformat(), dt.astimezone(UTC).isoformat()


def _compute_identity(link: Any, merchant_sku: Any, kaspi_sku: Any) -> tuple[str | None, str | None, str | None]:
    normalized_link = _normalize_link(link)
    if normalized_link:
        return "link", normalized_link, f"link:{normalized_link}"
    normalized_merchant = _normalize_identity_token(merchant_sku)
    if normalized_merchant:
        return "merchant_sku", normalized_merchant, f"merchant_sku:{normalized_merchant}"
    normalized_kaspi = _normalize_identity_token(kaspi_sku)
    if normalized_kaspi:
        return "kaspi_sku", normalized_kaspi, f"kaspi_sku:{normalized_kaspi}"
    return None, None, None


def _find_table_name(table_map: dict[str, str], prefix: str) -> str | None:
    matches: list[tuple[str, str]] = []
    for file_name, table_name in table_map.items():
        if file_name.startswith(prefix):
            matches.append((file_name, table_name))
    if not matches:
        return None
    matches.sort(key=lambda item: item[0])
    return matches[-1][1]


def _table_exists(conn: sqlite3.Connection, table_name: str) -> bool:
    row = conn.execute(
        "SELECT 1 FROM sqlite_master WHERE type='table' AND name=? LIMIT 1",
        (table_name,),
    ).fetchone()
    return row is not None


def _norm_header(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).strip()).lower()


def _header_index(header: tuple[Any, ...]) -> dict[str, int]:
    idx: dict[str, int] = {}
    for i, cell in enumerate(header):
        key = _norm_header(cell)
        if key and key not in idx:
            idx[key] = i
    return idx


def _cell_str(row: tuple[Any, ...], idx: int | None) -> str | None:
    if idx is None or idx >= len(row):
        return None
    value = row[idx]
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _backup_if_exists(path: Path) -> Path | None:
    if not path.exists():
        return None
    backup_dir = path.parent / "backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now(ALMATY_TZ).strftime("%Y%m%d_%H%M%S")
    backup_path = backup_dir / f"{path.stem}.{stamp}{path.suffix}"
    shutil.copy2(path, backup_path)
    return backup_path


def _save_workbook_atomic(workbook: Workbook, path: Path) -> None:
    _backup_if_exists(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(prefix=f"{path.stem}.", suffix=path.suffix, dir=path.parent, delete=False) as tmp:
        tmp_path = Path(tmp.name)
    try:
        workbook.save(tmp_path)
        tmp_path.replace(path)
    finally:
        if tmp_path.exists():
            tmp_path.unlink()


def _write_text_atomic(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(prefix=f"{path.stem}.", suffix=path.suffix, dir=path.parent, delete=False) as tmp:
        tmp_path = Path(tmp.name)
    try:
        tmp_path.write_text(content, encoding="utf-8")
        tmp_path.replace(path)
    finally:
        if tmp_path.exists():
            tmp_path.unlink()


def _load_store_maps(
    *,
    kaspi_accounts_path: str | Path,
    repricer_config_path: str | Path,
) -> tuple[dict[int, str], dict[str, int], dict[str, str]]:
    store_id_to_code: dict[int, str] = {}
    store_code_to_id: dict[str, int] = {}
    store_code_to_name: dict[str, str] = dict(DISPLAY_BY_CODE)

    kaspi_raw = yaml.safe_load(Path(kaspi_accounts_path).read_text(encoding="utf-8")) or {}
    for account in (kaspi_raw.get("kaspi", {}) or {}).get("accounts", []) or []:
        account_id = account.get("account_id")
        name = account.get("name")
        code = _normalize_store_token(name)
        if account_id is None or code is None:
            continue
        sid = int(str(account_id))
        store_id_to_code[sid] = code
        store_code_to_id[code] = sid
        store_code_to_name[code] = DISPLAY_BY_CODE.get(code, str(name))

    repricer_raw = yaml.safe_load(Path(repricer_config_path).read_text(encoding="utf-8")) or {}
    store_name_map = repricer_raw.get("store_name_map", {}) or {}
    for raw_name, raw_id in store_name_map.items():
        code = _normalize_store_token(raw_name)
        if code is None:
            continue
        sid = int(raw_id)
        store_id_to_code[sid] = code
        store_code_to_id.setdefault(code, sid)
        store_code_to_name.setdefault(code, DISPLAY_BY_CODE.get(code, str(raw_name)))

    return store_id_to_code, store_code_to_id, store_code_to_name


def _create_unified_tables(conn: sqlite3.Connection) -> None:
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS raw_repricer_items_snapshot (
            source_snapshot TEXT NOT NULL,
            source_priority INTEGER NOT NULL,
            fetched_at TEXT,
            fetched_at_almaty TEXT,
            fetched_at_utc TEXT,
            store_id INTEGER,
            store_name TEXT,
            store_code TEXT,
            row_id INTEGER,
            merchant_sku TEXT,
            kaspi_sku TEXT,
            merchant_title TEXT,
            link TEXT,
            price REAL,
            min_price REAL,
            max_price REAL,
            strategy TEXT,
            time_to_react TEXT,
            step TEXT,
            dumping INTEGER,
            active INTEGER,
            is_available INTEGER,
            preorder TEXT,
            position TEXT,
            brand TEXT,
            city_name TEXT,
            is_on_sale_flag INTEGER,
            raw_json TEXT
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS item_store_truth_history (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            source_name TEXT NOT NULL,
            source_priority INTEGER NOT NULL,
            source_snapshot TEXT,
            source_updated_at TEXT,
            fetched_at_almaty TEXT,
            fetched_at_utc TEXT,
            store_code TEXT,
            store_id INTEGER,
            store_name TEXT,
            row_id INTEGER,
            merchant_sku TEXT,
            kaspi_sku TEXT,
            merchant_title TEXT,
            kaspi_offer_name TEXT,
            link TEXT,
            price REAL,
            min_price REAL,
            max_price REAL,
            sku_key TEXT,
            sku_id TEXT,
            my_size TEXT,
            active INTEGER,
            is_available INTEGER,
            dumping INTEGER,
            is_on_sale_flag INTEGER,
            identity_method TEXT,
            identity_value TEXT,
            canonical_item_key TEXT,
            raw_source TEXT
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS item_store_truth_latest (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            history_id INTEGER NOT NULL,
            source_name TEXT NOT NULL,
            source_priority INTEGER NOT NULL,
            source_snapshot TEXT,
            source_updated_at TEXT,
            fetched_at_almaty TEXT,
            fetched_at_utc TEXT,
            store_code TEXT,
            store_id INTEGER,
            store_name TEXT,
            row_id INTEGER,
            merchant_sku TEXT,
            kaspi_sku TEXT,
            merchant_title TEXT,
            kaspi_offer_name TEXT,
            link TEXT,
            price REAL,
            min_price REAL,
            max_price REAL,
            sku_key TEXT,
            sku_id TEXT,
            my_size TEXT,
            active INTEGER,
            is_available INTEGER,
            dumping INTEGER,
            is_on_sale_flag INTEGER,
            identity_method TEXT,
            identity_value TEXT,
            canonical_item_key TEXT,
            raw_source TEXT
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS item_identity_bridge (
            canonical_item_key TEXT PRIMARY KEY,
            identity_method TEXT,
            identity_value TEXT,
            first_seen_at TEXT,
            last_seen_at TEXT,
            rows_count INTEGER,
            stores_count INTEGER,
            best_source_priority INTEGER,
            freshness_rank_source TEXT
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS export_meta (
            key TEXT PRIMARY KEY,
            value TEXT
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS raw_links_sku_base (
            sku TEXT,
            model TEXT,
            brand TEXT,
            price TEXT,
            shop_link TEXT,
            last_modified_date TEXT,
            source_document TEXT
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS raw_links_sku_catalog (
            store_name TEXT,
            sku_id TEXT,
            sku_id_ksp TEXT,
            sku_key TEXT,
            kaspi_offer_name TEXT,
            product_url TEXT,
            date_entered TEXT,
            sale_stop_date TEXT,
            sheet_name TEXT
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS raw_links_price_wars (
            sku_key TEXT,
            product_url TEXT,
            product_code TEXT,
            kaspi_offer_name TEXT,
            seller_name TEXT,
            price_kzt TEXT,
            sheet_name TEXT,
            file_name TEXT
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS link_backfill_map (
            match_key_type TEXT,
            match_key_value TEXT,
            resolved_link TEXT,
            resolved_from_source TEXT,
            source_priority INTEGER,
            created_at_almaty TEXT
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS item_store_truth_latest_enriched (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            latest_id INTEGER NOT NULL,
            fetched_at_almaty TEXT,
            fetched_at_utc TEXT,
            source_name TEXT,
            source_priority INTEGER,
            source_snapshot TEXT,
            source_updated_at TEXT,
            store_code TEXT,
            store_id INTEGER,
            store_name TEXT,
            row_id INTEGER,
            merchant_sku TEXT,
            kaspi_sku TEXT,
            merchant_title TEXT,
            kaspi_offer_name TEXT,
            link TEXT,
            effective_link TEXT,
            effective_link_source TEXT,
            is_link_backfilled INTEGER,
            price REAL,
            min_price REAL,
            max_price REAL,
            sku_key TEXT,
            sku_id TEXT,
            my_size TEXT,
            active INTEGER,
            is_available INTEGER,
            dumping INTEGER,
            is_on_sale_flag INTEGER,
            identity_method TEXT,
            identity_value TEXT,
            canonical_item_key TEXT,
            raw_source TEXT
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS control_scope_sold_90d (
            store_code TEXT,
            store_id INTEGER,
            sku_id TEXT,
            sku_key TEXT,
            last_sale_date TEXT,
            sold_units_90d REAL,
            sold_orders_90d INTEGER,
            latest_id INTEGER,
            merchant_sku TEXT,
            kaspi_sku TEXT,
            kaspi_offer_name TEXT,
            merchant_title TEXT,
            price REAL,
            min_price REAL,
            max_price REAL,
            link TEXT,
            effective_link TEXT,
            effective_link_source TEXT,
            is_link_backfilled INTEGER,
            control_ready_flag INTEGER
        )
        """
    )
    conn.execute(
        "CREATE INDEX IF NOT EXISTS idx_history_store ON item_store_truth_history(store_code, canonical_item_key)"
    )
    conn.execute(
        "CREATE INDEX IF NOT EXISTS idx_history_identity ON item_store_truth_history(identity_method, identity_value)"
    )
    conn.execute(
        "CREATE INDEX IF NOT EXISTS idx_link_backfill_key ON link_backfill_map(match_key_type, match_key_value)"
    )
    conn.execute(
        "CREATE INDEX IF NOT EXISTS idx_latest_enriched_store ON item_store_truth_latest_enriched(store_code, sku_id, sku_key)"
    )


def _reset_unified_tables(conn: sqlite3.Connection) -> None:
    conn.execute("DELETE FROM raw_repricer_items_snapshot")
    conn.execute("DELETE FROM item_store_truth_history")
    conn.execute("DELETE FROM item_store_truth_latest")
    conn.execute("DELETE FROM item_store_truth_latest_enriched")
    conn.execute("DELETE FROM item_identity_bridge")
    conn.execute("DELETE FROM raw_links_sku_base")
    conn.execute("DELETE FROM raw_links_sku_catalog")
    conn.execute("DELETE FROM raw_links_price_wars")
    conn.execute("DELETE FROM link_backfill_map")
    conn.execute("DELETE FROM control_scope_sold_90d")
    conn.execute("DELETE FROM export_meta")

    external_tables = conn.execute(
        """
        SELECT name
        FROM sqlite_master
        WHERE type = 'table'
          AND name LIKE 'raw_external_%'
        """
    ).fetchall()
    for (table_name,) in external_tables:
        conn.execute(f"DROP TABLE IF EXISTS {_quote_ident(table_name)}")


def _import_external_csvs(conn: sqlite3.Connection, external_truth_dir: str | Path) -> dict[str, dict[str, Any]]:
    base_dir = Path(external_truth_dir)
    if not base_dir.exists():
        raise RuntimeError(f"External truth directory not found: {base_dir}")

    imported: dict[str, dict[str, Any]] = {}
    for csv_path in sorted(base_dir.glob("*.csv")):
        table_name = _sanitize_table_name(csv_path.stem)
        with csv_path.open("r", encoding="utf-8", newline="") as fh:
            reader = csv.DictReader(fh)
            columns = reader.fieldnames or []
            if not columns:
                continue
            column_defs = ", ".join(f"{_quote_ident(col)} TEXT" for col in columns)
            conn.execute(f"CREATE TABLE {_quote_ident(table_name)} ({column_defs})")
            placeholders = ", ".join("?" for _ in columns)
            quoted_columns = ", ".join(_quote_ident(col) for col in columns)
            insert_sql = f"INSERT INTO {_quote_ident(table_name)} ({quoted_columns}) VALUES ({placeholders})"
            rows_inserted = 0
            batch: list[tuple[Any, ...]] = []
            for row in reader:
                batch.append(tuple(row.get(col) for col in columns))
                if len(batch) >= 1000:
                    conn.executemany(insert_sql, batch)
                    rows_inserted += len(batch)
                    batch.clear()
            if batch:
                conn.executemany(insert_sql, batch)
                rows_inserted += len(batch)
            imported[csv_path.name] = {"table_name": table_name, "rows": rows_inserted}
    return imported


def _import_links_sku_base(conn: sqlite3.Connection, path: str | Path) -> dict[str, Any]:
    workbook_path = Path(path)
    if not workbook_path.exists():
        return {"path": str(workbook_path), "rows": 0, "status": "missing"}

    from openpyxl import load_workbook

    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet_name = "Товары" if "Товары" in wb.sheetnames else wb.sheetnames[0]
        ws = wb[sheet_name]
        header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        idx = _header_index(header)

        rows = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            sku = _cell_str(row, idx.get("sku"))
            shop_link = _normalize_link(_cell_str(row, idx.get("shop_link")))
            if not sku and not shop_link:
                continue
            conn.execute(
                """
                INSERT INTO raw_links_sku_base (
                    sku, model, brand, price, shop_link, last_modified_date, source_document
                ) VALUES (?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    sku,
                    _cell_str(row, idx.get("model")),
                    _cell_str(row, idx.get("brand")),
                    _cell_str(row, idx.get("price")),
                    shop_link,
                    _cell_str(row, idx.get("last_modified_date")),
                    _cell_str(row, idx.get("source_document")),
                ),
            )
            rows += 1
    finally:
        wb.close()

    return {"path": str(workbook_path), "rows": rows, "status": "ok"}


def _import_links_sku_catalog(conn: sqlite3.Connection, path: str | Path) -> dict[str, Any]:
    workbook_path = Path(path)
    if not workbook_path.exists():
        return {"path": str(workbook_path), "rows": 0, "status": "missing"}

    from openpyxl import load_workbook

    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    total_rows = 0
    loaded_sheets: list[str] = []
    try:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
            idx = _header_index(header)
            if "product_url" not in idx:
                continue

            loaded_sheets.append(sheet_name)
            for row in ws.iter_rows(min_row=2, values_only=True):
                product_url = _normalize_link(_cell_str(row, idx.get("product_url")))
                if not product_url:
                    continue
                sku_id = _cell_str(row, idx.get("sku_id"))
                sku_id_ksp = _cell_str(row, idx.get("sku_id_ksp"))
                sku_key = _cell_str(row, idx.get("sku_key"))
                kaspi_offer_name = _cell_str(row, idx.get("kaspi_offer_name"))
                if not any((sku_id, sku_id_ksp, sku_key, kaspi_offer_name)):
                    continue
                conn.execute(
                    """
                    INSERT INTO raw_links_sku_catalog (
                        store_name, sku_id, sku_id_ksp, sku_key, kaspi_offer_name,
                        product_url, date_entered, sale_stop_date, sheet_name
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        _cell_str(row, idx.get("store_name")),
                        sku_id,
                        sku_id_ksp,
                        sku_key,
                        kaspi_offer_name,
                        product_url,
                        _cell_str(row, idx.get("date_entered")),
                        _cell_str(row, idx.get("sale_stop_date")),
                        sheet_name,
                    ),
                )
                total_rows += 1
    finally:
        wb.close()

    return {
        "path": str(workbook_path),
        "rows": total_rows,
        "loaded_sheets": loaded_sheets,
        "status": "ok",
    }


def _import_links_price_wars(conn: sqlite3.Connection, directory: str | Path) -> dict[str, Any]:
    result_dir = Path(directory)
    if not result_dir.exists():
        return {"path": str(result_dir), "rows": 0, "files": [], "status": "missing"}

    from openpyxl import load_workbook

    total_rows = 0
    loaded_files: list[str] = []
    for xlsx_path in sorted(result_dir.glob("*.xlsx")):
        if xlsx_path.name.startswith("~$"):
            continue
        wb = load_workbook(xlsx_path, read_only=True, data_only=True)
        try:
            file_rows = 0
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
                idx = _header_index(header)
                if "product_url" not in idx:
                    continue
                for row in ws.iter_rows(min_row=2, values_only=True):
                    product_url = _normalize_link(_cell_str(row, idx.get("product_url")))
                    sku_key = _cell_str(row, idx.get("sku_key"))
                    kaspi_offer_name = _cell_str(row, idx.get("kaspi_offer_name"))
                    if not product_url or not any((sku_key, kaspi_offer_name)):
                        continue
                    conn.execute(
                        """
                        INSERT INTO raw_links_price_wars (
                            sku_key, product_url, product_code, kaspi_offer_name,
                            seller_name, price_kzt, sheet_name, file_name
                        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                        """,
                        (
                            sku_key,
                            product_url,
                            _cell_str(row, idx.get("product_code")),
                            kaspi_offer_name,
                            _cell_str(row, idx.get("seller_name")),
                            _cell_str(row, idx.get("price_kzt")),
                            sheet_name,
                            xlsx_path.name,
                        ),
                    )
                    total_rows += 1
                    file_rows += 1
            if file_rows > 0:
                loaded_files.append(xlsx_path.name)
        finally:
            wb.close()

    return {"path": str(result_dir), "rows": total_rows, "files": loaded_files, "status": "ok"}


def _normalize_match_key(key_type: str, value: str | None) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    if key_type == "kaspi_offer_name_exact":
        return re.sub(r"\s+", " ", text).upper()
    return text.upper()


def _build_link_backfill_map(conn: sqlite3.Connection, created_at_almaty: str) -> dict[str, Any]:
    conn.execute("DELETE FROM link_backfill_map")
    candidates: dict[tuple[str, str], tuple[str, str, int]] = {}
    source_rows: dict[str, int] = {"sku_link_base": 0, "sku_catalog": 0, "price_wars": 0}

    def add_candidate(
        *,
        key_type: str,
        key_value: str | None,
        link: str | None,
        source: str,
        priority: int,
    ) -> None:
        normalized_key = _normalize_match_key(key_type, key_value)
        normalized_link = _normalize_link(link)
        if not normalized_key or not normalized_link:
            return
        key = (key_type, normalized_key)
        existing = candidates.get(key)
        if existing is None or priority < existing[2]:
            candidates[key] = (normalized_link, source, priority)

    for row in conn.execute("SELECT sku, shop_link FROM raw_links_sku_base"):
        add_candidate(
            key_type="merchant_sku",
            key_value=row[0],
            link=row[1],
            source="sku_link_base",
            priority=1,
        )

    for row in conn.execute(
        """
        SELECT sku_id_ksp, sku_id, sku_key, kaspi_offer_name, product_url
        FROM raw_links_sku_catalog
        """
    ):
        sku_id_ksp, sku_id, sku_key, kaspi_offer_name, product_url = row
        add_candidate(
            key_type="merchant_sku",
            key_value=sku_id_ksp,
            link=product_url,
            source="sku_catalog",
            priority=2,
        )
        add_candidate(
            key_type="sku_id",
            key_value=sku_id,
            link=product_url,
            source="sku_catalog",
            priority=2,
        )
        add_candidate(
            key_type="sku_key",
            key_value=sku_key,
            link=product_url,
            source="sku_catalog",
            priority=2,
        )
        add_candidate(
            key_type="kaspi_offer_name_exact",
            key_value=kaspi_offer_name,
            link=product_url,
            source="sku_catalog",
            priority=2,
        )

    for row in conn.execute("SELECT sku_key, kaspi_offer_name, product_url FROM raw_links_price_wars"):
        sku_key, kaspi_offer_name, product_url = row
        add_candidate(
            key_type="sku_key",
            key_value=sku_key,
            link=product_url,
            source="price_wars",
            priority=3,
        )
        add_candidate(
            key_type="kaspi_offer_name_exact",
            key_value=kaspi_offer_name,
            link=product_url,
            source="price_wars",
            priority=3,
        )

    for (key_type, key_value), (resolved_link, resolved_from_source, source_priority) in candidates.items():
        conn.execute(
            """
            INSERT INTO link_backfill_map (
                match_key_type, match_key_value, resolved_link, resolved_from_source, source_priority, created_at_almaty
            ) VALUES (?, ?, ?, ?, ?, ?)
            """,
            (
                key_type,
                key_value,
                resolved_link,
                resolved_from_source,
                source_priority,
                created_at_almaty,
            ),
        )
        source_rows[resolved_from_source] = source_rows.get(resolved_from_source, 0) + 1

    return {"rows": len(candidates), "source_rows": source_rows}


def _build_latest_enriched(conn: sqlite3.Connection) -> dict[str, int]:
    conn.execute("DELETE FROM item_store_truth_latest_enriched")

    lookup: dict[str, dict[str, tuple[str, str]]] = {}
    for key_type, key_value, resolved_link, resolved_from_source in conn.execute(
        """
        SELECT match_key_type, match_key_value, resolved_link, resolved_from_source
        FROM link_backfill_map
        ORDER BY source_priority ASC, rowid ASC
        """
    ):
        key_type_map = lookup.setdefault(str(key_type), {})
        key = str(key_value)
        if key not in key_type_map:
            key_type_map[key] = (str(resolved_link), str(resolved_from_source))

    rows_inserted = 0
    backfilled = 0
    native = 0
    latest_rows = conn.execute("SELECT * FROM item_store_truth_latest ORDER BY id").fetchall()
    for row in latest_rows:
        native_link = _normalize_link(row["link"])
        effective_link = native_link
        effective_link_source = "native" if native_link else None
        is_link_backfilled = 0

        if native_link:
            native += 1
        else:
            merchant_key = _normalize_match_key("merchant_sku", row["merchant_sku"])
            sku_id_key = _normalize_match_key("sku_id", row["sku_id"])
            sku_key_key = _normalize_match_key("sku_key", row["sku_key"])
            offer_key = _normalize_match_key("kaspi_offer_name_exact", row["kaspi_offer_name"] or row["merchant_title"])

            resolved: tuple[str, str] | None = None
            if merchant_key and merchant_key in lookup.get("merchant_sku", {}):
                resolved = lookup["merchant_sku"][merchant_key]
            elif sku_id_key and sku_id_key in lookup.get("sku_id", {}):
                resolved = lookup["sku_id"][sku_id_key]
            elif sku_key_key and sku_key_key in lookup.get("sku_key", {}):
                resolved = lookup["sku_key"][sku_key_key]
            elif offer_key and offer_key in lookup.get("kaspi_offer_name_exact", {}):
                resolved = lookup["kaspi_offer_name_exact"][offer_key]

            if resolved:
                effective_link = resolved[0]
                effective_link_source = resolved[1]
                is_link_backfilled = 1
                backfilled += 1

        conn.execute(
            """
            INSERT INTO item_store_truth_latest_enriched (
                latest_id, fetched_at_almaty, fetched_at_utc, source_name, source_priority, source_snapshot, source_updated_at,
                store_code, store_id, store_name, row_id, merchant_sku, kaspi_sku, merchant_title, kaspi_offer_name,
                link, effective_link, effective_link_source, is_link_backfilled, price, min_price, max_price, sku_key, sku_id,
                my_size, active, is_available, dumping, is_on_sale_flag, identity_method, identity_value, canonical_item_key,
                raw_source
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                row["id"],
                row["fetched_at_almaty"],
                row["fetched_at_utc"],
                row["source_name"],
                row["source_priority"],
                row["source_snapshot"],
                row["source_updated_at"],
                row["store_code"],
                row["store_id"],
                row["store_name"],
                row["row_id"],
                row["merchant_sku"],
                row["kaspi_sku"],
                row["merchant_title"],
                row["kaspi_offer_name"],
                row["link"],
                effective_link,
                effective_link_source,
                is_link_backfilled,
                row["price"],
                row["min_price"],
                row["max_price"],
                row["sku_key"],
                row["sku_id"],
                row["my_size"],
                row["active"],
                row["is_available"],
                row["dumping"],
                row["is_on_sale_flag"],
                row["identity_method"],
                row["identity_value"],
                row["canonical_item_key"],
                row["raw_source"],
            ),
        )
        rows_inserted += 1

    return {"rows": rows_inserted, "native_link_rows": native, "backfilled_link_rows": backfilled}


def _build_control_scope_sold_window(
    conn: sqlite3.Connection,
    sales_table_name: str | None,
    *,
    window_days: int,
) -> dict[str, Any]:
    conn.execute("DELETE FROM control_scope_sold_90d")
    if not sales_table_name:
        return {
            "sales_table_name": None,
            "window_days": window_days,
            "total_rows": 0,
            "control_ready_rows": 0,
            "missing_link_before": 0,
            "missing_link_after": 0,
            "backfilled_rows": 0,
        }

    enriched_rows = conn.execute(
        """
        SELECT *
        FROM item_store_truth_latest_enriched
        ORDER BY
            CASE WHEN COALESCE(effective_link, '') <> '' THEN 0 ELSE 1 END,
            source_priority ASC,
            latest_id DESC
        """
    ).fetchall()

    by_store_sku_id: dict[tuple[str, str], sqlite3.Row] = {}
    by_store_sku_key: dict[tuple[str, str], sqlite3.Row] = {}
    for row in enriched_rows:
        store_code = _normalize_store_token(row["store_code"])
        sku_id = _normalize_match_key("sku_id", row["sku_id"])
        sku_key = _normalize_match_key("sku_key", row["sku_key"])
        if store_code and sku_id and (store_code, sku_id) not in by_store_sku_id:
            by_store_sku_id[(store_code, sku_id)] = row
        if store_code and sku_key and (store_code, sku_key) not in by_store_sku_key:
            by_store_sku_key[(store_code, sku_key)] = row

    cutoff_sql = f"-{int(window_days)} days"
    sales_rows = conn.execute(
        f"""
        SELECT
            UPPER(TRIM(store_code)) AS store_code,
            NULLIF(TRIM(sku_id), '') AS sku_id,
            NULLIF(TRIM(sku_key), '') AS sku_key,
            MAX(date(sale_date)) AS last_sale_date,
            SUM(CAST(COALESCE(units, 0) AS REAL)) AS sold_units_90d,
            COUNT(DISTINCT order_id) AS sold_orders_90d
        FROM {_quote_ident(sales_table_name)}
        WHERE date(sale_date) >= date('now', ?)
        GROUP BY UPPER(TRIM(store_code)), NULLIF(TRIM(sku_id), ''), NULLIF(TRIM(sku_key), '')
        """,
        (cutoff_sql,),
    ).fetchall()

    total_rows = 0
    control_ready_rows = 0
    missing_before = 0
    missing_after = 0
    backfilled_rows = 0
    for row in sales_rows:
        store_code = _normalize_store_token(row["store_code"])
        sku_id_key = _normalize_match_key("sku_id", row["sku_id"])
        sku_key_key = _normalize_match_key("sku_key", row["sku_key"])
        latest_row = None
        if store_code and sku_id_key:
            latest_row = by_store_sku_id.get((store_code, sku_id_key))
        if latest_row is None and store_code and sku_key_key:
            latest_row = by_store_sku_key.get((store_code, sku_key_key))

        link = latest_row["link"] if latest_row is not None else None
        effective_link = latest_row["effective_link"] if latest_row is not None else None
        is_backfilled = int(latest_row["is_link_backfilled"]) if latest_row is not None else 0
        ready = int(bool(effective_link))

        conn.execute(
            """
            INSERT INTO control_scope_sold_90d (
                store_code, store_id, sku_id, sku_key, last_sale_date, sold_units_90d, sold_orders_90d, latest_id,
                merchant_sku, kaspi_sku, kaspi_offer_name, merchant_title, price, min_price, max_price,
                link, effective_link, effective_link_source, is_link_backfilled, control_ready_flag
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                store_code,
                latest_row["store_id"] if latest_row is not None else None,
                row["sku_id"],
                row["sku_key"],
                row["last_sale_date"],
                row["sold_units_90d"],
                row["sold_orders_90d"],
                latest_row["latest_id"] if latest_row is not None else None,
                latest_row["merchant_sku"] if latest_row is not None else None,
                latest_row["kaspi_sku"] if latest_row is not None else None,
                latest_row["kaspi_offer_name"] if latest_row is not None else None,
                latest_row["merchant_title"] if latest_row is not None else None,
                latest_row["price"] if latest_row is not None else None,
                latest_row["min_price"] if latest_row is not None else None,
                latest_row["max_price"] if latest_row is not None else None,
                link,
                effective_link,
                latest_row["effective_link_source"] if latest_row is not None else None,
                is_backfilled,
                ready,
            ),
        )

        total_rows += 1
        control_ready_rows += ready
        if not link:
            missing_before += 1
        if not effective_link:
            missing_after += 1
        if is_backfilled:
            backfilled_rows += 1

    return {
        "sales_table_name": sales_table_name,
        "window_days": window_days,
        "total_rows": total_rows,
        "control_ready_rows": control_ready_rows,
        "missing_link_before": missing_before,
        "missing_link_after": missing_after,
        "backfilled_rows": backfilled_rows,
    }


def _ingest_repricer_snapshot(
    *,
    conn: sqlite3.Connection,
    snapshot_path: Path,
    source_snapshot: str,
    source_priority: int,
    store_id_to_code: dict[int, str],
    store_code_to_name: dict[str, str],
    fallback_almaty: str,
    fallback_utc: str,
) -> int:
    if not snapshot_path.exists():
        return 0
    source_conn = sqlite3.connect(snapshot_path)
    source_conn.row_factory = sqlite3.Row
    try:
        tables = {
            row[0]
            for row in source_conn.execute(
                "SELECT name FROM sqlite_master WHERE type='table' AND name='repricer_items'"
            ).fetchall()
        }
        if "repricer_items" not in tables:
            return 0
        rows_inserted = 0
        cursor = source_conn.execute("SELECT * FROM repricer_items")
        for row in cursor:
            store_id = _as_int(row["store_id"])
            store_code = _normalize_store_code(
                store_code=None,
                store_name=row["store_name"],
                store_id=store_id,
                store_id_to_code=store_id_to_code,
            )
            if store_code is None and store_id is not None:
                store_code = store_id_to_code.get(store_id)
            store_name = store_code_to_name.get(store_code, row["store_name"])
            fetched_at_almaty, fetched_at_utc = _iso_pair_from_value(
                row["fetched_at"], fallback_almaty, fallback_utc
            )
            active = _as_int(row["active"])
            is_available = _as_int(row["is_available"])
            is_on_sale_flag = int(bool((active or 0) or (is_available or 0)))
            conn.execute(
                """
                INSERT INTO raw_repricer_items_snapshot (
                    source_snapshot, source_priority, fetched_at, fetched_at_almaty, fetched_at_utc,
                    store_id, store_name, store_code, row_id, merchant_sku, kaspi_sku,
                    merchant_title, link, price, min_price, max_price, strategy,
                    time_to_react, step, dumping, active, is_available, preorder, position,
                    brand, city_name, is_on_sale_flag, raw_json
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    source_snapshot,
                    source_priority,
                    row["fetched_at"],
                    fetched_at_almaty,
                    fetched_at_utc,
                    store_id,
                    store_name,
                    store_code,
                    _as_int(row["row_id"]),
                    row["merchant_sku"],
                    row["kaspi_sku"],
                    row["merchant_title"],
                    row["link"],
                    _as_float(row["price"]),
                    _as_float(row["min_price"]),
                    _as_float(row["max_price"]),
                    row["strategy"],
                    row["time_to_react"],
                    row["step"],
                    _as_int(row["dumping"]),
                    active,
                    is_available,
                    row["preorder"],
                    row["position"],
                    row["brand"],
                    row["city_name"],
                    is_on_sale_flag,
                    row["raw_json"],
                ),
            )
            rows_inserted += 1
    finally:
        source_conn.close()
    return rows_inserted


def _build_article_lookup(
    *,
    conn: sqlite3.Connection,
    table_name: str | None,
) -> tuple[dict[tuple[str, str], str], dict[tuple[str, str], str]]:
    by_store_sku_id: dict[tuple[str, str], str] = {}
    by_store_sku_key: dict[tuple[str, str], str] = {}
    if not table_name:
        return by_store_sku_id, by_store_sku_key

    query = f"""
        SELECT store_code, sku_id, sku_key, kaspi_article
        FROM {_quote_ident(table_name)}
        WHERE COALESCE(kaspi_article, '') <> ''
    """
    for store_code, sku_id, sku_key, kaspi_article in conn.execute(query):
        canonical_store = _normalize_store_token(store_code)
        article = str(kaspi_article).strip()
        if not canonical_store or not article:
            continue
        if sku_id and str(sku_id).strip():
            key = (canonical_store, str(sku_id).strip())
            by_store_sku_id.setdefault(key, article)
        if sku_key and str(sku_key).strip():
            key = (canonical_store, str(sku_key).strip())
            by_store_sku_key.setdefault(key, article)
    return by_store_sku_id, by_store_sku_key


def _build_size_lookup(conn: sqlite3.Connection, table_name: str | None) -> dict[str, str]:
    mapping: dict[str, str] = {}
    if not table_name:
        return mapping
    query = f"""
        SELECT sku_id, my_size
        FROM {_quote_ident(table_name)}
        WHERE COALESCE(sku_id, '') <> ''
    """
    for sku_id, my_size in conn.execute(query):
        sku = str(sku_id).strip()
        size = str(my_size).strip() if my_size is not None else ""
        if sku and size:
            mapping.setdefault(sku, size)
    return mapping


def _insert_history_row(
    *,
    conn: sqlite3.Connection,
    record: dict[str, Any],
    store_id_to_code: dict[int, str],
    store_code_to_id: dict[str, int],
    store_code_to_name: dict[str, str],
    default_almaty: str,
    default_utc: str,
) -> None:
    store_id = _as_int(record.get("store_id"))
    store_code = _normalize_store_code(
        store_code=record.get("store_code"),
        store_name=record.get("store_name"),
        store_id=store_id,
        store_id_to_code=store_id_to_code,
    )
    if store_code is None and store_id is not None:
        store_code = store_id_to_code.get(store_id)
    if store_id is None and store_code is not None:
        store_id = store_code_to_id.get(store_code)

    store_name = record.get("store_name")
    if store_code and (store_name is None or str(store_name).strip() == ""):
        store_name = store_code_to_name.get(store_code, DISPLAY_BY_CODE.get(store_code, store_code))

    source_updated_at = record.get("source_updated_at")
    fetched_at_almaty, fetched_at_utc = _iso_pair_from_value(
        record.get("fetched_at_almaty") or source_updated_at,
        default_almaty,
        default_utc,
    )

    active = _as_int(record.get("active"))
    is_available = _as_int(record.get("is_available"))
    is_on_sale_flag = record.get("is_on_sale_flag")
    if is_on_sale_flag is None:
        if active is None and is_available is None:
            is_on_sale_flag = None
        else:
            is_on_sale_flag = int(bool((active or 0) or (is_available or 0)))
    else:
        is_on_sale_flag = _as_int(is_on_sale_flag)

    identity_method, identity_value, canonical_item_key = _compute_identity(
        record.get("link"),
        record.get("merchant_sku"),
        record.get("kaspi_sku"),
    )

    raw_source = record.get("raw_source")
    if raw_source is not None and not isinstance(raw_source, str):
        raw_source = json.dumps(raw_source, ensure_ascii=False)

    conn.execute(
        """
        INSERT INTO item_store_truth_history (
            source_name, source_priority, source_snapshot, source_updated_at,
            fetched_at_almaty, fetched_at_utc, store_code, store_id, store_name,
            row_id, merchant_sku, kaspi_sku, merchant_title, kaspi_offer_name, link,
            price, min_price, max_price, sku_key, sku_id, my_size, active, is_available,
            dumping, is_on_sale_flag, identity_method, identity_value, canonical_item_key, raw_source
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            record.get("source_name"),
            _as_int(record.get("source_priority")) or 99,
            record.get("source_snapshot"),
            source_updated_at,
            fetched_at_almaty,
            fetched_at_utc,
            store_code,
            store_id,
            store_name,
            _as_int(record.get("row_id")),
            record.get("merchant_sku"),
            record.get("kaspi_sku"),
            record.get("merchant_title"),
            record.get("kaspi_offer_name"),
            record.get("link"),
            _as_float(record.get("price")),
            _as_float(record.get("min_price")),
            _as_float(record.get("max_price")),
            record.get("sku_key"),
            record.get("sku_id"),
            record.get("my_size"),
            active,
            is_available,
            _as_int(record.get("dumping")),
            is_on_sale_flag,
            identity_method,
            identity_value,
            canonical_item_key,
            raw_source,
        ),
    )


def _build_history_and_latest(
    *,
    conn: sqlite3.Connection,
    external_tables: dict[str, dict[str, Any]],
    store_id_to_code: dict[int, str],
    store_code_to_id: dict[str, int],
    store_code_to_name: dict[str, str],
    default_almaty: str,
    default_utc: str,
) -> dict[str, int]:
    table_map = {name: data["table_name"] for name, data in external_tables.items()}
    article_table = _find_table_name(table_map, "dim_kaspi_article_map_5stores_active")
    sku_size_table = _find_table_name(table_map, "dim_sku_size_active")
    store_offer_table = _find_table_name(table_map, "store_offer_sku_identity_current")
    sales_7m_table = _find_table_name(table_map, "sales_last_7_months_db_")
    sales_fact_table = _find_table_name(table_map, "sales_fact_v2_5stores_key_fields")
    orders_table = _find_table_name(table_map, "fact_orders_kaspi_5stores_key_fields")

    article_by_sku_id, article_by_sku_key = _build_article_lookup(conn=conn, table_name=article_table)
    size_by_sku_id = _build_size_lookup(conn, sku_size_table)

    inserted = {
        "history_from_repricer": 0,
        "history_from_dim_article": 0,
        "history_from_store_offer_identity": 0,
        "history_from_sales_7m": 0,
        "history_from_sales_fact": 0,
        "history_from_orders": 0,
    }

    repricer_rows = conn.execute(
        "SELECT * FROM raw_repricer_items_snapshot ORDER BY source_priority ASC, fetched_at_almaty DESC"
    ).fetchall()
    for row in repricer_rows:
        _insert_history_row(
            conn=conn,
            record={
                "source_name": "repricer_snapshot",
                "source_priority": row["source_priority"],
                "source_snapshot": row["source_snapshot"],
                "source_updated_at": row["fetched_at_almaty"],
                "fetched_at_almaty": row["fetched_at_almaty"],
                "fetched_at_utc": row["fetched_at_utc"],
                "store_code": row["store_code"],
                "store_id": row["store_id"],
                "store_name": row["store_name"],
                "row_id": row["row_id"],
                "merchant_sku": row["merchant_sku"],
                "kaspi_sku": row["kaspi_sku"],
                "merchant_title": row["merchant_title"],
                "kaspi_offer_name": row["merchant_title"],
                "link": row["link"],
                "price": row["price"],
                "min_price": row["min_price"],
                "max_price": row["max_price"],
                "active": row["active"],
                "is_available": row["is_available"],
                "dumping": row["dumping"],
                "is_on_sale_flag": row["is_on_sale_flag"],
                "raw_source": {
                    "source_snapshot": row["source_snapshot"],
                    "source": "repricer_items",
                },
            },
            store_id_to_code=store_id_to_code,
            store_code_to_id=store_code_to_id,
            store_code_to_name=store_code_to_name,
            default_almaty=default_almaty,
            default_utc=default_utc,
        )
        inserted["history_from_repricer"] += 1

    if article_table:
        query = f"""
            SELECT store_code, kaspi_article, kaspi_offer_name, sku_key, sku_id, updated_at, id
            FROM {_quote_ident(article_table)}
        """
        for store_code, kaspi_article, kaspi_offer_name, sku_key, sku_id, updated_at, row_id in conn.execute(query):
            _insert_history_row(
                conn=conn,
                record={
                    "source_name": "external_dim_kaspi_article_map",
                    "source_priority": 4,
                    "source_snapshot": "external_db_truth",
                    "source_updated_at": updated_at,
                    "store_code": store_code,
                    "merchant_sku": kaspi_article,
                    "merchant_title": kaspi_offer_name,
                    "kaspi_offer_name": kaspi_offer_name,
                    "sku_key": sku_key,
                    "sku_id": sku_id,
                    "my_size": size_by_sku_id.get(str(sku_id).strip(), None) if sku_id else None,
                    "row_id": row_id,
                    "raw_source": {"table": article_table, "id": row_id},
                },
                store_id_to_code=store_id_to_code,
                store_code_to_id=store_code_to_id,
                store_code_to_name=store_code_to_name,
                default_almaty=default_almaty,
                default_utc=default_utc,
            )
            inserted["history_from_dim_article"] += 1

    if store_offer_table:
        query = f"""
            SELECT store_code, kaspi_offer_name, sku_key, sku_id, my_size, avg_sell_price_kzt, first_seen, last_seen
            FROM {_quote_ident(store_offer_table)}
        """
        for store_code, offer_name, sku_key, sku_id, my_size, avg_price, first_seen, last_seen in conn.execute(query):
            canonical_store = _normalize_store_token(store_code)
            sku_id_norm = str(sku_id).strip() if sku_id else None
            sku_key_norm = str(sku_key).strip() if sku_key else None
            merchant_sku = None
            if canonical_store and sku_id_norm:
                merchant_sku = article_by_sku_id.get((canonical_store, sku_id_norm))
            if merchant_sku is None and canonical_store and sku_key_norm:
                merchant_sku = article_by_sku_key.get((canonical_store, sku_key_norm))
            _insert_history_row(
                conn=conn,
                record={
                    "source_name": "external_store_offer_sku_identity_current",
                    "source_priority": 3,
                    "source_snapshot": "external_db_truth",
                    "source_updated_at": last_seen or first_seen,
                    "store_code": store_code,
                    "merchant_sku": merchant_sku,
                    "merchant_title": offer_name,
                    "kaspi_offer_name": offer_name,
                    "sku_key": sku_key,
                    "sku_id": sku_id,
                    "my_size": my_size or (size_by_sku_id.get(sku_id_norm) if sku_id_norm else None),
                    "price": avg_price,
                    "raw_source": {"table": store_offer_table},
                },
                store_id_to_code=store_id_to_code,
                store_code_to_id=store_code_to_id,
                store_code_to_name=store_code_to_name,
                default_almaty=default_almaty,
                default_utc=default_utc,
            )
            inserted["history_from_store_offer_identity"] += 1

    if sales_7m_table:
        query = f"""
            SELECT store_code, sku_key, sku_id, my_size, units, net_rev_kzt, sale_date, order_id
            FROM {_quote_ident(sales_7m_table)}
        """
        for store_code, sku_key, sku_id, my_size, units, net_rev_kzt, sale_date, order_id in conn.execute(query):
            canonical_store = _normalize_store_token(store_code)
            sku_id_norm = str(sku_id).strip() if sku_id else None
            sku_key_norm = str(sku_key).strip() if sku_key else None
            merchant_sku = None
            if canonical_store and sku_id_norm:
                merchant_sku = article_by_sku_id.get((canonical_store, sku_id_norm))
            if merchant_sku is None and canonical_store and sku_key_norm:
                merchant_sku = article_by_sku_key.get((canonical_store, sku_key_norm))
            units_val = _as_float(units)
            net_rev_val = _as_float(net_rev_kzt)
            price = None
            if units_val and units_val > 0 and net_rev_val is not None:
                price = net_rev_val / units_val
            _insert_history_row(
                conn=conn,
                record={
                    "source_name": "external_sales_last_7_months",
                    "source_priority": 2,
                    "source_snapshot": "external_db_truth",
                    "source_updated_at": sale_date,
                    "store_code": store_code,
                    "merchant_sku": merchant_sku,
                    "sku_key": sku_key,
                    "sku_id": sku_id,
                    "my_size": my_size or (size_by_sku_id.get(sku_id_norm) if sku_id_norm else None),
                    "price": price,
                    "raw_source": {"table": sales_7m_table, "order_id": order_id},
                },
                store_id_to_code=store_id_to_code,
                store_code_to_id=store_code_to_id,
                store_code_to_name=store_code_to_name,
                default_almaty=default_almaty,
                default_utc=default_utc,
            )
            inserted["history_from_sales_7m"] += 1

    if sales_fact_table:
        query = f"""
            SELECT store_code, kaspi_offer_name, sku_key, sku_id, my_size, sell_price_kzt, order_date, order_id
            FROM {_quote_ident(sales_fact_table)}
        """
        for store_code, offer_name, sku_key, sku_id, my_size, sell_price, order_date, order_id in conn.execute(query):
            canonical_store = _normalize_store_token(store_code)
            sku_id_norm = str(sku_id).strip() if sku_id else None
            sku_key_norm = str(sku_key).strip() if sku_key else None
            merchant_sku = None
            if canonical_store and sku_id_norm:
                merchant_sku = article_by_sku_id.get((canonical_store, sku_id_norm))
            if merchant_sku is None and canonical_store and sku_key_norm:
                merchant_sku = article_by_sku_key.get((canonical_store, sku_key_norm))
            _insert_history_row(
                conn=conn,
                record={
                    "source_name": "external_sales_fact_v2",
                    "source_priority": 5,
                    "source_snapshot": "external_db_truth",
                    "source_updated_at": order_date,
                    "store_code": store_code,
                    "merchant_sku": merchant_sku,
                    "merchant_title": offer_name,
                    "kaspi_offer_name": offer_name,
                    "sku_key": sku_key,
                    "sku_id": sku_id,
                    "my_size": my_size or (size_by_sku_id.get(sku_id_norm) if sku_id_norm else None),
                    "price": sell_price,
                    "raw_source": {"table": sales_fact_table, "order_id": order_id},
                },
                store_id_to_code=store_id_to_code,
                store_code_to_id=store_code_to_id,
                store_code_to_name=store_code_to_name,
                default_almaty=default_almaty,
                default_utc=default_utc,
            )
            inserted["history_from_sales_fact"] += 1

    if orders_table:
        query = f"""
            SELECT store_code, kaspi_offer_name, sku_key, sku_id, my_size, assigned_size, unit_price_kzt,
                   created_at, updated_at, order_id
            FROM {_quote_ident(orders_table)}
        """
        for (
            store_code,
            offer_name,
            sku_key,
            sku_id,
            my_size,
            assigned_size,
            unit_price,
            created_at,
            updated_at,
            order_id,
        ) in conn.execute(query):
            canonical_store = _normalize_store_token(store_code)
            sku_id_norm = str(sku_id).strip() if sku_id else None
            sku_key_norm = str(sku_key).strip() if sku_key else None
            merchant_sku = None
            if canonical_store and sku_id_norm:
                merchant_sku = article_by_sku_id.get((canonical_store, sku_id_norm))
            if merchant_sku is None and canonical_store and sku_key_norm:
                merchant_sku = article_by_sku_key.get((canonical_store, sku_key_norm))
            effective_size = assigned_size or my_size or (size_by_sku_id.get(sku_id_norm) if sku_id_norm else None)
            _insert_history_row(
                conn=conn,
                record={
                    "source_name": "external_fact_orders_kaspi",
                    "source_priority": 6,
                    "source_snapshot": "external_db_truth",
                    "source_updated_at": created_at or updated_at,
                    "store_code": store_code,
                    "merchant_sku": merchant_sku,
                    "merchant_title": offer_name,
                    "kaspi_offer_name": offer_name,
                    "sku_key": sku_key,
                    "sku_id": sku_id,
                    "my_size": effective_size,
                    "price": unit_price,
                    "raw_source": {"table": orders_table, "order_id": order_id},
                },
                store_id_to_code=store_id_to_code,
                store_code_to_id=store_code_to_id,
                store_code_to_name=store_code_to_name,
                default_almaty=default_almaty,
                default_utc=default_utc,
            )
            inserted["history_from_orders"] += 1

    conn.execute("DELETE FROM item_store_truth_latest")
    conn.execute(
        """
        INSERT INTO item_store_truth_latest (
            history_id, source_name, source_priority, source_snapshot, source_updated_at,
            fetched_at_almaty, fetched_at_utc, store_code, store_id, store_name, row_id,
            merchant_sku, kaspi_sku, merchant_title, kaspi_offer_name, link, price,
            min_price, max_price, sku_key, sku_id, my_size, active, is_available,
            dumping, is_on_sale_flag, identity_method, identity_value, canonical_item_key, raw_source
        )
        WITH ranked AS (
            SELECT
                h.*,
                ROW_NUMBER() OVER (
                    PARTITION BY COALESCE(h.canonical_item_key, '__row__' || h.id), COALESCE(h.store_code, '')
                    ORDER BY
                        h.source_priority ASC,
                        COALESCE(h.source_updated_at, '') DESC,
                        COALESCE(h.fetched_at_almaty, '') DESC,
                        h.id DESC
                ) AS rn
            FROM item_store_truth_history h
        )
        SELECT
            id, source_name, source_priority, source_snapshot, source_updated_at,
            fetched_at_almaty, fetched_at_utc, store_code, store_id, store_name, row_id,
            merchant_sku, kaspi_sku, merchant_title, kaspi_offer_name, link, price,
            min_price, max_price, sku_key, sku_id, my_size, active, is_available,
            dumping, is_on_sale_flag, identity_method, identity_value, canonical_item_key, raw_source
        FROM ranked
        WHERE rn = 1
        """
    )

    conn.execute("DELETE FROM item_identity_bridge")
    conn.execute(
        """
        INSERT INTO item_identity_bridge (
            canonical_item_key, identity_method, identity_value, first_seen_at, last_seen_at,
            rows_count, stores_count, best_source_priority, freshness_rank_source
        )
        SELECT
            h.canonical_item_key,
            MAX(h.identity_method),
            MAX(h.identity_value),
            MIN(COALESCE(h.source_updated_at, h.fetched_at_almaty)),
            MAX(COALESCE(h.source_updated_at, h.fetched_at_almaty)),
            COUNT(*) AS rows_count,
            COUNT(DISTINCT h.store_code) AS stores_count,
            MIN(h.source_priority) AS best_source_priority,
            (
                SELECT h2.source_name
                FROM item_store_truth_history h2
                WHERE h2.canonical_item_key = h.canonical_item_key
                ORDER BY h2.source_priority ASC, COALESCE(h2.source_updated_at, '') DESC, h2.id DESC
                LIMIT 1
            ) AS freshness_rank_source
        FROM item_store_truth_history h
        WHERE h.canonical_item_key IS NOT NULL
        GROUP BY h.canonical_item_key
        """
    )

    return inserted


def _table_to_markdown(headers: list[str], rows: list[tuple[Any, ...]]) -> str:
    if not rows:
        rows = []
    header_line = "| " + " | ".join(headers) + " |"
    divider_line = "| " + " | ".join("---" for _ in headers) + " |"
    body_lines: list[str] = []
    for row in rows:
        body_lines.append("| " + " | ".join("" if v is None else str(v) for v in row) + " |")
    return "\n".join([header_line, divider_line, *body_lines])


def _write_rows_sheet(
    *,
    workbook: Workbook,
    sheet_name: str,
    headers: list[str],
    rows: list[tuple[Any, ...]],
) -> None:
    ws = workbook.create_sheet(title=sheet_name)
    ws.append(headers)
    for row in rows:
        ws.append(list(row))


def _collect_unified_report(conn: sqlite3.Connection, db_path: Path) -> dict[str, Any]:
    history_rows = conn.execute("SELECT COUNT(*) AS c FROM item_store_truth_history").fetchone()["c"]
    latest_rows = conn.execute("SELECT COUNT(*) AS c FROM item_store_truth_latest").fetchone()["c"]
    has_latest_enriched = _table_exists(conn, "item_store_truth_latest_enriched")
    has_control_scope = _table_exists(conn, "control_scope_sold_90d")
    has_link_backfill_map = _table_exists(conn, "link_backfill_map")
    latest_enriched_rows = (
        conn.execute("SELECT COUNT(*) AS c FROM item_store_truth_latest_enriched").fetchone()["c"]
        if has_latest_enriched
        else 0
    )
    bridge_rows = conn.execute("SELECT COUNT(*) AS c FROM item_identity_bridge").fetchone()["c"]
    unresolved_latest_rows = conn.execute(
        "SELECT COUNT(*) AS c FROM item_store_truth_latest WHERE canonical_item_key IS NULL"
    ).fetchone()["c"]

    store_coverage_rows = conn.execute(
        """
        SELECT
            COALESCE(store_code, 'UNKNOWN') AS store_code,
            MIN(store_id) AS store_id,
            COUNT(*) AS latest_rows,
            SUM(CASE WHEN COALESCE(is_on_sale_flag, 0) = 1 THEN 1 ELSE 0 END) AS on_sale_rows,
            SUM(CASE WHEN COALESCE(is_on_sale_flag, 0) = 0 THEN 1 ELSE 0 END) AS off_sale_rows,
            SUM(CASE WHEN canonical_item_key IS NULL THEN 1 ELSE 0 END) AS unresolved_identity_rows
        FROM item_store_truth_latest
        GROUP BY COALESCE(store_code, 'UNKNOWN')
        ORDER BY COALESCE(store_code, 'UNKNOWN')
        """
    ).fetchall()

    identity_method_rows = conn.execute(
        """
        SELECT
            COALESCE(identity_method, 'UNRESOLVED') AS identity_method,
            COUNT(*) AS rows_count
        FROM item_store_truth_latest
        GROUP BY COALESCE(identity_method, 'UNRESOLVED')
        ORDER BY rows_count DESC
        """
    ).fetchall()

    source_rows = conn.execute(
        """
        SELECT source_name, COUNT(*) AS rows_count
        FROM item_store_truth_latest
        GROUP BY source_name
        ORDER BY rows_count DESC
        """
    ).fetchall()

    if has_control_scope:
        sold_window_rows = conn.execute(
            """
            SELECT
                COUNT(*) AS total_rows,
                SUM(CASE WHEN control_ready_flag = 1 THEN 1 ELSE 0 END) AS control_ready_rows,
                SUM(CASE WHEN COALESCE(link, '') = '' THEN 1 ELSE 0 END) AS missing_link_before,
                SUM(CASE WHEN COALESCE(effective_link, '') = '' THEN 1 ELSE 0 END) AS missing_link_after,
                SUM(CASE WHEN COALESCE(is_link_backfilled, 0) = 1 THEN 1 ELSE 0 END) AS backfilled_rows
            FROM control_scope_sold_90d
            """
        ).fetchone()
        sold_window_by_store_rows = conn.execute(
            """
            SELECT
                COALESCE(store_code, 'UNKNOWN') AS store_code,
                COUNT(*) AS total_rows,
                SUM(CASE WHEN control_ready_flag = 1 THEN 1 ELSE 0 END) AS control_ready_rows,
                SUM(CASE WHEN COALESCE(link, '') = '' THEN 1 ELSE 0 END) AS missing_link_before,
                SUM(CASE WHEN COALESCE(effective_link, '') = '' THEN 1 ELSE 0 END) AS missing_link_after,
                SUM(CASE WHEN COALESCE(is_link_backfilled, 0) = 1 THEN 1 ELSE 0 END) AS backfilled_rows
            FROM control_scope_sold_90d
            GROUP BY COALESCE(store_code, 'UNKNOWN')
            ORDER BY COALESCE(store_code, 'UNKNOWN')
            """
        ).fetchall()
        backfill_source_rows = conn.execute(
            """
            SELECT
                COALESCE(effective_link_source, 'none') AS effective_link_source,
                COUNT(*) AS rows_count
            FROM control_scope_sold_90d
            WHERE COALESCE(is_link_backfilled, 0) = 1
            GROUP BY COALESCE(effective_link_source, 'none')
            ORDER BY rows_count DESC
            """
        ).fetchall()
        residual_rows = conn.execute(
            """
            SELECT
                store_code,
                sku_id,
                sku_key,
                merchant_sku,
                kaspi_offer_name,
                sold_units_90d,
                sold_orders_90d
            FROM control_scope_sold_90d
            WHERE COALESCE(effective_link, '') = ''
            ORDER BY sold_units_90d DESC, sold_orders_90d DESC, store_code
            LIMIT 50
            """
        ).fetchall()
    else:
        sold_window_rows = {
            "total_rows": 0,
            "control_ready_rows": 0,
            "missing_link_before": 0,
            "missing_link_after": 0,
            "backfilled_rows": 0,
        }
        sold_window_by_store_rows = []
        backfill_source_rows = []
        residual_rows = []

    if has_link_backfill_map:
        link_backfill_map_rows = conn.execute(
            """
            SELECT
                COALESCE(resolved_from_source, 'unknown') AS resolved_from_source,
                COUNT(*) AS rows_count
            FROM link_backfill_map
            GROUP BY COALESCE(resolved_from_source, 'unknown')
            ORDER BY rows_count DESC
            """
        ).fetchall()
    else:
        link_backfill_map_rows = []

    return {
        "db_path": str(db_path),
        "history_rows": history_rows,
        "latest_rows": latest_rows,
        "latest_enriched_rows": latest_enriched_rows,
        "bridge_rows": bridge_rows,
        "unresolved_latest_rows": unresolved_latest_rows,
        "store_coverage": [dict(row) for row in store_coverage_rows],
        "identity_methods": [dict(row) for row in identity_method_rows],
        "latest_source_usage": [dict(row) for row in source_rows],
        "sold_90d_total_rows": int(sold_window_rows["total_rows"] or 0),
        "sold_90d_control_ready_rows": int(sold_window_rows["control_ready_rows"] or 0),
        "sold_90d_missing_link_before": int(sold_window_rows["missing_link_before"] or 0),
        "sold_90d_missing_link_after": int(sold_window_rows["missing_link_after"] or 0),
        "sold_90d_backfilled_rows": int(sold_window_rows["backfilled_rows"] or 0),
        "sold_90d_by_store": [dict(row) for row in sold_window_by_store_rows],
        "sold_90d_backfill_sources": [dict(row) for row in backfill_source_rows],
        "link_backfill_map_sources": [dict(row) for row in link_backfill_map_rows],
        "sold_90d_residual_sample": [dict(row) for row in residual_rows],
    }


def export_repricer_unified_report(*, db_path: str | Path) -> dict[str, Any]:
    path = Path(db_path)
    if not path.exists():
        raise RuntimeError(f"Unified truth DB not found: {path}")
    conn = sqlite3.connect(path)
    conn.row_factory = sqlite3.Row
    try:
        return _collect_unified_report(conn, path)
    finally:
        conn.close()


def export_repricer_unified_truth(
    *,
    config_path: str | Path,
    db_out_path: str | Path,
    xlsx_out_path: str | Path,
    markdown_out_path: str | Path,
    external_truth_dir: str | Path,
    kaspi_accounts_path: str | Path,
    legacy_snapshot_path: str | Path,
    links_base_xlsx_path: str | Path,
    scrape_catalog_xlsx_path: str | Path,
    scrape_pricewars_dir: str | Path,
    sales_window_days: int,
    refresh_repricer: bool,
    headless: bool,
) -> dict[str, Any]:
    db_path = Path(db_out_path)
    xlsx_path = Path(xlsx_out_path)
    markdown_path = Path(markdown_out_path)
    db_path.parent.mkdir(parents=True, exist_ok=True)
    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    markdown_path.parent.mkdir(parents=True, exist_ok=True)

    run_almaty, run_utc = _now_pair()
    store_id_to_code, store_code_to_id, store_code_to_name = _load_store_maps(
        kaspi_accounts_path=kaspi_accounts_path,
        repricer_config_path=config_path,
    )

    refresh_snapshot_path = db_path.parent / "repricer_items_refresh.sqlite"
    refresh_summary: dict[str, Any] | None = None
    if refresh_repricer:
        refresh_summary = export_repricer_items_to_sqlite(
            config_path=config_path,
            output_path=refresh_snapshot_path,
            headless=headless,
            include_all_rows=True,
        )

    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    try:
        _create_unified_tables(conn)
        _reset_unified_tables(conn)

        external_tables = _import_external_csvs(conn, external_truth_dir)
        links_base_import = _import_links_sku_base(conn, links_base_xlsx_path)
        links_catalog_import = _import_links_sku_catalog(conn, scrape_catalog_xlsx_path)
        links_pricewars_import = _import_links_price_wars(conn, scrape_pricewars_dir)
        link_backfill_build = _build_link_backfill_map(conn, run_almaty)

        legacy_rows = _ingest_repricer_snapshot(
            conn=conn,
            snapshot_path=Path(legacy_snapshot_path),
            source_snapshot="legacy_repricer_snapshot",
            source_priority=2,
            store_id_to_code=store_id_to_code,
            store_code_to_name=store_code_to_name,
            fallback_almaty=run_almaty,
            fallback_utc=run_utc,
        )
        refresh_rows = 0
        if refresh_repricer:
            refresh_rows = _ingest_repricer_snapshot(
                conn=conn,
                snapshot_path=refresh_snapshot_path,
                source_snapshot="refresh_repricer_snapshot",
                source_priority=1,
                store_id_to_code=store_id_to_code,
                store_code_to_name=store_code_to_name,
                fallback_almaty=run_almaty,
                fallback_utc=run_utc,
            )

        history_counters = _build_history_and_latest(
            conn=conn,
            external_tables=external_tables,
            store_id_to_code=store_id_to_code,
            store_code_to_id=store_code_to_id,
            store_code_to_name=store_code_to_name,
            default_almaty=run_almaty,
            default_utc=run_utc,
        )
        latest_enriched_build = _build_latest_enriched(conn)
        table_map = {name: data["table_name"] for name, data in external_tables.items()}
        sales_7m_table = _find_table_name(table_map, "sales_last_7_months_db_")
        control_scope_build = _build_control_scope_sold_window(
            conn,
            sales_7m_table,
            window_days=int(sales_window_days),
        )

        report = _collect_unified_report(conn, db_path)

        latest_headers = [
            "fetched_at",
            "store_id",
            "store_name",
            "row_id",
            "merchant_sku",
            "kaspi_sku",
            "merchant_title",
            "link",
            "price",
            "min_price",
            "max_price",
            "link_count_store-c",
            "link_count_acmewear",
            "store_code",
            "sku_key",
            "sku_id",
            "my_size",
            "kaspi_offer_name",
            "is_on_sale_flag",
            "active",
            "is_available",
            "dumping",
            "identity_method",
            "canonical_item_key",
            "source_name",
            "source_priority",
            "source_updated_at",
        ]
        latest_rows = conn.execute(
            """
            WITH link_counts AS (
                SELECT
                    effective_link AS link,
                    SUM(CASE WHEN store_code = 'MELVIS' THEN 1 ELSE 0 END) AS link_count_store-c,
                    SUM(CASE WHEN store_code = 'ACMEWEAR' THEN 1 ELSE 0 END) AS link_count_acmewear
                FROM item_store_truth_latest_enriched
                WHERE COALESCE(effective_link, '') <> ''
                GROUP BY effective_link
            )
            SELECT
                l.fetched_at_almaty AS fetched_at,
                l.store_id,
                l.store_name,
                l.row_id,
                l.merchant_sku,
                l.kaspi_sku,
                l.merchant_title,
                l.effective_link AS link,
                l.price,
                l.min_price,
                l.max_price,
                COALESCE(c.link_count_store-c, 0) AS link_count_store-c,
                COALESCE(c.link_count_acmewear, 0) AS link_count_acmewear,
                l.store_code,
                l.sku_key,
                l.sku_id,
                l.my_size,
                l.kaspi_offer_name,
                l.is_on_sale_flag,
                l.active,
                l.is_available,
                l.dumping,
                l.identity_method,
                l.canonical_item_key,
                l.source_name,
                l.source_priority,
                l.source_updated_at
            FROM item_store_truth_latest_enriched l
            LEFT JOIN link_counts c ON c.link = l.effective_link
            ORDER BY COALESCE(l.store_code, ''), COALESCE(l.canonical_item_key, ''), l.id
            """
        ).fetchall()

        history_headers = [
            "id",
            "source_name",
            "source_priority",
            "source_snapshot",
            "source_updated_at",
            "fetched_at_almaty",
            "store_code",
            "store_id",
            "store_name",
            "merchant_sku",
            "kaspi_sku",
            "merchant_title",
            "kaspi_offer_name",
            "link",
            "price",
            "min_price",
            "max_price",
            "sku_key",
            "sku_id",
            "my_size",
            "identity_method",
            "canonical_item_key",
        ]
        history_rows = conn.execute(
            """
            SELECT
                id,
                source_name,
                source_priority,
                source_snapshot,
                source_updated_at,
                fetched_at_almaty,
                store_code,
                store_id,
                store_name,
                merchant_sku,
                kaspi_sku,
                merchant_title,
                kaspi_offer_name,
                link,
                price,
                min_price,
                max_price,
                sku_key,
                sku_id,
                my_size,
                identity_method,
                canonical_item_key
            FROM item_store_truth_history
            ORDER BY id
            """
        ).fetchall()

        bridge_headers = [
            "canonical_item_key",
            "identity_method",
            "identity_value",
            "first_seen_at",
            "last_seen_at",
            "rows_count",
            "stores_count",
            "best_source_priority",
            "freshness_rank_source",
        ]
        bridge_rows = conn.execute(
            """
            SELECT
                canonical_item_key,
                identity_method,
                identity_value,
                first_seen_at,
                last_seen_at,
                rows_count,
                stores_count,
                best_source_priority,
                freshness_rank_source
            FROM item_identity_bridge
            ORDER BY canonical_item_key
            """
        ).fetchall()

        coverage_headers = [
            "store_code",
            "store_id",
            "latest_rows",
            "on_sale_rows",
            "off_sale_rows",
            "unresolved_identity_rows",
        ]
        coverage_rows = conn.execute(
            """
            SELECT
                COALESCE(store_code, 'UNKNOWN') AS store_code,
                MIN(store_id) AS store_id,
                COUNT(*) AS latest_rows,
                SUM(CASE WHEN COALESCE(is_on_sale_flag, 0) = 1 THEN 1 ELSE 0 END) AS on_sale_rows,
                SUM(CASE WHEN COALESCE(is_on_sale_flag, 0) = 0 THEN 1 ELSE 0 END) AS off_sale_rows,
                SUM(CASE WHEN canonical_item_key IS NULL THEN 1 ELSE 0 END) AS unresolved_identity_rows
            FROM item_store_truth_latest
            GROUP BY COALESCE(store_code, 'UNKNOWN')
            ORDER BY COALESCE(store_code, 'UNKNOWN')
            """
        ).fetchall()

        identity_headers = ["identity_method", "rows_count"]
        identity_rows = conn.execute(
            """
            SELECT
                COALESCE(identity_method, 'UNRESOLVED') AS identity_method,
                COUNT(*) AS rows_count
            FROM item_store_truth_latest
            GROUP BY COALESCE(identity_method, 'UNRESOLVED')
            ORDER BY rows_count DESC
            """
        ).fetchall()

        source_usage_headers = ["source_name", "rows_count"]
        source_usage_rows = conn.execute(
            """
            SELECT source_name, COUNT(*) AS rows_count
            FROM item_store_truth_latest_enriched
            GROUP BY source_name
            ORDER BY rows_count DESC
            """
        ).fetchall()
        sold_window_headers = [
            "store_code",
            "store_id",
            "sku_id",
            "sku_key",
            "last_sale_date",
            "sold_units_90d",
            "sold_orders_90d",
            "merchant_sku",
            "kaspi_offer_name",
            "link",
            "effective_link",
            "effective_link_source",
            "is_link_backfilled",
            "control_ready_flag",
        ]
        sold_window_rows = conn.execute(
            """
            SELECT
                store_code,
                store_id,
                sku_id,
                sku_key,
                last_sale_date,
                sold_units_90d,
                sold_orders_90d,
                merchant_sku,
                kaspi_offer_name,
                link,
                effective_link,
                effective_link_source,
                is_link_backfilled,
                control_ready_flag
            FROM control_scope_sold_90d
            ORDER BY store_code, last_sale_date DESC, sku_id, sku_key
            """
        ).fetchall()
        sold_window_by_store_headers = [
            "store_code",
            "total_rows",
            "control_ready_rows",
            "missing_link_before",
            "missing_link_after",
            "backfilled_rows",
        ]
        sold_window_by_store_rows = conn.execute(
            """
            SELECT
                COALESCE(store_code, 'UNKNOWN') AS store_code,
                COUNT(*) AS total_rows,
                SUM(CASE WHEN control_ready_flag = 1 THEN 1 ELSE 0 END) AS control_ready_rows,
                SUM(CASE WHEN COALESCE(link, '') = '' THEN 1 ELSE 0 END) AS missing_link_before,
                SUM(CASE WHEN COALESCE(effective_link, '') = '' THEN 1 ELSE 0 END) AS missing_link_after,
                SUM(CASE WHEN COALESCE(is_link_backfilled, 0) = 1 THEN 1 ELSE 0 END) AS backfilled_rows
            FROM control_scope_sold_90d
            GROUP BY COALESCE(store_code, 'UNKNOWN')
            ORDER BY COALESCE(store_code, 'UNKNOWN')
            """
        ).fetchall()
        backfill_map_headers = [
            "match_key_type",
            "match_key_value",
            "resolved_link",
            "resolved_from_source",
            "source_priority",
            "created_at_almaty",
        ]
        backfill_map_rows = conn.execute(
            """
            SELECT
                match_key_type,
                match_key_value,
                resolved_link,
                resolved_from_source,
                source_priority,
                created_at_almaty
            FROM link_backfill_map
            ORDER BY source_priority, match_key_type, match_key_value
            """
        ).fetchall()
        sold_residual_headers = [
            "store_code",
            "sku_id",
            "sku_key",
            "merchant_sku",
            "kaspi_offer_name",
            "sold_units_90d",
            "sold_orders_90d",
        ]
        sold_residual_rows = conn.execute(
            """
            SELECT
                store_code,
                sku_id,
                sku_key,
                merchant_sku,
                kaspi_offer_name,
                sold_units_90d,
                sold_orders_90d
            FROM control_scope_sold_90d
            WHERE COALESCE(effective_link, '') = ''
            ORDER BY sold_units_90d DESC, sold_orders_90d DESC, store_code
            """
        ).fetchall()

        covered_store_codes = {str(row["store_code"]) for row in coverage_rows}
        quality_rows = [
            (
                "history_rows_nonzero",
                "PASS" if report["history_rows"] > 0 else "FAIL",
                f"history_rows={report['history_rows']}",
            ),
            (
                "latest_rows_nonzero",
                "PASS" if report["latest_rows"] > 0 else "FAIL",
                f"latest_rows={report['latest_rows']}",
            ),
            (
                "bridge_rows_nonzero",
                "PASS" if report["bridge_rows"] > 0 else "FAIL",
                f"bridge_rows={report['bridge_rows']}",
            ),
            (
                "expected_store_coverage",
                "PASS" if EXPECTED_STORE_CODES.issubset(covered_store_codes) else "WARN",
                f"covered={','.join(sorted(covered_store_codes))}",
            ),
            (
                "identity_chain_present",
                "PASS"
                if {"link", "merchant_sku", "kaspi_sku"} & {str(r['identity_method']) for r in identity_rows}
                else "WARN",
                "methods=" + ",".join(str(r["identity_method"]) for r in identity_rows),
            ),
            (
                "unresolved_rows_visible",
                "PASS",
                f"unresolved_latest_rows={report['unresolved_latest_rows']}",
            ),
            (
                "sold_90d_scope_nonzero",
                "PASS" if report["sold_90d_total_rows"] > 0 else "WARN",
                f"sold_90d_total_rows={report['sold_90d_total_rows']}",
            ),
            (
                "sold_90d_ready_coverage",
                "PASS" if report["sold_90d_missing_link_after"] == 0 else "WARN",
                f"missing_after={report['sold_90d_missing_link_after']}",
            ),
        ]

        workbook = Workbook()
        workbook.remove(workbook.active)
        _write_rows_sheet(
            workbook=workbook,
            sheet_name="items_latest_all_stores",
            headers=latest_headers,
            rows=[tuple(r) for r in latest_rows],
        )
        _write_rows_sheet(
            workbook=workbook,
            sheet_name="items_history",
            headers=history_headers,
            rows=[tuple(r) for r in history_rows],
        )
        _write_rows_sheet(
            workbook=workbook,
            sheet_name="identity_bridge",
            headers=bridge_headers,
            rows=[tuple(r) for r in bridge_rows],
        )
        _write_rows_sheet(
            workbook=workbook,
            sheet_name="store_coverage_summary",
            headers=coverage_headers,
            rows=[tuple(r) for r in coverage_rows],
        )
        _write_rows_sheet(
            workbook=workbook,
            sheet_name="quality_checks",
            headers=["check_name", "status", "details"],
            rows=quality_rows,
        )
        _write_rows_sheet(
            workbook=workbook,
            sheet_name="links_backfill_map",
            headers=backfill_map_headers,
            rows=[tuple(r) for r in backfill_map_rows],
        )
        _write_rows_sheet(
            workbook=workbook,
            sheet_name="control_scope_sold_90d",
            headers=sold_window_headers,
            rows=[tuple(r) for r in sold_window_rows],
        )
        _write_rows_sheet(
            workbook=workbook,
            sheet_name="control_scope_residual_gaps",
            headers=sold_residual_headers,
            rows=[tuple(r) for r in sold_residual_rows],
        )
        _save_workbook_atomic(workbook, xlsx_path)

        markdown_parts = [
            "# Repricer Unified Truth Export",
            "",
            f"- Generated at (Asia/Almaty): `{run_almaty}`",
            f"- Generated at (UTC): `{run_utc}`",
            f"- DB: `{db_path}`",
            f"- Workbook: `{xlsx_path}`",
            "",
            "## Store Coverage",
            _table_to_markdown(coverage_headers, [tuple(r) for r in coverage_rows]),
            "",
            "## Identity Method Coverage",
            _table_to_markdown(identity_headers, [tuple(r) for r in identity_rows]),
            "",
            "## Latest Source Usage",
            _table_to_markdown(source_usage_headers, [tuple(r) for r in source_usage_rows]),
            "",
            "## Sold Last 90 Days Coverage",
            _table_to_markdown(sold_window_by_store_headers, [tuple(r) for r in sold_window_by_store_rows]),
            "",
            "## Backfill Source Usage (Sold Last 90 Days)",
            _table_to_markdown(
                ["effective_link_source", "rows_count"],
                [tuple(r.values()) for r in report["sold_90d_backfill_sources"]],
            ),
            "",
            "## Residual Gaps (Sold Last 90 Days)",
            _table_to_markdown(sold_residual_headers, [tuple(r) for r in sold_residual_rows[:100]]),
            "",
            "## Quality Checks",
            _table_to_markdown(["check_name", "status", "details"], quality_rows),
        ]
        _write_text_atomic(markdown_path, "\n".join(markdown_parts) + "\n")

        conn.execute(
            "INSERT OR REPLACE INTO export_meta(key, value) VALUES(?, ?)",
            ("generated_at_almaty", run_almaty),
        )
        conn.execute(
            "INSERT OR REPLACE INTO export_meta(key, value) VALUES(?, ?)",
            ("generated_at_utc", run_utc),
        )
        conn.execute(
            "INSERT OR REPLACE INTO export_meta(key, value) VALUES(?, ?)",
            ("external_truth_dir", str(Path(external_truth_dir))),
        )
        conn.execute(
            "INSERT OR REPLACE INTO export_meta(key, value) VALUES(?, ?)",
            ("sales_window_days", str(int(sales_window_days))),
        )
        conn.execute(
            "INSERT OR REPLACE INTO export_meta(key, value) VALUES(?, ?)",
            ("links_base_xlsx", str(Path(links_base_xlsx_path))),
        )
        conn.execute(
            "INSERT OR REPLACE INTO export_meta(key, value) VALUES(?, ?)",
            ("scrape_catalog_xlsx", str(Path(scrape_catalog_xlsx_path))),
        )
        conn.execute(
            "INSERT OR REPLACE INTO export_meta(key, value) VALUES(?, ?)",
            ("scrape_pricewars_dir", str(Path(scrape_pricewars_dir))),
        )
        conn.commit()

        return {
            **report,
            **history_counters,
            "db_path": str(db_path),
            "xlsx_path": str(xlsx_path),
            "markdown_path": str(markdown_path),
            "generated_at_almaty": run_almaty,
            "generated_at_utc": run_utc,
            "imported_external_tables": external_tables,
            "links_base_import": links_base_import,
            "links_catalog_import": links_catalog_import,
            "links_pricewars_import": links_pricewars_import,
            "link_backfill_build": link_backfill_build,
            "latest_enriched_build": latest_enriched_build,
            "control_scope_build": control_scope_build,
            "legacy_snapshot_rows": legacy_rows,
            "refresh_snapshot_rows": refresh_rows,
            "refresh_summary": refresh_summary,
        }
    finally:
        conn.close()
