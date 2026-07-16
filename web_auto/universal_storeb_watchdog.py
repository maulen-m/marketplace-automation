from __future__ import annotations

import argparse
import csv
import json
import os
import sqlite3
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

import yaml
from openpyxl import load_workbook

from .kaspi_forbidden_cards import BERSERK_OWNER_DECISION_ID, forbidden_card_match_for_row
from .kaspi_merchant_common import load_env_assignments, normalize_store_name, resolve_store_credentials
from .kaspi_pricelist_ops import WAREHOUSE_COLUMNS, is_no_like, parse_int


SCHEMA_VERSION = "web_auto.universal_storeb_watchdog.v1"
DEFAULT_CONFIG_PATH = Path("config/tasks/universal_storeb_reactivation_watchdog.yaml")
DEFAULT_RUN_ROOT = Path("runs/universal_storeb_reactivation_watchdog")
ALLOWED_STORES = {"UNIVERSAL": 30000001, "STORE-B": 30000002}
FORBIDDEN_TOKENS = {
    "ACMEWEAR",
    "LINE51",
    "LINE61",
    "SUIT-21",
    "SUIT-31",
    "LINE-21",
    "LINE-31",
    "LS21",
    "LS31",
    "TS21",
    "TS31",
    "TRM",
}


class UniversalMGroupWatchdogError(ValueError):
    pass


def _now_almaty() -> datetime:
    return datetime.now(ZoneInfo("Asia/Almaty")).replace(microsecond=0)


def _repo_path(value: str | Path, repo_root: Path) -> Path:
    path = Path(str(value)).expanduser()
    if path.is_absolute():
        return path
    return repo_root / path


def _read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def _write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")


def _write_csv(path: Path, rows: list[dict[str, Any]], fieldnames: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            writer.writerow({field: row.get(field, "") for field in fieldnames})


def _read_csv_rows(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def _path_list(value: Any) -> list[str]:
    if value is None:
        return []
    if isinstance(value, (str, Path)):
        return [str(value)]
    if isinstance(value, list):
        return [str(item) for item in value if item is not None and str(item).strip()]
    raise UniversalMGroupWatchdogError("path list config values must be strings or lists")


def _walk_json_paths(value: Any) -> list[str]:
    paths: list[str] = []
    if isinstance(value, dict):
        for child in value.values():
            paths.extend(_walk_json_paths(child))
    elif isinstance(value, list):
        for child in value:
            paths.extend(_walk_json_paths(child))
    elif isinstance(value, str) and value.strip():
        paths.append(value.strip())
    return paths


def _repricer_sqlite_from_export_paths(path: Path, repo_root: Path) -> Path | None:
    if not path.exists():
        return None
    raw = _read_json(path)
    for candidate in _walk_json_paths(raw):
        if ".sqlite" not in candidate:
            continue
        resolved = _repo_path(candidate, repo_root)
        if resolved.exists():
            return resolved
    return None


def resolve_watchdog_sources(config: dict[str, Any], repo_root: Path) -> dict[str, Path | str | None]:
    ledgers = config.get("source_ledgers") or {}
    for raw_dir in _path_list(ledgers.get("agent1_output_dirs")):
        source_dir = _repo_path(raw_dir, repo_root)
        target = source_dir / "fresh_target_reconciliation.csv"
        pricelist_paths = source_dir / "fresh_pricelist_paths.json"
        if not target.exists() or not pricelist_paths.exists():
            continue
        export_paths = source_dir / "fresh_repricer_export_paths.json"
        return {
            "source": "agent1_output_dir",
            "source_dir": str(source_dir),
            "target_reconciliation_csv": target,
            "activation_ledger_csv": source_dir / "owner_decision_live_write_ledger.csv",
            "blocked_or_skipped_rows_csv": source_dir / "blocked_or_skipped_rows.csv",
            "fresh_pricelist_paths_json": pricelist_paths,
            "repricer_sqlite_path": _repricer_sqlite_from_export_paths(export_paths, repo_root),
            "fresh_repricer_export_paths_json": export_paths if export_paths.exists() else None,
        }

    target_value = ledgers.get("target_reconciliation_csv")
    pricelist_value = config.get("fresh_pricelist_paths_json")
    target = _repo_path(target_value, repo_root) if target_value else None
    pricelist_paths = _repo_path(pricelist_value, repo_root) if pricelist_value else None
    activation = ledgers.get("activation_ledger_csv")
    blocked = ledgers.get("blocked_or_skipped_rows_csv")
    repricer_sqlite = (config.get("repricer") or {}).get("sqlite_path")
    return {
        "source": "configured_paths",
        "source_dir": "",
        "target_reconciliation_csv": target,
        "activation_ledger_csv": _repo_path(activation, repo_root) if activation else None,
        "blocked_or_skipped_rows_csv": _repo_path(blocked, repo_root) if blocked else None,
        "fresh_pricelist_paths_json": pricelist_paths,
        "repricer_sqlite_path": _repo_path(repricer_sqlite, repo_root) if repricer_sqlite else None,
        "fresh_repricer_export_paths_json": None,
    }


def _configured_env_files(config: dict[str, Any], repo_root: Path) -> list[Path]:
    raw_files = config.get("env_files")
    if raw_files is None:
        raw_files = [config.get("env_file") or ".env"]
    elif isinstance(raw_files, (str, Path)):
        raw_files = [raw_files]
    elif not isinstance(raw_files, list):
        raise UniversalMGroupWatchdogError("env_files must be a string or list when provided")

    out: list[Path] = []
    seen: set[Path] = set()
    for raw in raw_files:
        if raw is None or str(raw).strip() == "":
            continue
        path = _repo_path(str(raw), repo_root)
        if path in seen:
            continue
        out.append(path)
        seen.add(path)
    return out


def _resolve_store_credentials_from_env_files(store_name: str, env_files: list[Path]) -> dict[str, str]:
    errors: list[str] = []
    for env_file in env_files:
        if not env_file.exists():
            errors.append(f"{env_file}: missing")
            continue
        try:
            return resolve_store_credentials(store_name, env_file)
        except Exception as exc:
            errors.append(f"{env_file}: {exc}")
    raise UniversalMGroupWatchdogError(
        f"missing merchant credentials for {normalize_store_name(store_name)} in configured env files: "
        + "; ".join(errors)
    )


def _load_required_env_keys_from_files(keys: list[str], env_files: list[Path]) -> list[str]:
    missing: list[str] = []
    for key in keys:
        if os.environ.get(key):
            continue
        loaded_value = ""
        for env_file in env_files:
            if not env_file.exists():
                continue
            assignments = load_env_assignments(env_file, [key])
            if assignments.get(key):
                loaded_value = assignments[key]
                break
        if loaded_value:
            os.environ[key] = loaded_value
        else:
            missing.append(key)
    return missing


def _truthy(value: Any) -> bool:
    return str(value or "").strip().lower() in {"1", "true", "yes", "y"}


def _stock_value(row: dict[str, Any]) -> int | None:
    for key in ("current_stock_estimate", "current_stock_snapshot", "stock", "approved_stock"):
        raw = row.get(key)
        if raw is None or str(raw).strip() == "":
            continue
        return parse_int(raw)
    return None


def _expected_pp1(row: dict[str, Any]) -> int | None:
    for key in ("chosen_platform_stock_value", "suggested_platform_PP1_stock", "source_PP1"):
        text = str(row.get(key) or "").strip()
        if not text:
            continue
        if "PP1=" in text:
            after = text.split("PP1=", 1)[1].split(";", 1)[0].split(",", 1)[0]
            value = parse_int(after)
        else:
            value = parse_int(text)
        if value > 0:
            return value
    return None


def _floor_value(row: dict[str, Any]) -> int | None:
    for key in (
        "floor_min_price",
        "target_floor_min_price",
        "price_floor_cogs_103pct",
        "cogs_floor_min_price",
        "min_floor_price",
    ):
        raw = row.get(key)
        if raw is None or str(raw).strip() == "":
            continue
        value = parse_int(raw)
        if value > 0:
            return value
    return None


def _normalize_store_code(value: Any) -> str:
    text = str(value or "").strip().upper().replace("-", "")
    if text == "STOREB":
        return "STOREB"
    if text == "UNIVERSAL":
        return "UNIVERSAL"
    return text


def _has_forbidden_token(row: dict[str, Any]) -> str:
    haystack = " ".join(
        str(row.get(key) or "")
        for key in ("store", "store_name", "store_code", "merchant_sku", "sku_key", "sku_id", "model", "mapped_model")
    ).upper()
    for token in sorted(FORBIDDEN_TOKENS):
        if token in haystack:
            return token
    return ""


def load_watchdog_config(path: str | Path = DEFAULT_CONFIG_PATH, *, repo_root: Path | None = None) -> dict[str, Any]:
    root = (repo_root or Path.cwd()).resolve()
    config_path = _repo_path(path, root)
    raw = yaml.safe_load(config_path.read_text(encoding="utf-8")) or {}
    if not isinstance(raw, dict):
        raise UniversalMGroupWatchdogError("watchdog config must be a mapping")
    if str(raw.get("schema_version") or "").strip() != SCHEMA_VERSION:
        raise UniversalMGroupWatchdogError("unsupported or missing schema_version")

    stores = raw.get("stores")
    if not isinstance(stores, list) or not stores:
        raise UniversalMGroupWatchdogError("stores must be a non-empty list")
    seen: set[str] = set()
    for store in stores:
        if not isinstance(store, dict):
            raise UniversalMGroupWatchdogError("each store must be a mapping")
        store_name = normalize_store_name(store.get("store_name"))
        store_id = int(store.get("store_id") or 0)
        if store_name not in ALLOWED_STORES:
            raise UniversalMGroupWatchdogError(f"unsupported store in watchdog config: {store_name}")
        if ALLOWED_STORES[store_name] != store_id:
            raise UniversalMGroupWatchdogError(f"store_id mismatch for {store_name}: {store_id}")
        if store_name in seen:
            raise UniversalMGroupWatchdogError(f"duplicate store: {store_name}")
        seen.add(store_name)
    if seen != set(ALLOWED_STORES):
        raise UniversalMGroupWatchdogError("watchdog config must include exactly UNIVERSAL and STORE-B")

    auto_correction = raw.get("auto_correction") or {}
    if bool(auto_correction.get("enabled", False)):
        raise UniversalMGroupWatchdogError("auto_correction is not implemented and must remain disabled")

    sources = resolve_watchdog_sources(raw, root)
    required_paths = [
        ("source_ledgers.target_reconciliation_csv", sources["target_reconciliation_csv"]),
        ("fresh_pricelist_paths_json", sources["fresh_pricelist_paths_json"]),
    ]
    for label, path in required_paths:
        if not path or not Path(path).exists():
            raise UniversalMGroupWatchdogError(f"configured path does not exist: {label}={path}")
    return raw


def _load_workbook_rows(path: Path, *, store_name: str, sale_state: str) -> dict[str, dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Лист1" not in workbook.sheetnames:
            raise UniversalMGroupWatchdogError(f"missing required sheet Лист1 in {path}")
        sheet = workbook["Лист1"]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            raise UniversalMGroupWatchdogError(f"empty workbook sheet: {path}:Лист1")
        columns = [str(cell or "").strip() for cell in header]
        index = {name: idx for idx, name in enumerate(columns)}
        required = ["SKU", "price", *WAREHOUSE_COLUMNS, "preorder"]
        missing = [column for column in required if column not in index]
        if missing:
            raise UniversalMGroupWatchdogError(f"missing column(s) in {path}: {', '.join(missing)}")

        out: dict[str, dict[str, Any]] = {}
        for row_number, values in enumerate(rows, start=2):
            merchant_sku = str(values[index["SKU"]] or "").strip()
            if not merchant_sku:
                continue
            warehouse = {col: values[index[col]] for col in WAREHOUSE_COLUMNS}
            positive_columns = [col for col, value in warehouse.items() if parse_int(value) > 0]
            out[merchant_sku] = {
                "store_name": store_name,
                "source_state": sale_state,
                "row_number": row_number,
                "merchant_sku": merchant_sku,
                "price": parse_int(values[index["price"]]),
                "PP1": values[index["PP1"]],
                "PP2": values[index["PP2"]],
                "PP3": values[index["PP3"]],
                "PP4": values[index["PP4"]],
                "PP5": values[index["PP5"]],
                "preorder": values[index["preorder"]],
                "positive_stock": bool(positive_columns),
                "positive_stock_columns": ",".join(positive_columns),
                "active_stock_total": sum(parse_int(warehouse[col]) for col in positive_columns),
            }
        return out
    finally:
        workbook.close()


def _load_pricelist_paths(config: dict[str, Any], repo_root: Path, sources: dict[str, Path | str | None]) -> dict[str, dict[str, Path]]:
    raw = _read_json(Path(sources["fresh_pricelist_paths_json"]))
    stores = raw.get("stores") or {}
    out: dict[str, dict[str, Path]] = {}
    for store_name in ALLOWED_STORES:
        store_paths = stores.get(store_name) or {}
        active = store_paths.get("ACTIVE")
        archive = store_paths.get("ARCHIVE")
        if not active or not archive:
            raise UniversalMGroupWatchdogError(f"missing ACTIVE/ARCHIVE path for {store_name}")
        out[store_name] = {
            "ACTIVE": _repo_path(active, repo_root),
            "ARCHIVE": _repo_path(archive, repo_root),
        }
    return out


def _load_pricelist_state(paths: dict[str, dict[str, Path]]) -> dict[str, dict[str, dict[str, dict[str, Any]]]]:
    state: dict[str, dict[str, dict[str, dict[str, Any]]]] = {}
    for store_name, by_state in paths.items():
        state[store_name] = {
            "ACTIVE": _load_workbook_rows(by_state["ACTIVE"], store_name=store_name, sale_state="ACTIVE"),
            "ARCHIVE": _load_workbook_rows(by_state["ARCHIVE"], store_name=store_name, sale_state="ARCHIVE"),
        }
    return state


def _load_repricer_state(path: Path | None) -> dict[tuple[str, str], dict[str, Any]]:
    if path is None or not path.exists():
        return {}
    conn = sqlite3.connect(path)
    conn.row_factory = sqlite3.Row
    try:
        rows = conn.execute(
            """
            SELECT store_id, store_name, merchant_sku, price, min_price, max_price, dumping, active, is_available
            FROM repricer_items
            """
        ).fetchall()
    finally:
        conn.close()
    out: dict[tuple[str, str], dict[str, Any]] = {}
    for row in rows:
        store_name = normalize_store_name(row["store_name"])
        merchant_sku = str(row["merchant_sku"] or "").strip()
        if store_name in ALLOWED_STORES and merchant_sku:
            out[(store_name, merchant_sku)] = dict(row)
    return out


def _load_target_rows(
    config: dict[str, Any],
    repo_root: Path,
    sources: dict[str, Path | str | None],
    *,
    suppression_stats: dict[str, int] | None = None,
) -> list[dict[str, Any]]:
    target_path = Path(sources["target_reconciliation_csv"])
    rows = _read_csv_rows(target_path)
    activation_path = sources.get("activation_ledger_csv")
    activation_keys: set[tuple[str, str]] = set()
    if activation_path:
        path = Path(activation_path)
        if path.exists():
            for row in _read_csv_rows(path):
                activation_keys.add((normalize_store_name(row.get("store") or row.get("target_store")), str(row.get("merchant_sku") or "").strip()))

    out: list[dict[str, Any]] = []
    denied_family_suppressed = 0
    for row in rows:
        store = normalize_store_name(row.get("store") or row.get("target_store") or row.get("store_name"))
        merchant_sku = str(row.get("merchant_sku") or "").strip()
        if not store or not merchant_sku:
            continue
        stock = _stock_value(row)
        live_write_candidate = _truthy(row.get("live_write_candidate"))
        target_for_activation = _truthy(row.get("target_for_activation"))
        active_now = str(row.get("fresh_state_classification") or "").strip() == "active_now"
        row["store"] = store
        row["store_code"] = _normalize_store_code(row.get("store_code") or store)
        row["stock_value"] = stock
        row["expected_pp1"] = _expected_pp1(row)
        if stock is not None and stock > 0 and row["expected_pp1"] is None:
            row["expected_pp1"] = 500
        row["floor_min_price_value"] = _floor_value(row)
        row["is_activation_ledger_row"] = (store, merchant_sku) in activation_keys
        row["is_watch_target"] = live_write_candidate or target_for_activation or active_now or row["is_activation_ledger_row"]
        denied_match = forbidden_card_match_for_row(row)
        if denied_match is not None and denied_match.owner_decision == BERSERK_OWNER_DECISION_ID:
            denied_family_suppressed += 1
            continue
        out.append(row)
    if suppression_stats is not None:
        suppression_stats["denied_family_suppressed"] = denied_family_suppressed
    return out


def _active_has_expected_pp1(active_row: dict[str, Any], expected_pp1: int | None) -> bool:
    if expected_pp1 is None:
        return True
    return parse_int(active_row.get("PP1")) == expected_pp1 and all(is_no_like(active_row.get(col)) for col in ("PP2", "PP3", "PP4", "PP5"))


def _finding(
    *,
    severity: str,
    finding_type: str,
    row: dict[str, Any],
    message: str,
    kaspi_state: str = "",
    repricer_state: str = "",
) -> dict[str, Any]:
    return {
        "severity": severity,
        "finding_type": finding_type,
        "store": row.get("store", ""),
        "merchant_sku": row.get("merchant_sku", ""),
        "sku_key": row.get("sku_key", ""),
        "my_size_key": row.get("my_size_key", ""),
        "stock_value": "" if row.get("stock_value") is None else row.get("stock_value"),
        "kaspi_state": kaspi_state,
        "repricer_state": repricer_state,
        "message": message,
    }


def build_watchdog_report(
    *,
    config: dict[str, Any],
    repo_root: Path,
    pricelist_paths: dict[str, dict[str, Path]],
    repricer_sqlite_path: Path | None,
    run_dir: Path,
    generated_at: str,
) -> dict[str, Any]:
    pricelists = _load_pricelist_state(pricelist_paths)
    repricer = _load_repricer_state(repricer_sqlite_path)
    sources = resolve_watchdog_sources(config, repo_root)
    suppression_stats: dict[str, int] = {}
    target_rows = _load_target_rows(config, repo_root, sources, suppression_stats=suppression_stats)

    findings: list[dict[str, Any]] = []
    repricer_store_counts = Counter(store for store, _merchant_sku in repricer.keys())
    if repricer_sqlite_path is None:
        findings.append(
            _finding(
                severity="YELLOW",
                finding_type="repricer_import_missing_or_incomplete",
                row={},
                message="no Repricer sqlite path is available from config or Agent 1 export paths",
            )
        )
    elif not repricer_sqlite_path.exists():
        findings.append(
            _finding(
                severity="YELLOW",
                finding_type="repricer_import_missing_or_incomplete",
                row={},
                message=f"Repricer sqlite path does not exist: {repricer_sqlite_path}",
            )
        )
    elif not repricer:
        findings.append(
            _finding(
                severity="YELLOW",
                finding_type="repricer_import_missing_or_incomplete",
                row={},
                message=f"Repricer sqlite has no scoped Universal/STORE-B rows: {repricer_sqlite_path}",
            )
        )
    else:
        missing_repricer_stores = [store for store in ALLOWED_STORES if repricer_store_counts.get(store, 0) <= 0]
        if missing_repricer_stores:
            findings.append(
                _finding(
                    severity="YELLOW",
                    finding_type="repricer_import_missing_or_incomplete",
                    row={},
                    message="Repricer sqlite is missing scoped rows for: " + ", ".join(missing_repricer_stores),
                )
            )

    checked_rows = 0
    for row in target_rows:
        store = normalize_store_name(row.get("store"))
        merchant_sku = str(row.get("merchant_sku") or "").strip()
        if not row.get("is_watch_target"):
            continue
        checked_rows += 1
        forbidden = _has_forbidden_token(row)
        if store not in ALLOWED_STORES:
            findings.append(
                _finding(
                    severity="RED",
                    finding_type="forbidden_store_scope",
                    row=row,
                    message=f"row is outside Universal/STORE-B scope: {store}",
                )
            )
            continue
        if forbidden:
            findings.append(
                _finding(
                    severity="RED",
                    finding_type="forbidden_product_scope",
                    row=row,
                    message=f"row contains forbidden token: {forbidden}",
                )
            )
            continue

        active_row = pricelists[store]["ACTIVE"].get(merchant_sku)
        archive_row = pricelists[store]["ARCHIVE"].get(merchant_sku)
        if active_row:
            kaspi_state = "ACTIVE_POSITIVE" if active_row["positive_stock"] else "ACTIVE_NO_POSITIVE_STOCK"
        elif archive_row:
            kaspi_state = "ARCHIVE"
        else:
            kaspi_state = "MISSING"
        repricer_row = repricer.get((store, merchant_sku))
        if repricer_row:
            repricer_state = "active=%s available=%s dumping=%s" % (
                repricer_row.get("active"),
                repricer_row.get("is_available"),
                repricer_row.get("dumping"),
            )
        else:
            repricer_state = "MISSING"

        stock = row.get("stock_value")
        floor = row.get("floor_min_price_value")
        positive_expected = stock is not None and stock > 0
        zero_expected = stock is not None and stock <= 0
        if positive_expected:
            if not active_row:
                findings.append(
                    _finding(
                        severity="YELLOW",
                        finding_type="positive_stock_target_not_active_kaspi",
                        row=row,
                        kaspi_state=kaspi_state,
                        repricer_state=repricer_state,
                        message="positive-stock watch target is not ACTIVE in Kaspi pricelist state",
                    )
                )
            elif not active_row["positive_stock"]:
                findings.append(
                    _finding(
                        severity="YELLOW",
                        finding_type="positive_stock_target_active_without_platform_stock",
                        row=row,
                        kaspi_state=kaspi_state,
                        repricer_state=repricer_state,
                        message="positive-stock watch target is ACTIVE but has no positive PP stock carrier",
                    )
                )
            elif not _active_has_expected_pp1(active_row, row.get("expected_pp1")):
                findings.append(
                    _finding(
                        severity="YELLOW",
                        finding_type="platform_stock_policy_mismatch",
                        row=row,
                        kaspi_state=kaspi_state,
                        repricer_state=repricer_state,
                        message="ACTIVE Kaspi row does not match expected PP1 facade stock policy",
                    )
                )
            if repricer and not repricer_row:
                findings.append(
                    _finding(
                        severity="YELLOW",
                        finding_type="positive_stock_target_missing_repricer",
                        row=row,
                        kaspi_state=kaspi_state,
                        repricer_state=repricer_state,
                        message="positive-stock watch target is missing from scoped Repricer export",
                    )
                )
            elif repricer_row and (parse_int(repricer_row.get("active")) <= 0 or parse_int(repricer_row.get("is_available")) <= 0):
                findings.append(
                    _finding(
                        severity="YELLOW",
                        finding_type="positive_stock_target_repricer_off",
                        row=row,
                        kaspi_state=kaspi_state,
                        repricer_state=repricer_state,
                        message="positive-stock watch target is present in Repricer but not active/available",
                    )
                )
            if repricer_row and parse_int(repricer_row.get("active")) > 0 and parse_int(repricer_row.get("is_available")) > 0:
                if parse_int(repricer_row.get("dumping")) <= 0:
                    findings.append(
                        _finding(
                            severity="YELLOW",
                            finding_type="active_target_dumping_disabled",
                            row=row,
                            kaspi_state=kaspi_state,
                            repricer_state=repricer_state,
                            message="positive-stock active target is active/available in Repricer but dumping is disabled",
                        )
                    )
                if floor is not None:
                    current_price = parse_int(repricer_row.get("price"))
                    min_price = parse_int(repricer_row.get("min_price"))
                    if current_price > 0 and current_price < floor:
                        findings.append(
                            _finding(
                                severity="YELLOW",
                                finding_type="repricer_current_price_below_floor",
                                row=row,
                                kaspi_state=kaspi_state,
                                repricer_state=repricer_state,
                                message=f"Repricer current price {current_price} is below floor {floor}",
                            )
                        )
                    if min_price > 0 and min_price < floor:
                        findings.append(
                            _finding(
                                severity="YELLOW",
                                finding_type="repricer_min_price_below_floor",
                                row=row,
                                kaspi_state=kaspi_state,
                                repricer_state=repricer_state,
                                message=f"Repricer min price {min_price} is below floor {floor}",
                            )
                        )
            if active_row and floor is not None:
                active_price = parse_int(active_row.get("price"))
                if active_price > 0 and active_price < floor:
                    findings.append(
                        _finding(
                            severity="YELLOW",
                            finding_type="kaspi_active_price_below_floor",
                            row=row,
                            kaspi_state=kaspi_state,
                            repricer_state=repricer_state,
                            message=f"Kaspi ACTIVE price {active_price} is below floor {floor}",
                        )
                    )
        if zero_expected:
            if active_row and active_row["positive_stock"]:
                findings.append(
                    _finding(
                        severity="RED",
                        finding_type="zero_stock_target_active_kaspi",
                        row=row,
                        kaspi_state=kaspi_state,
                        repricer_state=repricer_state,
                        message="zero-stock/OOS watch target is still ACTIVE with positive Kaspi stock",
                    )
                )
            if repricer_row and parse_int(repricer_row.get("active")) > 0 and parse_int(repricer_row.get("is_available")) > 0:
                findings.append(
                    _finding(
                        severity="RED",
                        finding_type="zero_stock_target_active_repricer",
                        row=row,
                        kaspi_state=kaspi_state,
                        repricer_state=repricer_state,
                        message="zero-stock/OOS watch target is still active/available in Repricer",
                    )
                )

    severity_counts = Counter(item["severity"] for item in findings)
    type_counts = Counter(item["finding_type"] for item in findings)
    gate = "RED" if severity_counts.get("RED") else "YELLOW" if findings else "GREEN"
    summary = {
        "schema_version": SCHEMA_VERSION,
        "status": "failed" if gate == "RED" else "warning" if gate == "YELLOW" else "success",
        "gate": gate,
        "generated_at": generated_at,
        "run_dir": str(run_dir),
        "stores": list(ALLOWED_STORES.keys()),
        "production_write_action_executed": False,
        "auto_correction_enabled": False,
        "target_rows_loaded": len(target_rows),
        "denied_family_suppressed": suppression_stats.get("denied_family_suppressed", 0),
        "watch_rows_checked": checked_rows,
        "findings_count": len(findings),
        "severity_counts": dict(severity_counts),
        "finding_type_counts": dict(type_counts),
        "repricer_sqlite_path": str(repricer_sqlite_path or ""),
        "repricer_store_counts": dict(repricer_store_counts),
        "source_paths": {key: str(value or "") for key, value in sources.items()},
        "pricelist_paths": {
            store: {state: str(path) for state, path in by_state.items()} for store, by_state in pricelist_paths.items()
        },
    }

    fieldnames = [
        "severity",
        "finding_type",
        "store",
        "merchant_sku",
        "sku_key",
        "my_size_key",
        "stock_value",
        "kaspi_state",
        "repricer_state",
        "message",
    ]
    _write_csv(run_dir / "WATCHDOG_FINDINGS.csv", findings, fieldnames)
    _write_json(run_dir / "WATCHDOG_SUMMARY.json", summary)
    (run_dir / "WATCHDOG_REPORT.md").write_text(_render_report(summary, findings), encoding="utf-8")
    return summary


def _render_report(summary: dict[str, Any], findings: list[dict[str, Any]]) -> str:
    lines = [
        "# Universal/STORE-B Reactivation Watchdog Report",
        "",
        f"- generated_at: {summary['generated_at']}",
        f"- run_dir: {summary['run_dir']}",
        "- scope: UNIVERSAL and STORE-B only",
        "- production_write_action_executed: false",
        "- auto_correction_enabled: false",
        f"- target_rows_loaded: {summary['target_rows_loaded']}",
        f"- denied_family_suppressed: {summary['denied_family_suppressed']}",
        f"- watch_rows_checked: {summary['watch_rows_checked']}",
        f"- findings_count: {summary['findings_count']}",
        "",
    ]
    if summary.get("source_errors"):
        lines.extend(["## Source Errors", ""])
        lines.extend([f"- {item}" for item in summary.get("source_errors") or []])
        lines.append("")
    lines.extend(["## Findings", ""])
    if not findings:
        lines.append("No watchdog findings.")
    else:
        lines.extend(
            [
                "| severity | type | store | merchant_sku | stock | kaspi_state | repricer_state |",
                "| --- | --- | --- | --- | ---: | --- | --- |",
            ]
        )
        for item in findings:
            lines.append(
                "| {severity} | {finding_type} | {store} | {merchant_sku} | {stock_value} | {kaspi_state} | {repricer_state} |".format(
                    **item
                )
            )
    lines.extend(
        [
            "",
            "This watchdog is read-only/reporting. No price, stock, Kaspi, Repricer, scheduler, CRM, database, or Autonomous_business write was executed.",
            "",
            f"Gate: {summary['gate']}",
            "",
        ]
    )
    return "\n".join(lines)


def _write_scoped_repricer_config(config: dict[str, Any], repo_root: Path, run_dir: Path) -> Path:
    source_config = _repo_path((config.get("repricer") or {}).get("source_config", "config/tasks/repricer_competitors.yaml"), repo_root)
    raw = yaml.safe_load(source_config.read_text(encoding="utf-8")) or {}
    allowed_ids = set(ALLOWED_STORES.values())
    raw["accounts"] = [
        {**account, "stores": [int(store) for store in account.get("stores", []) if int(store) in allowed_ids]}
        for account in raw.get("accounts", [])
    ]
    raw["accounts"] = [account for account in raw["accounts"] if account.get("stores")]
    raw["store_name_map"] = {name: sid for name, sid in (raw.get("store_name_map") or {}).items() if int(sid) in allowed_ids}
    raw.setdefault("run", {})
    raw["run"]["checkpoint_path"] = str(run_dir / "repricer_checkpoint.json")
    raw["run"]["artifacts_dir"] = str(run_dir / "repricer_artifacts")
    out = run_dir / "repricer_scoped_universal_storeb.yaml"
    out.write_text(yaml.safe_dump(raw, sort_keys=False, allow_unicode=True), encoding="utf-8")
    return out


def _refresh_repricer(config: dict[str, Any], repo_root: Path, run_dir: Path, *, headless: bool) -> tuple[Path | None, dict[str, Any]]:
    from .repricer_items_export import export_repricer_items_to_sqlite

    scoped_config = _write_scoped_repricer_config(config, repo_root, run_dir)
    scoped_raw = yaml.safe_load(scoped_config.read_text(encoding="utf-8")) or {}
    token_keys = [str(account.get("token_env") or "").strip() for account in scoped_raw.get("accounts", [])]
    token_keys = [key for key in token_keys if key]
    missing = _load_required_env_keys_from_files(token_keys, _configured_env_files(config, repo_root))
    if missing:
        raise UniversalMGroupWatchdogError("missing Repricer env var(s): " + ", ".join(sorted(set(missing))))
    db_path = run_dir / "repricer_items.sqlite"
    summary = export_repricer_items_to_sqlite(
        config_path=scoped_config,
        output_path=db_path,
        headless=headless,
        include_all_rows=True,
    )
    _write_json(run_dir / "repricer_refresh_summary.json", summary)
    return db_path, summary


def _refresh_kaspi(config: dict[str, Any], repo_root: Path, run_dir: Path, *, headless: bool) -> tuple[dict[str, dict[str, Path]], dict[str, Any]]:
    from .kaspi_pricelist_download import run_kaspi_pricelist_download

    env_files = _configured_env_files(config, repo_root)
    out_paths: dict[str, dict[str, Path]] = {}
    summaries: dict[str, Any] = {}
    for store_name in ALLOWED_STORES:
        creds = _resolve_store_credentials_from_env_files(store_name, env_files)
        store_dir = run_dir / f"kaspi_{store_name.lower().replace('-', '')}"
        summary = run_kaspi_pricelist_download(
            store_name=creds["store_name"],
            email=creds["email"],
            password=creds["password"],
            run_dir=store_dir,
            headless=headless,
        )
        summaries[store_name] = summary
        if summary.get("status") != "success":
            raise UniversalMGroupWatchdogError(f"Kaspi pricelist refresh failed for {store_name}: {summary.get('error')}")
        downloads = {item["sale_state"]: Path(item["saved_path"]) for item in summary.get("downloads", [])}
        out_paths[store_name] = {"ACTIVE": downloads["ACTIVE"], "ARCHIVE": downloads["ARCHIVE"]}
    _write_json(run_dir / "kaspi_refresh_summary.json", summaries)
    return out_paths, summaries


def run_watchdog(
    *,
    config_path: str | Path = DEFAULT_CONFIG_PATH,
    repo_root: Path | None = None,
    run_root: str | Path | None = None,
    timestamp: str | None = None,
    refresh_kaspi: bool = False,
    refresh_repricer: bool = False,
    headless: bool = True,
) -> dict[str, Any]:
    root = (repo_root or Path.cwd()).resolve()
    config = load_watchdog_config(config_path, repo_root=root)
    sources = resolve_watchdog_sources(config, root)
    generated = _now_almaty()
    ts = timestamp or generated.strftime("%Y%m%d_%H%M%S")
    output_root = _repo_path(run_root or (config.get("report") or {}).get("run_root") or DEFAULT_RUN_ROOT, root)
    run_dir = output_root / ts
    run_dir.mkdir(parents=True, exist_ok=False)

    source_errors: list[str] = []
    pricelist_paths = _load_pricelist_paths(config, root, sources)
    if refresh_kaspi:
        try:
            pricelist_paths, _ = _refresh_kaspi(config, root, run_dir, headless=headless)
        except Exception as exc:
            source_errors.append(str(exc))
    repricer_sqlite: Path | None = None
    if sources.get("repricer_sqlite_path"):
        repricer_sqlite = Path(sources["repricer_sqlite_path"])
    if refresh_repricer:
        try:
            repricer_sqlite, _ = _refresh_repricer(config, root, run_dir, headless=headless)
        except Exception as exc:
            source_errors.append(str(exc))

    summary = build_watchdog_report(
        config=config,
        repo_root=root,
        pricelist_paths=pricelist_paths,
        repricer_sqlite_path=repricer_sqlite,
        run_dir=run_dir,
        generated_at=generated.isoformat(),
    )
    if source_errors:
        summary["source_errors"] = source_errors
        if summary["gate"] == "GREEN":
            summary["gate"] = "YELLOW"
            summary["status"] = "warning"
        _write_json(run_dir / "WATCHDOG_SUMMARY.json", summary)
        findings = _read_csv_rows(run_dir / "WATCHDOG_FINDINGS.csv")
        (run_dir / "WATCHDOG_REPORT.md").write_text(_render_report(summary, findings), encoding="utf-8")
    return summary


def validate_watchdog_config(*, config_path: str | Path = DEFAULT_CONFIG_PATH, repo_root: Path | None = None) -> dict[str, Any]:
    root = (repo_root or Path.cwd()).resolve()
    config = load_watchdog_config(config_path, repo_root=root)
    return {
        "schema_version": SCHEMA_VERSION,
        "status": "ok",
        "gate": "GREEN",
        "stores": [normalize_store_name(store.get("store_name")) for store in config.get("stores", [])],
        "auto_correction_enabled": False,
        "production_write_action_authorized": False,
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(prog="python -m web_auto.universal_storeb_watchdog")
    parser.add_argument("--json", action="store_true", help="Print JSON output")
    subparsers = parser.add_subparsers(dest="command", required=True)

    validate_parser = subparsers.add_parser("validate", help="Validate watchdog config")
    validate_parser.add_argument("--config", default=str(DEFAULT_CONFIG_PATH))

    run_parser = subparsers.add_parser("run", help="Run watchdog report")
    run_parser.add_argument("--config", default=str(DEFAULT_CONFIG_PATH))
    run_parser.add_argument("--run-root", default=None)
    run_parser.add_argument("--timestamp", default=None)
    run_parser.add_argument("--refresh-kaspi", action="store_true", help="Download fresh Kaspi ACTIVE/ARCHIVE state")
    run_parser.add_argument("--refresh-repricer", action="store_true", help="Export fresh scoped Repricer state")
    run_parser.add_argument("--headless", action="store_true", help="Use headless browser for refresh operations")
    run_parser.add_argument("--headed", action="store_true", help="Use headed browser for refresh operations")

    args = parser.parse_args(argv)
    try:
        if args.command == "validate":
            payload = validate_watchdog_config(config_path=args.config)
        else:
            headless = True
            if args.headed:
                headless = False
            if args.headless:
                headless = True
            payload = run_watchdog(
                config_path=args.config,
                run_root=args.run_root,
                timestamp=args.timestamp,
                refresh_kaspi=args.refresh_kaspi,
                refresh_repricer=args.refresh_repricer,
                headless=headless,
            )
    except Exception as exc:
        payload = {"schema_version": SCHEMA_VERSION, "status": "failed", "gate": "RED", "error": str(exc)}
        if args.json:
            print(json.dumps(payload, ensure_ascii=False, indent=2))
        else:
            print(str(exc))
        return 4
    if args.json:
        print(json.dumps(payload, ensure_ascii=False, indent=2, default=str))
    else:
        print(f"status={payload.get('status')} gate={payload.get('gate')} run_dir={payload.get('run_dir', '')}")
    return 4 if payload.get("gate") == "RED" else 0


if __name__ == "__main__":
    raise SystemExit(main())
